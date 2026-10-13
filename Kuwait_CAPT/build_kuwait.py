"""
Kuwait CAPT winning bids 2024-2026 -> KW_CAPT_Awards.xlsx. VOLUME SIGNALS ONLY (Phase 4 rule): awards are tender-level
(subject, organisation, winner, total KD); no item lines or unit prices, so no line-level value is built.
Arabic subjects are translated (translate.googleapis.com, cached) and mapped AR-first then EN with the vocab layer.
Health = organisation or subject mentions health / medical / hospital (EN or AR).
"""
import glob, html, json, os, re, sys
from collections import Counter

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))
from vocab_rules import Vocab  # noqa: E402
from translate import translate_many  # noqa: E402

HEALTH = re.compile(r"health|medical|hospital|صحة|صحي|طبي|مستشفى", re.I)
KWD_PER_USD = 0.307  # approx. (KWD is pegged to a basket; ~0.305-0.308 per USD in 2024-26)


def kd(s):
    m = re.search(r"([\d,]+\.\d+|\d[\d,]*)", s or "")
    return float(m.group(1).replace(",", "")) if m else None


def popup_items(pid):
    f = os.path.join(HERE, "capt_popups", re.sub(r"[^\w-]", "_", pid) + ".html")
    if not os.path.exists(f):
        return []
    s = open(f, encoding="utf8", errors="ignore").read()
    out = []
    for r in re.findall(r'<div class="table-row tbody">(.*?)</div>\s*</div>', s, re.S):
        c = [re.sub(r"\s+", " ", html.unescape(re.sub(r"<[^>]+>", " ", x))).strip() for x in re.findall(r'<div class="table-cell[^"]*">(.*?)</div>', r + "</div>", re.S)]
        if len(c) >= 3:
            out.append({"item": c[1], "contractor": c[2]})
    return out


def main():
    v = Vocab()
    data = json.load(open(os.path.join(HERE, "capt_awards.json"), encoding="utf8"))
    rows = [dict(r, month=k) for k, rs in sorted(data.items()) for r in rs]
    en = dict(zip([r["subject"] for r in rows], translate_many([r["subject"] for r in rows], src="ar")))
    out = []
    for r in rows:
        subj_en = en.get(r["subject"], "")
        pm = v.product(r["subject"]) or v.product(subj_en) or {}
        health = bool(HEALTH.search(r["org"] + " " + r["subject"] + " " + subj_en))
        val = kd(r["total"])
        items = popup_items(r["popup"]) if r["has_items"] else []
        out.append({"Month": r["month"], "Meeting Date": r["date"], "Tender No": r["tender_no"], "Organisation": r["org"],
                    "Subject (AR)": r["subject"], "Subject (EN, machine)": subj_en, "Winner": r["winner"],
                    "Total KD": val, "Total USD (approx)": round(val / KWD_PER_USD, 0) if val else None,
                    "Item-level winners (popup)": "; ".join(f"{i['item']} -> {i['contractor']}" for i in items)[:2000],
                    "Health-related": "Yes" if health else "", "OU": pm.get("ou"), "Sub-OU": pm.get("sub_ou"), "Device Category": pm.get("product_category"),
                    "Tender Stage": "Awarded", "Price Evidence": "Contract-level only", "Manufacturer Evidence": "None",
                    "Use": "Volume signal only (Phase 4)"})
    hl = [o for o in out if o["Health-related"]]
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    hdr = PatternFill("solid", fgColor="1F4E78")

    def sheet(title, data, widths):
        ws = wb.create_sheet(title)
        cols = list(data[0].keys())
        ws.append(cols)
        for c in ws[1]:
            c.font, c.fill, c.alignment = Font(bold=True, color="FFFFFF"), hdr, Alignment(wrap_text=True)
        for r in data:
            ws.append([r[c] for c in cols])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(data)+1}"
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    ws = wb.create_sheet("README")
    by = Counter(o["Month"][:4] for o in out)
    for line in [
        "Kuwait CAPT (Central Agency for Public Tenders) winning bids, meeting dates 2024-01 .. 2026-09",
        f"Awards: {len(out)} (by year {dict(by)}); health-related: {len(hl)}; mapped to a taxonomy category: {sum(1 for o in out if o['Device Category'])}.",
        "VOLUME SIGNAL ONLY: tender-level award totals in KD; no item quantities or unit prices; no line-level value is built (Phase 4 rule).",
        "CAPT handles the large tenders; MOH's own direct purchases/practices are not on CAPT. Jan-May 2024 has almost no published awards on CAPT.",
        "Subjects are Arabic, machine-translated for mapping. KD->USD at ~0.307 KWD/USD (approximate).",
        "Accessed via the user's Edge (Cloudflare) after the user approved the site's terms-of-use popup.",
    ]:
        ws.append([line])
    ws.column_dimensions["A"].width = 150
    sheet("Health_Awards", hl or out[:1], [8, 13, 18, 30, 50, 50, 40, 14, 14, 60, 6, 25, 30, 35, 9, 18, 12, 22])
    sheet("All_Awards", out, [8, 13, 18, 30, 50, 50, 40, 14, 14, 60, 6, 25, 30, 35, 9, 18, 12, 22])
    wb.save(os.path.join(HERE, "KW_CAPT_Awards.xlsx"))
    print(f"awards {len(out)}; health {len(hl)}; mapped {sum(1 for o in out if o['Device Category'])}")
    for o in hl[:40]:
        print(" ", o["Meeting Date"], "|", o["Organisation"][:28], "|", o["Subject (EN, machine)"][:80], "|", o["Total KD"], "|", o["Device Category"])


if __name__ == "__main__":
    main()
