from __future__ import annotations

import html
import json
import re
import subprocess
from datetime import date
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from urllib.parse import quote, unquote

import requests

from harness.shared.normalize import normalize_text, utc_now_iso
from harness.shared.schema import OutputRow
from harness.shared.text_extract import first_match, strip_html_tags
from harness.sources.base import SmokeProbeResult
from harness.sources.registry import SOURCES

CONFIG = SOURCES["czechia_nen"]
NOTICE_PATH = Path("evidence") / "czechia_notice_N006-23-V00034551.html"
RESULT_PATH = Path("evidence") / "czechia_result_2217224876.html"
LIVE_BASE_URL = "https://nen.nipez.cz"
LIVE_LISTING_PATH = "/en/verejne-zakazky"
LIVE_HEADERS = {
    "User-Agent": "Mozilla/5.0",
    "Accept-Language": "en-US,en;q=0.9",
}
LIVE_MEDICAL_KEYWORDS = (
    "lůžko",
    "zdravot",
    "nemocnic",
    "lék",
    "lékař",
    "laborator",
    "diagnost",
    "steriliz",
    "ultrazv",
    "rentgen",
    "rehabilit",
    "krev",
    "ambul",
)
LIVE_MEDICAL_TERMS = (
    "zdravot",
    "nemocnic",
    "lék",
    "lékař",
    "laborator",
    "diagnost",
    "steriliz",
    "ultrazv",
    "rentgen",
    "rehabilit",
    "krev",
    "ambul",
    "medical",
    "medic",
    "hospital",
    "clinic",
    "pharma",
)
LIVE_ITEM_UOM_PATTERNS = (
    r"ks",
    r"kus",
    r"kusy",
    r"sada",
    r"sady",
    r"set",
    r"bal",
    r"balení",
    r"lahv",
    r"flakon",
    r"amp",
    r"ml",
    r"l",
    r"g",
    r"kg",
    r"mg",
    r"cm",
    r"mm",
    r"m2",
    r"m3",
)


def _load_text(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="ignore")


def _fetch_live_text(path: str) -> str:
    url = f"{LIVE_BASE_URL}{path}"
    try:
        response = requests.get(url, headers=LIVE_HEADERS, timeout=60)
        response.raise_for_status()
        return response.text
    except requests.RequestException:
        result = subprocess.run(
            ["curl.exe", "-L", "--max-time", "60", url],
            check=True,
            capture_output=True,
        )
        return result.stdout.decode("utf-8", errors="ignore")


def _extract_initial_redux_state(text: str) -> dict:
    encoded = first_match(r'<meta name="initialReduxState" content="([^"]+)"', text, re.S)
    if not encoded:
        return {}
    decoded = unquote(html.unescape(encoded))
    try:
        return json.loads(decoded)
    except json.JSONDecodeError:
        return {}


def _build_live_listing_path(keyword: str, date_from: str, date_to: str, page: int = 1) -> str:
    filter_value = f"query={quote(keyword)}&stavZP=zadana&datumPrvniUver={date_from}%2C"
    if page > 1:
        filter_value = f"{filter_value}&page={page}"
    return f"{LIVE_LISTING_PATH}/p:vz:{filter_value}"


def _normalize_notice_code(code: str) -> str:
    normalized = normalize_text(code)
    return normalized.replace("/", "-")


def _extract_last_listing_page(text: str) -> int:
    pages = [int(match) for match in re.findall(r"[?&]page=(\d+)\b", text)]
    return max(pages, default=1)


def _parse_cz_date(value: str) -> date | None:
    match = re.search(r"(\d{1,2})\.\s*(\d{1,2})\.\s*(\d{4})", value)
    if not match:
        return None
    day, month, year = [int(part) for part in match.groups()]
    try:
        return date(year, month, day)
    except ValueError:
        return None


def _within_date_window(value: str, date_from: str, date_to: str) -> bool:
    parsed = _parse_cz_date(value)
    if not parsed:
        return True
    start = date.fromisoformat(date_from)
    end = date.fromisoformat(date_to)
    return start <= parsed <= end


def _parse_decimal_text(value: str) -> Decimal | None:
    cleaned = normalize_text(value).replace("\xa0", "").replace(" ", "").replace(",", ".")
    if not cleaned:
        return None
    try:
        return Decimal(cleaned)
    except InvalidOperation:
        return None


def _format_decimal_text(value: Decimal) -> str:
    quantized = value.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    return f"{quantized:.2f}".replace(".", ",")


def _extract_live_listing_rows(text: str) -> list[dict[str, str]]:
    redux = _extract_initial_redux_state(text)
    collection = (
        redux.get("collectionStore", {})
        .get("collections", {})
        .get("verejne-zakazky-seznam", {})
        .get("collection", [])
    )
    rows: list[dict[str, str]] = []
    for item in collection:
        rows.append(
            {
                "notice_code": normalize_text(item.get("kod")),
                "title": normalize_text(item.get("nazev")),
                "status": normalize_text(item.get("stavZP")),
                "closing_date": normalize_text(item.get("podaniLhuta")),
            }
        )
    return rows


def _extract_live_notice_codes(text: str) -> list[str]:
    codes: list[str] = []
    for href in re.findall(r'href="([^"]*detail-zakazky/N006-[^"]+)"', text):
        match = re.search(r"detail-zakazky/(N006-\d{2}-V\d+)", html.unescape(href))
        if not match:
            continue
        code = normalize_text(match.group(1))
        if code and code not in codes:
            codes.append(code)
    return codes


def _extract_live_notice_refs(text: str) -> list[tuple[str, str]]:
    refs: list[tuple[str, str]] = []
    seen: set[str] = set()
    for href in re.findall(r'href="([^"]*detail-zakazky/N006-[^"]+)"', text):
        decoded_href = html.unescape(href)
        match = re.search(r"detail-zakazky/(N006-\d{2}-V\d+)", decoded_href)
        if not match:
            continue
        code = normalize_text(match.group(1))
        if not code or code in seen:
            continue
        seen.add(code)
        refs.append((code, decoded_href))
    return refs


def _extract_live_notice_fields(text: str, notice_code: str) -> dict[str, str]:
    title = normalize_text(
        first_match(r'PROCUREMENT PROCEDURE NAME</h3><p title="([^"]*)"', text, re.S)
        or first_match(r"<title>(.*?)</title>", text, re.S)
    )
    publication_date = normalize_text(
        first_match(r'DATE OF PUBLICATION ON PROFILE</h3><p title="([^"]*)"', text, re.S)
    )
    closing_date = normalize_text(
        first_match(r'DEADLINE FOR SUBMITTING TENDERS</h3><p title="([^"]*)"', text, re.S)
    )
    description = normalize_text(
        html.unescape(first_match(r'SUBJECT-MATTER DESCRIPTION</h3><p title="([^"]*)"', text, re.S))
    )
    subject_name = normalize_text(
        html.unescape(first_match(r'SUBJECT-MATTER NAME</h3><p title="([^"]*)"', text, re.S))
    )
    buyer = normalize_text(
        html.unescape(first_match(r"Contracting authority:\s*(.*?)\.\s*SUBJECT-MATTER DESCRIPTION:", text, re.S))
    )
    cpv_code = normalize_text(first_match(r'Code from the CPV code list</h3><p title="([^"]*)"', text, re.S))
    cpv_name = normalize_text(first_match(r'Name from the CPV code list</h3><p title="([^"]*)"', text, re.S))
    nipez_code = normalize_text(first_match(r'Code from the NIPEZ code list</h3><p title="([^"]*)"', text, re.S))
    nipez_name = normalize_text(first_match(r'Name from the NIPEZ code list</h3><p title="([^"]*)"', text, re.S))
    amount = normalize_text(first_match(r'Estimated value \(excl\. VAT\)</h3><p title="([^"]*)"', text, re.S))
    currency = normalize_text(first_match(r'Currency</h3><p title="([^"]*)"', text, re.S))
    classification_parts = [part for part in [cpv_code, cpv_name, nipez_code, nipez_name] if part]
    classification = ", ".join(dict.fromkeys(classification_parts))

    item_quantity = ""
    item_uom = ""
    item_desc = subject_name or title
    quantity_match = re.search(
        rf"\b(\d+(?:[.,]\d+)?)\s+((?:{'|'.join(LIVE_ITEM_UOM_PATTERNS)}))\b",
        description,
        re.I,
    )
    if quantity_match:
        item_quantity = normalize_text(quantity_match.group(1))
        item_uom = normalize_text(quantity_match.group(2))

    return {
        "title": title,
        "buyer": buyer,
        "notice_id": notice_code,
        "publication_date": publication_date,
        "closing_date": closing_date,
        "description": description,
        "classification": classification,
        "amount": amount,
        "currency": currency,
        "item_no": "1" if item_desc else "",
        "item_desc": item_desc,
        "item_uom": item_uom,
        "item_quantity": item_quantity,
    }


def _extract_live_contracts(result_text: str) -> list[dict[str, str]]:
    pattern = re.compile(
        r'detail-uverejneni/(?P<detail_id>\d+)".*?'
        r'data-title="CONTRACT REFERENCE NUMBER">(?P<contract_ref>.*?)</td>.*?'
        r'data-title="DATE OF CONTRACT CONCLUSION">(?P<contract_date>.*?)</td>.*?'
        r'data-title="Official name">(?P<supplier>.*?)</td>.*?'
        r'data-title="Contractual price incl\. VAT">(?P<with_vat>.*?)</td>.*?'
        r'data-title="Contractual price excl\. VAT">(?P<without_vat>.*?)</td>.*?'
        r'data-title="CONTRACTUAL PRICE WITH ADDITIONS INCL\. VAT">(?P<updated_with_vat>.*?)</td>.*?'
        r'data-title="CONTRACTUAL PRICE WITH ADDITIONS EXCL\. VAT">(?P<updated_without_vat>.*?)</td>.*?'
        r'data-title="Currency">(?P<currency>.*?)</td>',
        re.S,
    )
    rows: list[dict[str, str]] = []
    for match in pattern.finditer(result_text):
        rows.append(
            {
                "contract_id": normalize_text(match.group("detail_id")),
                "contract_ref": normalize_text(strip_html_tags(match.group("contract_ref"))),
                "contract_period": normalize_text(strip_html_tags(match.group("contract_date"))),
                "supplier": normalize_text(html.unescape(strip_html_tags(match.group("supplier")))),
                "awarded_currency": normalize_text(strip_html_tags(match.group("currency"))),
                "awarded_value": normalize_text(strip_html_tags(match.group("without_vat"))),
                "item_award": normalize_text(strip_html_tags(match.group("updated_with_vat")))
                or normalize_text(strip_html_tags(match.group("with_vat"))),
            }
        )
    return rows


def _is_medical_notice(notice: dict[str, str]) -> bool:
    text = normalize_text(
        " ".join(
            [
                notice.get("title", ""),
                notice.get("description", ""),
                notice.get("classification", ""),
                notice.get("buyer", ""),
                notice.get("item_desc", ""),
            ]
        )
    ).lower()
    return any(term.lower() in text for term in LIVE_MEDICAL_TERMS)


def _derive_item_unit_price(quantity_text: str, awarded_value_text: str) -> str:
    quantity = _parse_decimal_text(quantity_text)
    awarded_value = _parse_decimal_text(awarded_value_text)
    if not quantity or not awarded_value or quantity <= 0:
        return ""
    return _format_decimal_text(awarded_value / quantity)


def _extract_notice_fields(text: str) -> dict[str, str]:
    title = normalize_text(first_match(r"<h1 class=\"\">(.*?)</h1>", text, re.S))
    buyer = normalize_text(
        strip_html_tags(
            first_match(r"<h3 class=\"gov-title--delta\">Zadavatel</h3><a class=\"gov-link\"[^>]*>(.*?)</a>", text, re.S)
        )
    )
    notice_id = normalize_text(
        first_match(r"<h3 class=\"gov-title--delta\">Systémové číslo NEN</h3><p title=\"([^\"]+)\"", text, re.S)
    )
    amount = normalize_text(
        first_match(r"<h3 class=\"gov-title--delta\">Předpokládaná hodnota \(bez DPH\)</h3><p title=\"([^\"]+)\"", text, re.S)
    )
    currency = normalize_text(
        first_match(r"<h3 class=\"gov-title--delta\">Měna</h3><p title=\"([^\"]+)\"", text, re.S)
    )
    publication_date = normalize_text(
        first_match(r"<h3 class=\"gov-title--delta\">Datum uveřejnění ZP na profil</h3><p title=\"([^\"]+)\"", text, re.S)
    )
    closing_date = normalize_text(
        first_match(r"<h3 class=\"gov-title--delta\">Lhůta pro podání nabídek</h3><p title=\"([^\"]+)\"", text, re.S)
    )
    description = normalize_text(
        html.unescape(first_match(r"<h3 class=\"gov-title--delta\">Popis předmětu</h3><p title=\"(.*?)\" class=\"text gov-note\">", text, re.S))
    )
    vvz = normalize_text(
        first_match(r"<h3 class=\"gov-title--delta\">Evidenční číslo zakázky ve VVZ</h3><p title=\"([^\"]+)\"", text, re.S)
    )
    return {
        "title": title,
        "buyer": buyer,
        "notice_id": notice_id,
        "amount": amount,
        "currency": currency,
        "publication_date": publication_date,
        "closing_date": closing_date,
        "description": description,
        "classification": vvz,
    }


def _extract_item_fields(text: str) -> dict[str, str]:
    qty_match = re.search(r"jednorázová dodávka\s+(\d+)\s+ks\s+sterilizačních kontejnerů", html.unescape(text), re.I)
    item_quantity = qty_match.group(1) if qty_match else ""
    return {
        "item_no": "1" if item_quantity else "",
        "item_desc": "Kontejner sterilizační vč. sterilizačního síta",
        "item_uom": "ks" if item_quantity else "",
        "item_quantity": item_quantity,
    }


def _extract_contracts(result_text: str) -> list[dict[str, str]]:
    encoded = first_match(r'<meta name="initialReduxState" content="([^"]+)"', result_text, re.S)
    decoded = unquote(html.unescape(encoded))
    pattern = re.compile(
        r'"id":(?P<id>\d+),"cisloSmlouva":"(?P<contract>[^"]+)","datumUzavreni":"(?P<date>[^"]+)",'
        r'"uredniNazev":"(?P<name>[^"]+)","smluvniCenaS":(?P<with_vat>\d+),"smluvniCenaBez":(?P<without_vat>\d+).*?'
        r'"menaUverKod":"(?P<currency>[A-Z]+)"'
    )
    rows = []
    for match in pattern.finditer(decoded):
        rows.append(
            {
                "contract_id": match.group("id"),
                "contract_period": normalize_text(match.group("date")),
                "supplier": normalize_text(match.group("name")),
                "awarded_currency": normalize_text(match.group("currency")),
                "awarded_value": normalize_text(match.group("without_vat")),
                "item_award": normalize_text(match.group("with_vat")),
            }
        )
    return rows


def map_notice_and_result_to_rows(query_text: str = "") -> list[dict[str, str]]:
    notice_text = _load_text(NOTICE_PATH)
    result_text = _load_text(RESULT_PATH)
    notice = _extract_notice_fields(notice_text)
    item = _extract_item_fields(notice_text)
    contracts = _extract_contracts(result_text)

    rows: list[dict[str, str]] = []
    for index, contract in enumerate(contracts, start=1):
        row = OutputRow(
            source=CONFIG.source,
            country=CONFIG.country,
            country_cc=CONFIG.country_cc,
            publication_date=notice["publication_date"],
            closing_date=notice["closing_date"],
            title=notice["title"],
            description=notice["description"],
            buyer=notice["buyer"],
            classification=notice["classification"],
            status="Awarded",
            currency=notice["currency"],
            amount=notice["amount"],
            awarding_supplier_name=contract["supplier"],
            awarded_currency=contract["awarded_currency"],
            awarded_value=contract["awarded_value"],
            contract_period=contract["contract_period"],
            item_no=item["item_no"],
            item_desc=item["item_desc"],
            item_uom=item["item_uom"],
            item_quantity=item["item_quantity"],
            item_award=contract["item_award"],
            notice_id=notice["notice_id"],
            notice_url="https://nen.nipez.cz/verejne-zakazky/detail-zakazky/N006-23-V00034551",
            query_text=query_text,
            scraped_at=utc_now_iso(),
            dedup_key=f"{CONFIG.key}|{notice['notice_id']}|{index}|{contract['supplier']}",
        )
        rows.append(row.to_dict())
    return rows


def smoke_probe() -> SmokeProbeResult:
    rows = map_notice_and_result_to_rows(query_text="fixture:N006-23-V00034551")
    ok = bool(rows) and all(
        [
            rows[0]["notice_id"],
            rows[0]["title"],
            rows[0]["buyer"],
            rows[0]["awarding_supplier_name"],
            rows[0]["awarded_currency"],
            rows[0]["item_quantity"],
        ]
    )
    return SmokeProbeResult(
        source_key=CONFIG.key,
        ok=ok,
        mode="fixture",
        detail="Mapped one public NEN notice plus result page into normalized awarded rows.",
        sample_rows=rows,
    )


def live_medical_awarded_rows(date_from: str, date_to: str, page_size: int = 50) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    seen_notice_codes: set[str] = set()

    for keyword in LIVE_MEDICAL_KEYWORDS:
        first_listing_path = _build_live_listing_path(keyword, date_from, date_to, page=1)
        try:
            first_listing_text = _fetch_live_text(first_listing_path)
        except Exception:
            continue
        last_page = _extract_last_listing_page(first_listing_text)

        for page in range(1, last_page + 1):
            listing_path = _build_live_listing_path(keyword, date_from, date_to, page=page)
            try:
                listing_text = first_listing_text if page == 1 else _fetch_live_text(listing_path)
            except Exception:
                continue
            notice_refs = _extract_live_notice_refs(listing_text)
            for notice_code, detail_path in notice_refs:
                if not notice_code or notice_code in seen_notice_codes:
                    continue
                seen_notice_codes.add(notice_code)

                result_path = f"{detail_path}/vysledek"
                try:
                    detail_text = _fetch_live_text(detail_path)
                    result_text = _fetch_live_text(result_path)
                except Exception:
                    continue

                notice = _extract_live_notice_fields(detail_text, notice_code=notice_code)
                if not _within_date_window(notice["publication_date"], date_from, date_to):
                    continue
                if not _is_medical_notice(notice):
                    continue

                contracts = _extract_live_contracts(result_text)
                if not contracts:
                    continue

                query_text = f"live-medical-awarded:{keyword}:{date_from}..{date_to}:page={page}"
                for contract in contracts:
                    row = OutputRow(
                        source=CONFIG.source,
                        country=CONFIG.country,
                        country_cc=CONFIG.country_cc,
                        publication_date=notice["publication_date"],
                        closing_date=notice["closing_date"],
                        title=notice["title"],
                        description=notice["description"],
                        buyer=notice["buyer"],
                        classification=notice["classification"],
                        status="Awarded",
                        currency=notice["currency"],
                        amount=notice["amount"],
                        awarding_supplier_name=contract["supplier"],
                        awarded_currency=contract["awarded_currency"],
                        awarded_value=contract["awarded_value"],
                        contract_period=contract["contract_period"],
                        item_no=notice["item_no"],
                        item_desc=notice["item_desc"],
                        item_uom=notice["item_uom"],
                        item_quantity=notice["item_quantity"],
                        item_unit_price=_derive_item_unit_price(notice["item_quantity"], contract["awarded_value"]),
                        item_award=contract["item_award"],
                        notice_id=notice["notice_id"],
                        notice_url=f"{LIVE_BASE_URL}{detail_path}",
                        query_text=query_text,
                        scraped_at=utc_now_iso(),
                        dedup_key=f"{CONFIG.key}|{notice['notice_id']}|{contract['contract_ref'] or contract['contract_id']}|{contract['supplier']}",
                    )
                    rows.append(row.to_dict())

    return rows


def cached_live_medical_awarded_rows(
    listing_paths: list[Path],
    cache_dir: Path,
    date_from: str = "",
    date_to: str = "",
    query_text_prefix: str = "cached-live-medical-awarded",
) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    seen_row_keys: set[str] = set()

    for listing_path in listing_paths:
        listing_text = listing_path.read_text(encoding="utf-8", errors="ignore")
        for notice_code, _detail_href in _extract_live_notice_refs(listing_text):
            detail_path = cache_dir / f"{notice_code}_detail.html"
            result_path = cache_dir / f"{notice_code}_result.html"
            if not detail_path.exists() or not result_path.exists():
                continue

            detail_text = detail_path.read_text(encoding="utf-8", errors="ignore")
            result_text = result_path.read_text(encoding="utf-8", errors="ignore")
            notice = _extract_live_notice_fields(detail_text, notice_code=notice_code)
            if date_from and date_to and not _within_date_window(notice["publication_date"], date_from, date_to):
                continue
            if not _is_medical_notice(notice):
                continue

            contracts = _extract_live_contracts(result_text)
            if not contracts:
                continue

            for contract in contracts:
                dedup_key = (
                    f"{CONFIG.key}|{notice['notice_id']}|"
                    f"{contract['contract_ref'] or contract['contract_id']}|{contract['supplier']}"
                )
                if dedup_key in seen_row_keys:
                    continue
                seen_row_keys.add(dedup_key)

                row = OutputRow(
                    source=CONFIG.source,
                    country=CONFIG.country,
                    country_cc=CONFIG.country_cc,
                    publication_date=notice["publication_date"],
                    closing_date=notice["closing_date"],
                    title=notice["title"],
                    description=notice["description"],
                    buyer=notice["buyer"],
                    classification=notice["classification"],
                    status="Awarded",
                    currency=notice["currency"],
                    amount=notice["amount"],
                    awarding_supplier_name=contract["supplier"],
                    awarded_currency=contract["awarded_currency"],
                    awarded_value=contract["awarded_value"],
                    contract_period=contract["contract_period"],
                    item_no=notice["item_no"],
                    item_desc=notice["item_desc"],
                    item_uom=notice["item_uom"],
                    item_quantity=notice["item_quantity"],
                    item_unit_price=_derive_item_unit_price(notice["item_quantity"], contract["awarded_value"]),
                    item_award=contract["item_award"],
                    notice_id=notice["notice_id"],
                    notice_url=f"{LIVE_BASE_URL}/en/verejne-zakazky/detail-zakazky/{notice['notice_id']}",
                    query_text=f"{query_text_prefix}:{listing_path.name}",
                    scraped_at=utc_now_iso(),
                    dedup_key=dedup_key,
                )
                rows.append(row.to_dict())

    return rows


def workbook_medical_awarded_rows(workbook_path: Path) -> list[dict[str, str]]:
    from openpyxl import load_workbook

    accepted_log_by_notice: dict[str, dict[str, str]] = {}
    log_path = workbook_path.with_name("scrape_log.jsonl")
    if log_path.exists():
        for line in log_path.read_text(encoding="utf-8", errors="ignore").splitlines():
            try:
                payload = json.loads(line)
            except json.JSONDecodeError:
                continue
            if payload.get("stage") != "accepted":
                continue
            notice_key = _normalize_notice_code(payload.get("nen_url_id", ""))
            if notice_key:
                accepted_log_by_notice[notice_key] = {
                    "publication_date": normalize_text(payload.get("publication_date")),
                    "status": normalize_text(payload.get("status")),
                }

    wb = load_workbook(workbook_path, data_only=True, read_only=True)
    ws = wb.active
    row_iter = ws.iter_rows(values_only=True)
    header_row = next(row_iter, None)
    if not header_row:
        return []

    headers = [normalize_text(value) for value in header_row]
    header_index = {name: idx for idx, name in enumerate(headers) if name}

    def cell(values: tuple[object, ...], name: str) -> str:
        idx = header_index.get(name)
        if idx is None or idx >= len(values):
            return ""
        value = values[idx]
        return normalize_text("" if value is None else str(value))

    rows: list[dict[str, str]] = []
    for values in row_iter:
        title = cell(values, "title")
        buyer = cell(values, "buyer")
        item_desc = cell(values, "item_description")
        notice_id = cell(values, "notice_id")
        notice_url = cell(values, "notice_url")
        if not any([title, buyer, item_desc, notice_id, notice_url]):
            continue

        classification_parts = [
            part
            for part in [
                cell(values, "classification_code"),
                cell(values, "classification"),
            ]
            if part
        ]
        classification = " | ".join(dict.fromkeys(classification_parts))
        dedup_key = cell(values, "dedup_key")
        item_no = cell(values, "item_no")
        awarding_supplier_name = cell(values, "awarding_supplier")
        log_fields = accepted_log_by_notice.get(_normalize_notice_code(notice_id), {})

        row = OutputRow(
            source=CONFIG.source,
            country=CONFIG.country,
            country_cc=CONFIG.country_cc,
            publication_date=cell(values, "publication_date") or log_fields.get("publication_date", ""),
            closing_date=cell(values, "closing_date"),
            title=title,
            description=cell(values, "description"),
            buyer=buyer,
            classification=classification,
            status=cell(values, "status") or log_fields.get("status", "") or "Awarded",
            currency=cell(values, "currency"),
            amount=cell(values, "amount"),
            awarding_supplier_name=awarding_supplier_name,
            awarded_currency="",
            awarded_value="",
            contract_period="",
            item_no=item_no,
            item_desc=item_desc,
            item_uom=cell(values, "item_uom"),
            item_quantity=cell(values, "item_quantity"),
            item_unit_price="",
            item_award="",
            notice_id=notice_id,
            notice_url=notice_url,
            query_text=cell(values, "query_text"),
            scraped_at=cell(values, "scraped_at") or utc_now_iso(),
            dedup_key=dedup_key
            or "|".join(
                part
                for part in [
                    CONFIG.key,
                    notice_id,
                    item_no,
                    awarding_supplier_name,
                    item_desc[:120],
                ]
                if part
            ),
        )
        rows.append(row.to_dict())

    return rows
