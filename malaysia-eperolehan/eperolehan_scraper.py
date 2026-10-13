import argparse
import csv
import hashlib
import io
import json
import re
import sys
from dataclasses import dataclass
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Optional
from urllib.parse import quote, urljoin

from openpyxl import Workbook
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from playwright.sync_api import Download, Error as PlaywrightError, Page, TimeoutError as PlaywrightTimeoutError, sync_playwright
from pypdf import PdfReader


URL = "https://www.eperolehan.gov.my/quotation-tender-notice"


@dataclass(frozen=True)
class SearchConfig:
    target_tab: str = "SELESAI"
    date_from: str = "01/01/2024"
    date_to: str = date.today().strftime("%d/%m/%Y")
    bidang_code: Optional[str] = "05"
    sub_bidang_code: Optional[str] = "01"
    pecahan_sub_bidang_code: Optional[str] = "01"
    headless: bool = False
    slow_mo_ms: int = 500
    timeout_ms: int = 60000
    output_xlsx: str = "eperolehan_tenders.xlsx"
    progress_json: str = "eperolehan_progress.json"
    restart_from_first_page: bool = False
    debug: bool = True
    debug_dir: str = "eperolehan_debug"
    max_pages: Optional[int] = None
    max_rows_per_page: Optional[int] = None
    country: str = "Malaysia"
    country_cc: str = "MY"
    currency: str = "MYR"
    source: str = "ePerolehan"
    status: str = "awarded"


CONFIG = SearchConfig()


OUTPUT_FIELDS = [
    "source",
    "country",
    "country_cc",
    "publication_date",
    "closing_date",
    "title",
    "title_en",
    "buyer",
    "buyer_en",
    "status",
    "currency",
    "amount",
    "supplier_name",
    "contract_period",
    "item_no",
    "item_desc",
    "item_desc_en",
    "item_uom",
    "item_quantity",
    "kod_bidang_code",
    "main_category",
    "main_category_en",
    "sub_category",
    "sub_category_en",
    "category_detail",
    "category_detail_en",
    "specification_document_name",
    "specification_text",
    "notice_id",
    "notice_url",
    "query_text",
    "scraped_at",
    "dedup_key",
    "detail_control_id",
]


def clean(text: str) -> str:
    return re.sub(r"\s+", " ", text or "").strip()


def css_id(raw_id: str) -> str:
    return f'[id="{raw_id}"]'


def css_id_suffix(raw_suffix: str) -> str:
    return f'[id$="{raw_suffix}"]'


def selector_from_target(target: str) -> str:
    if target.startswith("[") or target.startswith(".") or target.startswith("#") or target.startswith("select"):
        return target
    return css_id(target)


def debug_log(message: str) -> None:
    if CONFIG.debug:
        log_line = f"[{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}] {message}"
        try:
            print(log_line, flush=True)
        except UnicodeEncodeError:
            safe_line = log_line.encode("cp1252", errors="replace").decode("cp1252")
            print(safe_line, flush=True)


def is_target_closed_error(exc: Exception) -> bool:
    return "Target page, context or browser has been closed" in str(exc)


TRANSLATION_MAP = {
    "SEBUT HARGA": "QUOTATION",
    "PEROLEHAN": "PROCUREMENT",
    "MEMBEKAL": "SUPPLY",
    "MENGHANTAR": "DELIVER",
    "MEMASANG": "INSTALL",
    "MENGUJI": "TEST",
    "MENTAULIAH": "COMMISSION",
    "PERALATAN PERUBATAN": "MEDICAL EQUIPMENT",
    "PERALATAN HOSPITAL": "HOSPITAL EQUIPMENT",
    "PERUBATAN": "MEDICAL",
    "UBAT-UBATAN": "MEDICINES",
    "FARMASEUTIKAL": "PHARMACEUTICAL",
    "BAHAN DAN KELENGKAPAN PERUBATAN": "MEDICAL MATERIALS AND EQUIPMENT",
    "PEJABAT KESIHATAN": "HEALTH OFFICE",
    "PERGIGIAN": "DENTAL",
    "BAHAGIAN": "DIVISION",
    "KEMENTERIAN KESIHATAN": "MINISTRY OF HEALTH",
}


def translate_text_basic(text: str) -> str:
    translated = clean(text)
    for source, target in sorted(TRANSLATION_MAP.items(), key=lambda item: len(item[0]), reverse=True):
        translated = re.sub(re.escape(source), target, translated, flags=re.IGNORECASE)
    return translated


def save_debug_artifacts(page: Page, name: str) -> None:
    if not CONFIG.debug:
        return

    debug_dir = Path(CONFIG.debug_dir)
    debug_dir.mkdir(parents=True, exist_ok=True)
    safe_name = re.sub(r"[^A-Za-z0-9_.-]+", "_", name).strip("_")

    try:
        page.screenshot(path=str(debug_dir / f"{safe_name}.png"), full_page=True)
    except Exception as exc:
        debug_log(f"Could not save screenshot for {name}: {exc}")

    try:
        (debug_dir / f"{safe_name}.html").write_text(page.content(), encoding="utf-8")
    except Exception as exc:
        debug_log(f"Could not save HTML for {name}: {exc}")


def wait_for_ajax(page: Page, pause_ms: int = 1500) -> None:
    debug_log("Waiting for page/ajax to settle")
    try:
        page.wait_for_load_state("networkidle", timeout=min(CONFIG.timeout_ms, 10000))
    except PlaywrightTimeoutError:
        debug_log("Network idle wait timed out. Continuing after a short pause.")
    except PlaywrightError as exc:
        if "Target page, context or browser has been closed" in str(exc):
            debug_log("Page closed while waiting for network idle. Continuing.")
            return
        raise

    try:
        page.wait_for_timeout(pause_ms)
    except PlaywrightError as exc:
        if "Target page, context or browser has been closed" in str(exc):
            debug_log("Page closed during settle pause. Continuing.")
            return
        raise


def switch_to_target_tab(page: Page, tab_name: str) -> None:
    debug_log(f"Switching to tab: {tab_name}")
    tab = page.get_by_text(tab_name, exact=True)
    tab.wait_for(state="visible", timeout=CONFIG.timeout_ms)
    tab.click()
    wait_for_ajax(page, pause_ms=3000)
    page.wait_for_function(
        """expectedTab => {
            const active = document.querySelector('li.ui-tabs-selected');
            if (!active) return false;
            return (active.textContent || '').replace(/\\s+/g, ' ').trim() === expectedTab;
        }""",
        arg=tab_name,
        timeout=CONFIG.timeout_ms,
    )


def set_text_value(page: Page, element_id: str, value: str) -> None:
    debug_log(f"Setting text value for {element_id}: {value}")
    locator = page.locator(selector_from_target(element_id))
    locator.wait_for(state="attached", timeout=CONFIG.timeout_ms)
    locator.evaluate(
        """(element, newValue) => {
            element.removeAttribute('readonly');
            element.value = newValue;
            element.dispatchEvent(new Event('input', { bubbles: true }));
            element.dispatchEvent(new Event('change', { bubbles: true }));
        }""",
        value,
    )


def select_option_by_code(page: Page, select_id: str, code: str) -> str:
    locator = page.locator(selector_from_target(select_id))
    locator.wait_for(state="visible", timeout=CONFIG.timeout_ms)
    options = locator.locator("option").evaluate_all(
        """nodes => nodes.map(node => ({
            value: node.value,
            text: (node.textContent || '').replace(/\\s+/g, ' ').trim()
        }))"""
    )

    normalized = str(code).zfill(2)
    for option in options:
        if option["text"].startswith(f"{normalized} -"):
            locator.select_option(option["value"])
            return option["text"]

    raise ValueError(f"Could not find option code {normalized} in {select_id}. Options: {options}")


def wait_for_option_code(page: Page, select_id: str, code: str) -> None:
    normalized = str(code).zfill(2)
    page.wait_for_function(
        """([selector, expectedCode]) => {
            const select = document.querySelector(selector);
            if (!select) return false;
            return Array.from(select.options).some(option =>
                (option.textContent || '').replace(/\\s+/g, ' ').trim().startsWith(`${expectedCode} -`)
            );
        }""",
        arg=[selector_from_target(select_id), normalized],
        timeout=CONFIG.timeout_ms,
    )


def resolve_popup_id(page: Page, suffix: str) -> str:
    locator = page.locator(css_id_suffix(suffix)).last
    locator.wait_for(state="attached", timeout=CONFIG.timeout_ms)
    element_id = locator.get_attribute("id")
    if not element_id:
        raise RuntimeError(f"Could not resolve popup element id for suffix {suffix}")
    return element_id


def result_row_locator(page: Page):
    try:
        panel_id = active_results_panel_id(page)
        if panel_id:
            locator = page.locator(f"[id='{panel_id}'] tbody[id*='nbsearchresults_data'] > tr")
            if locator.count() > 0:
                return locator
    except Exception:
        pass
    return page.locator("tbody[id*='nbsearchresults_data'] > tr")


def result_title_link_locator(row):
    return row.locator("a[id*='nbsearchresults']").first


def active_results_panel_id(page: Page) -> str:
    try:
        panel_id = page.evaluate(
            """() => {
                const panels = Array.from(document.querySelectorAll("div[id*='nbresultTabs']"));
                for (const panel of panels) {
                    const style = window.getComputedStyle(panel);
                    if (style.display === 'none' || style.visibility === 'hidden') continue;
                    const rect = panel.getBoundingClientRect();
                    if (rect.width === 0 || rect.height === 0) continue;
                    if (panel.getAttribute('aria-hidden') === 'false' || panel.style.display === 'block') {
                        return panel.id || '';
                    }
                }
                return '';
            }"""
        )
        if isinstance(panel_id, str) and panel_id:
            return panel_id
    except Exception:
        pass
    return ""


def visible_paginator_state(page: Page) -> dict:
    panel_id = active_results_panel_id(page)
    try:
        state = page.evaluate(
            """panelId => {
                const scope = panelId ? document.getElementById(panelId) : document;
                if (!scope) return { index: -1, activePage: 0, panelId: panelId || '' };
                const paginators = Array.from(scope.querySelectorAll('.ui-paginator'));
                for (let index = 0; index < paginators.length; index += 1) {
                    const paginator = paginators[index];
                    const style = window.getComputedStyle(paginator);
                    if (style.display === 'none' || style.visibility === 'hidden') continue;
                    const rect = paginator.getBoundingClientRect();
                    if (rect.width === 0 || rect.height === 0) continue;

                    const activeLink = paginator.querySelector('.ui-paginator-page.ui-state-active');
                    const activeText = activeLink ? (activeLink.textContent || '').trim() : '';
                    const activePage = parseInt(activeText, 10);

                    return {
                        index,
                        activePage: Number.isNaN(activePage) ? 0 : activePage,
                        panelId: panelId || '',
                    };
                }
                return { index: -1, activePage: 0, panelId: panelId || '' };
            }""",
            panel_id,
        )
        if isinstance(state, dict):
            return state
    except Exception:
        pass
    return {"index": -1, "activePage": 0, "panelId": panel_id}


def resolve_listing_row(page: Page, listing_row: dict):
    rows = result_row_locator(page)
    row_count = rows.count()
    title = clean(listing_row.get("title", ""))
    if title:
        normalized_target = normalize_compare_text(title)
        for index in range(row_count):
            candidate = rows.nth(index)
            try:
                candidate_title = clean(result_title_link_locator(candidate).text_content() or "")
            except Exception:
                candidate_title = ""
            normalized_candidate = normalize_compare_text(candidate_title)
            if normalized_target and normalized_candidate == normalized_target:
                return candidate
        for index in range(row_count):
            candidate = rows.nth(index)
            try:
                candidate_title = clean(result_title_link_locator(candidate).text_content() or "")
            except Exception:
                candidate_title = ""
            normalized_candidate = normalize_compare_text(candidate_title)
            if normalized_target and normalized_target in normalized_candidate:
                return candidate
            if normalized_target and normalized_candidate in normalized_target and normalized_candidate:
                return candidate

    row_index = listing_row["row_index"]
    if row_index < row_count:
        return rows.nth(row_index)

    raise RuntimeError(
        f"Could not resolve listing row for page={listing_row['page_number']} "
        f"row_index={row_index} title={title!r}. Current row count={row_count}"
    )


def first_result_title(page: Page) -> str:
    rows = result_row_locator(page)
    if rows.count() == 0:
        return ""

    first_row = rows.first
    title_link = result_title_link_locator(first_row)
    if title_link.count() > 0:
        return clean(title_link.text_content() or "")

    cells = first_row.locator("td")
    if cells.count():
        return clean(cells.nth(0).text_content() or "")

    return ""


def current_results_page_number(page: Page) -> int:
    paginator_state = visible_paginator_state(page)
    try:
        current_page = page.evaluate(
            """([panelId, paginatorIndex]) => {
                const scope = panelId ? document.getElementById(panelId) : document;
                if (!scope) return 0;
                const paginators = Array.from(scope.querySelectorAll('.ui-paginator'));
                const paginator = paginatorIndex >= 0 ? paginators[paginatorIndex] : null;
                if (!paginator) return 0;
                const activePages = Array.from(paginator.querySelectorAll('.ui-paginator-page.ui-state-active'));
                for (const pageLink of activePages) {
                    const text = (pageLink.textContent || '').trim();
                    const style = window.getComputedStyle(pageLink);
                    if (style.display === 'none' || style.visibility === 'hidden') continue;
                    const pageNumber = parseInt(text, 10);
                    if (!Number.isNaN(pageNumber)) return pageNumber;
                }
                return 0;
            }""",
            [paginator_state.get("panelId", ""), paginator_state.get("index", -1)],
        )
        return int(current_page or 0)
    except Exception:
        return 0


def open_kod_bidang_popup(page: Page) -> None:
    debug_log("Opening Kod Bidang popup")
    icon = page.locator(css_id("_scNoticeBoard_WAR_NGePportlet_:form:catCodeImg"))
    icon.wait_for(state="attached", timeout=CONFIG.timeout_ms)
    icon.click(force=True)
    wait_for_ajax(page, pause_ms=2000)
    save_debug_artifacts(page, "after_kod_bidang_popup_click")
    page.locator(css_id_suffix(":catCodeLvl1")).last.wait_for(
        state="visible",
        timeout=CONFIG.timeout_ms,
    )


def apply_kod_bidang_filter(page: Page, config: SearchConfig) -> None:
    debug_log("Applying Kod Bidang filter")
    if not config.bidang_code:
        return

    open_kod_bidang_popup(page)

    search_by_code_radio = page.locator(css_id(resolve_popup_id(page, ":searchBy:0")))
    if not search_by_code_radio.is_checked():
        search_by_code_radio.check()

    lvl1_id = resolve_popup_id(page, ":catCodeLvl1")
    lvl2_id = resolve_popup_id(page, ":catCodeLvl2")
    lvl3_id = resolve_popup_id(page, ":catCodeLvl3")

    select_option_by_code(page, lvl1_id, config.bidang_code)

    if config.sub_bidang_code:
        wait_for_option_code(page, lvl2_id, config.sub_bidang_code)
        select_option_by_code(page, lvl2_id, config.sub_bidang_code)

    if config.pecahan_sub_bidang_code:
        wait_for_option_code(page, lvl3_id, config.pecahan_sub_bidang_code)
        select_option_by_code(page, lvl3_id, config.pecahan_sub_bidang_code)

    page.get_by_text("OK", exact=True).last.click()
    wait_for_ajax(page)


def expand_tarikh_iklan_panel(page: Page) -> None:
    date_from = page.locator(css_id("_scNoticeBoard_WAR_NGePportlet_:form:pubFromDt_input"))
    if date_from.is_visible():
        return

    page.locator(css_id("_scNoticeBoard_WAR_NGePportlet_:form:publishdtl_toggler")).click()
    date_from.wait_for(state="visible", timeout=CONFIG.timeout_ms)
    page.wait_for_timeout(1000)


def apply_date_filter(page: Page, config: SearchConfig) -> None:
    debug_log(f"Applying date filter: {config.date_from} -> {config.date_to}")
    expand_tarikh_iklan_panel(page)
    set_text_value(page, "_scNoticeBoard_WAR_NGePportlet_:form:pubFromDt_input", config.date_from)
    set_text_value(page, "_scNoticeBoard_WAR_NGePportlet_:form:pubToDt_input", config.date_to)


def click_search(page: Page) -> None:
    debug_log("Clicking search button")
    row_selector = "tbody[id*='nbsearchresults_data'] > tr"
    previous_first_title = first_result_title(page)

    page.locator(css_id("_scNoticeBoard_WAR_NGePportlet_:form:nbSearchButton")).click()
    wait_for_ajax(page, pause_ms=2500)

    if previous_first_title:
        try:
            page.wait_for_function(
                """([selector, previousTitle]) => {
                    const firstRow = document.querySelector(selector);
                    if (!firstRow) return true;
                    const titleAnchor = firstRow.querySelector("a[id*='nbsearchresults']");
                    const text = titleAnchor ? titleAnchor.textContent : firstRow.textContent;
                    return (text || '').replace(/\\s+/g, ' ').trim() !== previousTitle;
                }""",
                arg=[row_selector, previous_first_title],
                timeout=5000,
            )
        except PlaywrightTimeoutError:
            debug_log("First result did not change after search; continuing because the site may reuse the same table content")

    result_row_locator(page).first.wait_for(state="attached", timeout=CONFIG.timeout_ms)
    debug_log(f"Search results attached. Result rows found: {result_row_locator(page).count()}. First title: {first_result_title(page)}")
    save_debug_artifacts(page, "after_search_success")


def extract_listing_rows_from_current_page(page: Page, current_page: int) -> list[dict]:
    rows = []
    result_rows = result_row_locator(page)
    count = result_rows.count()
    debug_log(f"Extracting page {current_page}. Result rows found: {count}")
    # Use the same scoped locator so action IDs come from the active tab only
    action_link_ids = result_row_locator(page).evaluate_all(
        """rows => rows.map(row => {
            const lastCell = row.querySelector('td:last-child');
            const actionLink = lastCell ? lastCell.querySelector('a') : null;
            return actionLink ? actionLink.id : '';
        })"""
    )

    for index in range(count):
        if CONFIG.max_rows_per_page is not None and index >= CONFIG.max_rows_per_page:
            break

        row = result_rows.nth(index)
        title_link = result_title_link_locator(row)
        cells = row.locator("td")
        values = [clean(cells.nth(i).inner_text()) for i in range(cells.count())]

        if len(values) < 5:
            debug_log(f"Skipping row {index + 1} on page {current_page}: expected at least 5 cells, found {len(values)}")
            continue

        rows.append(
            {
                "page_number": current_page,
                "row_index": index,
                "title": values[0],
                "ptj": values[1],
                "publish_date": values[2],
                "closing_date": values[3],
                "action_text": values[4],
                "detail_control_id": title_link.get_attribute("id") or "",
                "action_control_id": action_link_ids[index] if index < len(action_link_ids) else "",
            }
        )

    return rows


def click_next_page(page: Page) -> bool:
    first_title = first_result_title(page)
    previous_page_number = current_results_page_number(page)
    target_page_number = previous_page_number + 1 if previous_page_number else 0
    paginator_state = visible_paginator_state(page)
    paginator_index = paginator_state.get("index", -1)
    panel_id = paginator_state.get("panelId", "")

    clicked = page.evaluate(
        """([panelId, paginatorIndex, targetPageNumber]) => {
            const scope = panelId ? document.getElementById(panelId) : document;
            if (!scope) return false;
            const paginators = Array.from(scope.querySelectorAll('.ui-paginator'));
            const paginator = paginatorIndex >= 0 ? paginators[paginatorIndex] : null;
            if (!paginator) return false;

            if (targetPageNumber) {
                const numberedPages = Array.from(paginator.querySelectorAll('.ui-paginator-page'));
                for (const pageLink of numberedPages) {
                    const style = window.getComputedStyle(pageLink);
                    if (style.display === 'none' || style.visibility === 'hidden') continue;
                    const text = (pageLink.textContent || '').trim();
                    if (parseInt(text, 10) === targetPageNumber) {
                        pageLink.click();
                        return true;
                    }
                }
            }

            const nextButtons = Array.from(paginator.querySelectorAll('.ui-paginator-next'));
            for (const button of nextButtons) {
                const style = window.getComputedStyle(button);
                if (style.display === 'none' || style.visibility === 'hidden') continue;
                const classes = button.className || '';
                if (classes.includes('ui-state-disabled')) return false;
                button.click();
                return true;
            }
            return false;
        }""",
        [panel_id, paginator_index, target_page_number],
    )

    if not clicked:
        return False

    wait_for_ajax(page, pause_ms=2500)

    try:
        page.wait_for_function(
            """([previousTitle, previousPageNumber, panelId, paginatorIndex]) => {
                const scope = panelId ? document.getElementById(panelId) : document;
                if (!scope) return false;
                const paginators = Array.from(scope.querySelectorAll('.ui-paginator'));
                const paginator = paginatorIndex >= 0 ? paginators[paginatorIndex] : null;
                if (!paginator) return false;
                const firstRow = scope.querySelector("tbody[id*='nbsearchresults_data'] > tr");
                const activePages = Array.from(paginator.querySelectorAll('.ui-paginator-page.ui-state-active'));
                let currentPageNumber = 0;
                for (const pageLink of activePages) {
                    const text = (pageLink.textContent || '').trim();
                    const style = window.getComputedStyle(pageLink);
                    if (style.display === 'none' || style.visibility === 'hidden') continue;
                    const parsed = parseInt(text, 10);
                    if (!Number.isNaN(parsed)) {
                        currentPageNumber = parsed;
                        break;
                    }
                }
                if (previousPageNumber && currentPageNumber && currentPageNumber !== previousPageNumber) return true;
                if (!firstRow) return false;
                const titleAnchor = firstRow.querySelector("a[id*='nbsearchresults']");
                const text = titleAnchor ? titleAnchor.textContent : firstRow.textContent;
                return (text || '').replace(/\\s+/g, ' ').trim() !== previousTitle;
            }""",
            arg=[first_title, previous_page_number, panel_id, paginator_index],
            timeout=CONFIG.timeout_ms,
        )
    except PlaywrightTimeoutError:
        debug_log("Paginator next click did not move to a new results page.")
        save_debug_artifacts(page, "pagination_next_failed")
        return False

    return True


def go_to_result_page(page: Page, target_page: int) -> None:
    if target_page <= 1:
        return

    for _ in range(1, target_page):
        if not click_next_page(page):
            break


def ensure_results_page(page: Page, target_page: int) -> None:
    if result_row_locator(page).count():
        current_page = current_results_page_number(page)
        if target_page <= 1 or current_page == target_page:
            return
        debug_log(f"Results are visible but paginator is on page {current_page}; rebuilding page {target_page}.")

    debug_log(f"Returning to search results for page {target_page}")
    page.goto(URL, wait_until="domcontentloaded", timeout=CONFIG.timeout_ms)
    wait_for_ajax(page, pause_ms=2500)
    switch_to_target_tab(page, CONFIG.target_tab)
    apply_kod_bidang_filter(page, CONFIG)
    apply_date_filter(page, CONFIG)
    click_search(page)
    go_to_result_page(page, target_page)


def safe_go_back_to_results(page: Page, target_page: int) -> None:
    if target_page > 1:
        debug_log(f"Rebuilding filtered results directly for page {target_page}")
        ensure_results_page(page, target_page)
        return

    if result_row_locator(page).count():
        current_page = current_results_page_number(page)
        if target_page <= 1 or current_page == target_page:
            return
        debug_log(f"Results returned on page {current_page} instead of page {target_page}; correcting.")

    try:
        back_button = page.get_by_text("Kembali", exact=False).first
        if back_button.count():
            debug_log(f"Returning to search results with the page back button for page {target_page}")
            back_button.click()
            wait_for_ajax(page, pause_ms=2000)
            if result_row_locator(page).count():
                current_page = current_results_page_number(page)
                if target_page <= 1 or current_page == target_page:
                    return
                debug_log(f"Back button returned results page {current_page}; need page {target_page}.")
    except Exception as exc:
        debug_log(f"Page back button did not return results directly: {exc}")

    try:
        debug_log("Falling back to browser history back")
        page.go_back(wait_until="domcontentloaded", timeout=CONFIG.timeout_ms)
        wait_for_ajax(page, pause_ms=2000)
        if result_row_locator(page).count():
            current_page = current_results_page_number(page)
            if target_page <= 1 or current_page == target_page:
                return
            debug_log(f"Browser history returned results page {current_page}; need page {target_page}.")
    except Exception as exc:
        debug_log(f"Browser history back did not return results directly: {exc}")

    ensure_results_page(page, target_page)


def open_notice_detail(page: Page, listing_row: dict) -> None:
    ensure_results_page(page, listing_row["page_number"])
    row = resolve_listing_row(page, listing_row)
    title_link = result_title_link_locator(row)
    debug_log(f"Opening notice detail: {listing_row['title']}")
    try:
        title_link.click(timeout=5000)
    except PlaywrightTimeoutError:
        debug_log("Title link was not considered visible by Playwright. Falling back to DOM click.")
        row.evaluate(
            """rowElement => {
                const anchor = rowElement.querySelector("a[id*='nbsearchresults']");
                if (!anchor) {
                    throw new Error("Notice detail link not found in resolved row");
                }
                anchor.click();
            }"""
        )
    wait_for_ajax(page, pause_ms=3000)

    try:
        page.wait_for_function(
            """() => {
                const bodyText = (document.body.innerText || '').toUpperCase();
                return bodyText.includes('KOD BIDANG') || bodyText.includes('SENARAI DOKUMEN');
            }""",
            timeout=CONFIG.timeout_ms,
        )
    except PlaywrightTimeoutError:
        debug_log("Detail page markers did not appear in time. Continuing with best-effort parsing.")
        save_debug_artifacts(page, "notice_detail_wait_failed")

    save_debug_artifacts(page, f"detail_page_{listing_row['page_number']}_{listing_row['row_index'] + 1}")


def extract_tables(page: Page) -> list[dict]:
    return page.locator("table").evaluate_all(
        """tables => tables.map((table, index) => ({
            index,
            headers: Array.from(table.querySelectorAll('th')).map(th =>
                (th.textContent || '').replace(/\\s+/g, ' ').trim()
            ).filter(Boolean),
            rows: Array.from(table.querySelectorAll('tbody tr')).map(tr =>
                Array.from(tr.querySelectorAll('td')).map(td =>
                    (td.textContent || '').replace(/\\s+/g, ' ').trim()
                )
            ).filter(row => row.some(Boolean))
        }))"""
    )


def find_table_by_headers(page: Page, required_headers: list[str]) -> Optional[dict]:
    required = {header.upper() for header in required_headers}
    for table in extract_tables(page):
        headers = {header.upper() for header in table["headers"]}
        if required.issubset(headers):
            return table
    return None


def parse_kod_bidang_detail(page: Page) -> dict:
    table = find_table_by_headers(page, ["Kod Bidang", "Bidang", "Sub Bidang", "Pecahan Sub Bidang"])
    if not table or not table["rows"]:
        return {
            "kod_bidang_code": "",
            "bidang": "",
            "sub_bidang": "",
            "pecahan_sub_bidang": "",
        }

    first_row = table["rows"][0]
    padded = first_row + ["", "", "", ""]
    return {
        "kod_bidang_code": padded[0],
        "bidang": padded[1],
        "sub_bidang": padded[2],
        "pecahan_sub_bidang": padded[3],
    }


def extract_tender_detail_fields(page: Page) -> dict:
    fields = {}

    try:
        detail_rows = page.locator("table[id$='publishTabQTDtls'] tr")
        row_count = detail_rows.count()
    except Exception:
        row_count = 0

    for index in range(row_count):
        row = detail_rows.nth(index)
        cells = row.locator("td")
        if cells.count() < 2:
            continue
        label = clean(cells.nth(0).inner_text())
        value = clean(cells.nth(1).inner_text())
        if label:
            fields[label] = value

    return fields


def extract_detail_value(page: Page, label: str) -> str:
    detail_fields = extract_tender_detail_fields(page)
    if label in detail_fields and clean(detail_fields[label]):
        return clean(detail_fields[label])

    lines = [clean(line) for line in page.locator("body").inner_text().splitlines() if clean(line)]
    normalized_label = clean(label).lower()

    for index, line in enumerate(lines):
        lower_line = line.lower()
        if normalized_label not in lower_line:
            continue

        if ":" in line:
            before, after = line.split(":", 1)
            if normalized_label in clean(before).lower():
                if clean(after):
                    return clean(after)
                if index + 1 < len(lines):
                    return clean(lines[index + 1])

        if index + 1 < len(lines):
            return clean(lines[index + 1])

    return ""


def is_valid_notice_number(value: str) -> bool:
    return bool(re.fullmatch(r"QT\d{6,}", clean(value), flags=re.IGNORECASE))


def is_invalid_supplier_text(value: str) -> bool:
    upper = clean(value).upper()
    if not upper:
        return True
    invalid_fragments = [
        "EPEROLEHAN, KEMENTERIAN KEWANGAN",
        "DISELENGGARA OLEH COMMERCE DOT COM",
        "TAJUK PEROLEHAN",
        "KRITERIA CARIAN",
    ]
    return any(fragment in upper for fragment in invalid_fragments)


def normalize_compare_text(value: str) -> str:
    return re.sub(r"[^A-Z0-9]+", " ", clean(value).upper()).strip()


def looks_like_valid_detail_page(page: Page, listing_row: dict) -> bool:
    detail_fields = extract_tender_detail_fields(page)
    try:
        body_text = clean(page.locator("body").inner_text())
    except Exception:
        return False

    title = clean(listing_row.get("title", ""))
    detail_title = clean(detail_fields.get("Tajuk Perolehan", ""))
    normalized_title = normalize_compare_text(title)
    normalized_detail_title = normalize_compare_text(detail_title)
    normalized_body = normalize_compare_text(body_text)

    if normalized_title and normalized_detail_title:
        if normalized_title not in normalized_detail_title and normalized_detail_title not in normalized_title:
            debug_log(
                f"Detail validation title mismatch. Listing title={title!r}, detail title={detail_title!r}"
            )
            return False
    elif normalized_title and normalized_title not in normalized_body:
        debug_log(f"Detail validation could not find listing title in page body for: {title}")
        return False

    notice_number = clean(detail_fields.get("No. Sebut Harga / Tender", "")) or extract_detail_value(page, "No. Sebut Harga / Tender")
    has_kod_bidang = bool(parse_kod_bidang_detail(page).get("kod_bidang_code", ""))
    has_specification_row = find_specification_row(page).count() > 0
    has_detail_title = bool(detail_title) or "TAJUK PEROLEHAN" in body_text.upper()

    if is_valid_notice_number(notice_number):
        return True

    if has_detail_title and (has_kod_bidang or has_specification_row or len(detail_fields) >= 4):
        debug_log(
            "Detail validation accepted page without QT notice number because other detail markers were present."
        )
        return True

    debug_log(
        "Detail validation failed: "
        f"notice_number={notice_number!r}, has_detail_title={has_detail_title}, "
        f"has_kod_bidang={has_kod_bidang}, has_specification_row={has_specification_row}, "
        f"detail_field_count={len(detail_fields)}"
    )
    return False


def find_specification_row(page: Page):
    return page.locator("tr").filter(has_text="Spesifikasi").first


def read_download_bytes(download: Download) -> bytes:
    download_path = download.path()
    if not download_path:
        return b""
    return Path(download_path).read_bytes()


def extract_pdf_text_from_bytes(pdf_bytes: bytes) -> str:
    if not pdf_bytes:
        return ""

    try:
        reader = PdfReader(io.BytesIO(pdf_bytes))
        pages = [clean(page.extract_text() or "") for page in reader.pages]
        return "\n".join(page for page in pages if page).strip()
    except Exception as exc:
        debug_log(f"Could not extract text from PDF bytes: {exc}")
        return ""


def extract_items_from_spec_text(specification_text: str) -> list[dict]:
    """Extract item rows from ePerolehan specification text.

    The site uses several PDF table layouts, so this parser tries multiple
    patterns and returns every distinct item it can find:

    1. Inline rows like "PORTABLE DENTAL UNIT unit Biasa Standard 10"
    2. Multiline rows where description comes before a separate UOM row
    3. Item-code blocks like "ZC" + "2514" + "unit 1 1"
    4. Service-style one-line rows with description + UOM + quantity
    """
    if not specification_text:
        return []

    uom_pattern = r"(unit|units|set|sets|lot|lots|kit|kits|pair|pairs|service|services|servis|buah|botol|pack|packs|box|boxes)"
    lines = [clean(line) for line in specification_text.replace("\r", "\n").splitlines() if clean(line)]
    items: list[dict] = []

    def add_item(item_desc: str, item_uom: str, item_quantity: str, item_no: Optional[str] = None) -> None:
        normalized_desc = clean(item_desc)
        normalized_uom = clean(item_uom)
        normalized_quantity = clean(item_quantity)
        normalized_item_no = clean(item_no or str(len(items) + 1))
        if not normalized_desc:
            return
        item_key = (
            normalized_item_no.casefold(),
            normalized_desc.casefold(),
            normalized_uom.casefold(),
            normalized_quantity.casefold(),
        )
        seen_keys = {
            (
                clean(existing.get("item_no", "")).casefold(),
                clean(existing.get("item_desc", "")).casefold(),
                clean(existing.get("item_uom", "")).casefold(),
                clean(existing.get("item_quantity", "")).casefold(),
            )
            for existing in items
        }
        if item_key in seen_keys:
            return
        items.append(
            {
                "item_no": normalized_item_no,
                "item_desc": normalized_desc,
                "item_uom": normalized_uom,
                "item_quantity": normalized_quantity,
            }
        )

    header_re = re.compile(r"Item\s+UOM\s+Jenis\s+Harga\s+Kuantiti\s+Spesifikasi", flags=re.IGNORECASE)
    inline_item_re = re.compile(
        rf"(?P<item>[A-Z][^.]*?)\s+(?P<uom>{uom_pattern})\s+Biasa\s+Standard\s+(?P<quantity>\d+(?:\.\d+)?)",
        flags=re.IGNORECASE,
    )
    segments = header_re.split(specification_text)
    for seg in segments[1:]:
        pos = 0
        seg_stripped = seg.lstrip()
        while pos < len(seg_stripped):
            match = inline_item_re.search(seg_stripped, pos)
            if not match:
                break
            item_text = clean(match.group("item"))
            if len(item_text) <= 120 and not re.search(r"\d{1,2}\.\d", item_text):
                item_text = re.sub(
                    r"^[A-Z]\.\s+APPLICATION FEATURES\s*[-–]\s*",
                    "",
                    item_text,
                    flags=re.IGNORECASE,
                ).strip()
                add_item(item_text, match.group("uom"), match.group("quantity"))
            pos = match.end()
            next_start = re.search(r"\d\s+(?=[A-Z])", seg_stripped[pos:])
            if next_start:
                pos += next_start.end()
            else:
                break

    pending_desc: list[str] = []
    started = False

    def is_header_line(text: str) -> bool:
        upper_text = text.upper()
        return (
            upper_text.startswith("PAGE ")
            or upper_text in {"ITEM", "SPESIFIKASI"}
            or upper_text.startswith("TAJUK TEMPLAT")
            or upper_text.startswith("JENIS ITEM")
            or upper_text.startswith("JENIS BARANGAN")
            or upper_text.startswith("UOM JENIS")
            or upper_text.startswith("UOM ")
        )

    def parse_item_row(text: str) -> Optional[tuple[str, str]]:
        match = re.match(
            rf"^(?P<uom>{uom_pattern})(?:\s+\S+){{1,5}}\s+(?P<quantity>\d+(?:\.\d+)?)$",
            text,
            flags=re.IGNORECASE,
        )
        if not match:
            return None
        return clean(match.group("uom")), clean(match.group("quantity"))

    def looks_like_item_desc_line(text: str) -> bool:
        if len(text) > 120:
            return False
        if re.match(r"^\d+\.", text):
            return False
        letters = [char for char in text if char.isalpha()]
        if not letters:
            return False
        uppercase_ratio = sum(1 for char in letters if char.isupper()) / len(letters)
        return uppercase_ratio >= 0.6

    index = 0
    while index < len(lines):
        line = lines[index]
        upper = line.upper()

        if is_header_line(line):
            if "ITEM" in upper or upper.startswith("UOM "):
                started = True
            index += 1
            continue

        if not started:
            index += 1
            continue

        row_data = parse_item_row(line)
        consumed_next = False
        if not row_data and index + 1 < len(lines):
            combined = clean(f"{line} {lines[index + 1]}")
            row_data = parse_item_row(combined)
            if row_data:
                consumed_next = True

        if row_data:
            item_uom, item_quantity = row_data
            item_desc = clean(" ".join(pending_desc))
            if item_desc:
                add_item(item_desc, item_uom, item_quantity)
            pending_desc = []
            index += 2 if consumed_next else 1
            continue

        if re.match(r"^\d+\.", line):
            index += 1
            continue

        if not is_header_line(line) and looks_like_item_desc_line(line):
            if not pending_desc or pending_desc[-1] != line:
                pending_desc.append(line)
            if len(pending_desc) > 8:
                pending_desc = pending_desc[-8:]

        index += 1

    code_block_starts: list[tuple[int, str]] = []
    for idx in range(len(lines) - 2):
        if re.fullmatch(r"[A-Z]{1,5}", lines[idx]) and re.fullmatch(r"\d{2,6}", lines[idx + 1]):
            if re.match(rf"^{uom_pattern}\b", lines[idx + 2], flags=re.IGNORECASE):
                code_block_starts.append((idx, clean(f"{lines[idx]} {lines[idx + 1]}")))

    for block_index, (start_idx, item_code) in enumerate(code_block_starts):
        end_idx = code_block_starts[block_index + 1][0] if block_index + 1 < len(code_block_starts) else len(lines)
        block = lines[start_idx:end_idx]
        item_uom = ""
        item_quantity = ""
        for block_line in block[:6]:
            block_match = re.match(
                rf"^(?P<uom>{uom_pattern})(?:\s+\S+){{0,5}}\s+(?P<quantity>\d+(?:\.\d+)?)$",
                block_line,
                flags=re.IGNORECASE,
            )
            if block_match:
                item_uom = clean(block_match.group("uom"))
                item_quantity = clean(block_match.group("quantity"))
                break
        add_item(item_code, item_uom, item_quantity, item_no=item_code)

    inline_service_re = re.compile(
        rf"^(?P<item>[A-Z0-9 /()&.,'-]{{4,160}}?)\s+(?P<uom>{uom_pattern})\s+(?:Biasa|Standard|\d+)(?:\s+\S+){{0,3}}\s+(?P<quantity>\d+(?:\.\d+)?)$",
        flags=re.IGNORECASE,
    )
    for line in lines:
        service_match = inline_service_re.match(line)
        if not service_match:
            continue
        item_desc = clean(service_match.group("item"))
        if re.match(r"^\d+\.", item_desc):
            continue
        add_item(item_desc, service_match.group("uom"), service_match.group("quantity"))

    if items:
        return items

    fallback_desc = ""
    for line in lines:
        upper = line.upper()
        if upper.startswith("PAGE ") or upper in {"SPESIFIKASI", "ITEM", "UOM", "KUANTITI"}:
            continue
        if upper.startswith("TAJUK TEMPLAT") or upper.startswith("JENIS ITEM") or upper.startswith("JENIS BARANGAN"):
            continue
        if re.match(r"^\d+\.", line):
            break
        if re.fullmatch(r"[A-Z0-9 /()&.,'-]+", line):
            fallback_desc = clean(f"{fallback_desc} {line}")
            if len(fallback_desc) > 120:
                break
    return [
        {
            "item_no": "1",
            "item_desc": fallback_desc,
            "item_uom": "",
            "item_quantity": "",
        }
    ]


def download_specification_pdf(page: Page) -> tuple[str, bytes, str]:
    row = find_specification_row(page)
    if row.count() == 0:
        return "", b"", ""

    document_name = clean(row.locator("td").first.inner_text())
    buttons = row.locator("a, button, span.ui-button-icon-left, span.ui-button-icon-only")
    button_count = buttons.count()
    if button_count == 0:
        return document_name, b"", ""

    debug_log(f"Downloading specification PDF for: {document_name}")
    try:
        download_timeout_ms = min(CONFIG.timeout_ms, 15000)
        with page.expect_download(timeout=download_timeout_ms) as download_info:
            buttons.first.click(force=True)
        download = download_info.value
        pdf_bytes = read_download_bytes(download)
        return document_name, pdf_bytes, extract_pdf_text_from_bytes(pdf_bytes)
    except PlaywrightTimeoutError:
        debug_log(
            f"Specification PDF download did not start within {download_timeout_ms}ms for: {document_name}. "
            "Continuing without PDF."
        )
        save_debug_artifacts(page, "specification_download_timeout")
        return document_name, b"", ""
    except Exception as exc:
        debug_log(f"Specification PDF download failed: {exc}")
        save_debug_artifacts(page, "specification_download_failed")
        return document_name, b"", ""


def extract_supplier_page_fields(page: Page) -> dict:
    company_pattern = re.compile(
        r"\b[A-Z][A-Z0-9&.,'()/ -]{2,}(?:SDN\.?\s*BHD\.?|BHD\.?|ENTERPRISE|TRADING|RESOURCES|SERVICES|HOLDINGS|INDUSTRIES|SUPPLIES)\b",
        flags=re.IGNORECASE,
    )

    def looks_like_supplier_name(text: str) -> bool:
        upper = clean(text).upper()
        if not upper:
            return False
        if any(token in upper for token in ["KEMENTERIAN", "JABATAN", "PEJABAT", "HOSPITAL", "KLINIK", "KESIHATAN", "PTJ"]):
            return False
        return bool(company_pattern.search(upper))

    for table in extract_tables(page):
        headers = [header.upper() for header in table["headers"]]
        rows = table["rows"]
        if not rows:
            continue
        if not any("PEMBEKAL" in header for header in headers):
            continue

        supplier_index = next((i for i, h in enumerate(headers) if "PEMBEKAL" in h), None)
        item_index = next((i for i, h in enumerate(headers) if h == "ITEM"), None)
        amount_index = next((i for i, h in enumerate(headers) if "DENGAN LOA" in h), None)
        if supplier_index is None:
            continue

        for row in rows:
            supplier_name = clean(row[supplier_index]) if supplier_index < len(row) else ""
            item_desc = clean(row[item_index]) if item_index is not None and item_index < len(row) else ""
            amount = clean(row[amount_index]) if amount_index is not None and amount_index < len(row) else ""
            if looks_like_supplier_name(supplier_name) and len(supplier_name) <= 120 and not is_invalid_supplier_text(supplier_name):
                return {
                    "supplier_name": supplier_name,
                    "item_desc": item_desc,
                    "amount": amount,
                }

    for table in extract_tables(page):
        for row in table["rows"]:
            for cell in row:
                cell_text = clean(cell)
                if looks_like_supplier_name(cell_text) and not is_invalid_supplier_text(cell_text):
                    return {
                        "supplier_name": cell_text,
                        "item_desc": "",
                        "amount": "",
                    }

    body_text = clean(page.locator("body").inner_text())
    supplier_match = company_pattern.search(body_text)
    return {
        "supplier_name": clean(supplier_match.group(0)) if supplier_match and not is_invalid_supplier_text(supplier_match.group(0)) else "",
        "item_desc": "",
        "amount": "",
    }


def extract_notice_id(notice_number: str, listing_row: dict) -> str:
    return clean(notice_number) or listing_row["detail_control_id"]


def extract_notice_url(page: Page, notice_number: str) -> str:
    form = page.locator("form[action*='nbQuotationNTendenDetail.xhtml']").first
    if form.count():
        action = form.get_attribute("action") or ""
        if action:
            base_url = urljoin(URL, action)
            if notice_number:
                return f"{base_url}#notice_id={quote(notice_number)}"
            return base_url

    if "nbQuotationNTendenDetail.xhtml" in page.url:
        if notice_number:
            return f"{page.url}#notice_id={quote(notice_number)}"
        return page.url

    if notice_number:
        return f"{URL}#notice_id={quote(notice_number)}"
    return page.url


def build_query_text(config: SearchConfig) -> str:
    return (
        f"tab={config.target_tab};"
        f"date_from={config.date_from};"
        f"date_to={config.date_to};"
        f"bidang={config.bidang_code or ''};"
        f"sub_bidang={config.sub_bidang_code or ''};"
        f"pecahan_sub_bidang={config.pecahan_sub_bidang_code or ''}"
    )


def build_description(specification_text: str, title: str) -> str:
    return ""


def build_dedup_key(record: dict) -> str:
    key = "|".join(
        [
            record["source"],
            record["notice_id"],
            record["title"],
            record["publication_date"],
            record["closing_date"],
            record.get("item_no", ""),
            record.get("item_desc", ""),
        ]
    )
    return hashlib.sha256(key.encode("utf-8")).hexdigest()


def build_listing_fallback_records(listing_row: dict, supplier_fields: dict) -> list[dict]:
    supplier_name = supplier_fields.get("supplier_name", "")
    if is_invalid_supplier_text(supplier_name):
        supplier_name = ""

    supplier_item_desc = clean(supplier_fields.get("item_desc", ""))
    amount = clean(supplier_fields.get("amount", ""))

    record = {
        "source": CONFIG.source,
        "country": CONFIG.country,
        "country_cc": CONFIG.country_cc,
        "publication_date": listing_row["publish_date"],
        "closing_date": listing_row["closing_date"],
        "title": listing_row["title"],
        "title_en": translate_text_basic(listing_row["title"]),
        "buyer": listing_row["ptj"],
        "buyer_en": translate_text_basic(listing_row["ptj"]),
        "status": CONFIG.status,
        "currency": CONFIG.currency,
        "amount": amount,
        "supplier_name": supplier_name,
        "contract_period": "",
        "item_no": "1",
        "item_desc": supplier_item_desc,
        "item_desc_en": translate_text_basic(supplier_item_desc),
        "item_uom": "",
        "item_quantity": "",
        "kod_bidang_code": "",
        "main_category": "",
        "main_category_en": "",
        "sub_category": "",
        "sub_category_en": "",
        "category_detail": "",
        "category_detail_en": "",
        "specification_document_name": "",
        "specification_text": "",
        "notice_id": "",
        "notice_url": "",
        "query_text": build_query_text(CONFIG),
        "scraped_at": datetime.now(timezone.utc).isoformat(),
        "detail_control_id": listing_row["detail_control_id"],
    }
    record["dedup_key"] = build_dedup_key(record)
    return [record]


def deduplicate_notice_items(items: list[dict], supplier_item_desc: str = "") -> list[dict]:
    unique_items = []
    seen_keys = set()

    for item in items:
        item_no = clean(item.get("item_no", ""))
        item_desc = clean(item.get("item_desc", "") or supplier_item_desc)
        item_uom = clean(item.get("item_uom", ""))
        item_quantity = clean(item.get("item_quantity", ""))

        if not item_desc and not item_uom and not item_quantity:
            continue

        item_key = (
            item_desc.casefold(),
            item_uom.casefold(),
            item_quantity.casefold(),
        )
        if item_key in seen_keys:
            continue

        seen_keys.add(item_key)
        unique_items.append(
            {
                "item_no": item_no or str(len(unique_items) + 1),
                "item_desc": item_desc,
                "item_uom": item_uom,
                "item_quantity": item_quantity,
            }
        )

    if unique_items:
        for index, item in enumerate(unique_items, start=1):
            if not clean(item.get("item_no", "")):
                item["item_no"] = str(index)
        return unique_items

    fallback_desc = clean(supplier_item_desc)
    if fallback_desc:
        return [
            {
                "item_no": "1",
                "item_desc": fallback_desc,
                "item_uom": "",
                "item_quantity": "",
            }
        ]

    return []


def record_completeness_score(record: dict) -> int:
    fields = [
        "supplier_name",
        "amount",
        "contract_period",
        "item_desc",
        "item_uom",
        "item_quantity",
        "notice_url",
        "specification_document_name",
    ]
    score = sum(1 for field in fields if clean(record.get(field, "")))
    score += len(clean(record.get("specification_text", ""))) // 500
    return score


def deduplicate_output_rows(rows: list[dict]) -> list[dict]:
    best_by_key: dict[str, dict] = {}

    for row in rows:
        dedup_key = clean(row.get("dedup_key", ""))
        if not dedup_key:
            dedup_key = build_dedup_key(row)
            row["dedup_key"] = dedup_key

        existing = best_by_key.get(dedup_key)
        if existing is None or record_completeness_score(row) > record_completeness_score(existing):
            best_by_key[dedup_key] = row

    return list(best_by_key.values())


def load_existing_rows_from_excel(output_xlsx: str) -> list[dict]:
    output_path = Path(output_xlsx).resolve()
    if not output_path.exists():
        return []

    workbook = load_workbook(output_path)
    worksheet = workbook.active
    values = list(worksheet.values)
    if not values:
        return []

    headers = [clean(str(value or "")) for value in values[0]]
    rows = []
    for row_values in values[1:]:
        row = {}
        for index, header in enumerate(headers):
            if not header:
                continue
            cell_value = row_values[index] if index < len(row_values) else ""
            row[header] = "" if cell_value is None else str(cell_value)
        if any(clean(str(value)) for value in row.values()):
            rows.append(row)
    return rows


def load_progress(progress_json: str) -> dict:
    progress_path = Path(progress_json).resolve()
    if not progress_path.exists():
        return {"resume_page": 1, "completed_titles": {}}

    try:
        return json.loads(progress_path.read_text(encoding="utf-8"))
    except Exception as exc:
        debug_log(f"Could not read progress file {progress_path}: {exc}")
        return {"resume_page": 1, "completed_titles": {}}


def save_progress(progress_json: str, progress: dict) -> None:
    progress_path = Path(progress_json).resolve()
    progress_path.write_text(json.dumps(progress, ensure_ascii=True, indent=2), encoding="utf-8")


def completed_titles_for_page(progress: dict, page_number: int) -> set[str]:
    completed = progress.get("completed_titles", {})
    page_titles = completed.get(str(page_number), [])
    return {clean(title) for title in page_titles if clean(title)}


def mark_listing_row_completed(progress: dict, listing_row: dict) -> None:
    page_key = str(listing_row["page_number"])
    completed = progress.setdefault("completed_titles", {})
    titles = completed.setdefault(page_key, [])
    title = clean(listing_row.get("title", ""))
    if title and title not in titles:
        titles.append(title)
    progress["resume_page"] = listing_row["page_number"]


def mark_page_completed(progress: dict, page_number: int) -> None:
    completed = progress.setdefault("completed_titles", {})
    completed.pop(str(page_number), None)
    progress["resume_page"] = page_number + 1


def reset_progress(progress_json: str) -> None:
    progress_path = Path(progress_json).resolve()
    if progress_path.exists():
        progress_path.unlink()


def open_supplier_detail(page: Page, listing_row: dict) -> dict:
    ensure_results_page(page, listing_row["page_number"])
    debug_log(f"Opening supplier detail for: {listing_row['title']}")
    row = resolve_listing_row(page, listing_row)
    action_control_id = row.evaluate(
        """row => {
            const lastCell = row.querySelector('td:last-child');
            const actionLink = lastCell ? lastCell.querySelector('a') : null;
            return actionLink ? actionLink.id : '';
        }"""
    )
    if not action_control_id:
        return {"supplier_name": "", "item_desc": "", "amount": ""}

    page.evaluate(
        """elementId => {
            const element = document.getElementById(elementId);
            if (!element) {
                throw new Error(`Supplier action element not found: ${elementId}`);
            }
            element.click();
        }""",
        action_control_id,
    )
    wait_for_ajax(page, pause_ms=2500)
    try:
        page.wait_for_function(
            """() => {
                const bodyText = (document.body.innerText || '').toUpperCase();
                if (bodyText.includes('PEMBEKAL BERJAYA')) return true;
                return Array.from(document.querySelectorAll('table')).some(table => {
                    const text = (table.innerText || '').toUpperCase();
                    return text.includes('PEMBEKAL');
                });
            }""",
            timeout=CONFIG.timeout_ms,
        )
    except PlaywrightTimeoutError:
        debug_log("Supplier page markers did not appear in time. Continuing with best-effort parsing.")
        save_debug_artifacts(page, f"supplier_page_wait_failed_{listing_row['page_number']}_{listing_row['row_index'] + 1}")
    supplier_fields = extract_supplier_page_fields(page)
    if not supplier_fields.get("supplier_name"):
        save_debug_artifacts(page, f"supplier_page_missing_{listing_row['page_number']}_{listing_row['row_index'] + 1}")
    safe_go_back_to_results(page, listing_row["page_number"])
    return supplier_fields


def scrape_notice_detail(page: Page, listing_row: dict, supplier_fields: dict) -> list[dict]:
    open_notice_detail(page, listing_row)

    if not looks_like_valid_detail_page(page, listing_row):
        debug_log(f"Invalid detail page on first attempt for: {listing_row['title']}. Retrying from rebuilt results page.")
        save_debug_artifacts(page, f"invalid_detail_page_first_attempt_{listing_row['page_number']}_{listing_row['row_index'] + 1}")
        ensure_results_page(page, listing_row["page_number"])
        open_notice_detail(page, listing_row)

    if not looks_like_valid_detail_page(page, listing_row):
        debug_log(f"Saving listing-level fallback row for: {listing_row['title']}")
        save_debug_artifacts(page, f"invalid_detail_page_{listing_row['page_number']}_{listing_row['row_index'] + 1}")
        return build_listing_fallback_records(listing_row, supplier_fields)

    detail_fields = extract_tender_detail_fields(page)
    notice_number = clean(detail_fields.get("No. Sebut Harga / Tender", "")) or extract_detail_value(page, "No. Sebut Harga / Tender")
    notice_url = extract_notice_url(page, notice_number)
    kod_bidang = parse_kod_bidang_detail(page)
    specification_document_name, specification_bytes, specification_text = download_specification_pdf(page)
    items = extract_items_from_spec_text(specification_text)
    amount = clean(detail_fields.get("Jumlah Harga Indikatif (RM)", "")) or extract_detail_value(page, "Jumlah Harga Indikatif (RM)")
    contract_period = clean(detail_fields.get("Tempoh Sah Laku Tawaran (Hari)", "")) or extract_detail_value(page, "Tempoh Sah Laku Tawaran (Hari)")
    supplier_name = supplier_fields.get("supplier_name", "")
    if is_invalid_supplier_text(supplier_name):
        supplier_name = ""
    supplier_item_desc = supplier_fields.get("item_desc", "")
    items = deduplicate_notice_items(items, supplier_item_desc)
    if not amount:
        amount = supplier_fields.get("amount", "")

    base_record = {
        "source": CONFIG.source,
        "country": CONFIG.country,
        "country_cc": CONFIG.country_cc,
        "publication_date": listing_row["publish_date"],
        "closing_date": listing_row["closing_date"],
        "title": listing_row["title"],
        "title_en": translate_text_basic(listing_row["title"]),
        "buyer": listing_row["ptj"],
        "buyer_en": translate_text_basic(listing_row["ptj"]),
        "status": CONFIG.status,
        "currency": CONFIG.currency,
        "amount": amount,
        "supplier_name": supplier_name,
        "contract_period": contract_period,
        "kod_bidang_code": kod_bidang["kod_bidang_code"],
        "main_category": kod_bidang["bidang"],
        "main_category_en": translate_text_basic(kod_bidang["bidang"]),
        "sub_category": kod_bidang["sub_bidang"],
        "sub_category_en": translate_text_basic(kod_bidang["sub_bidang"]),
        "category_detail": kod_bidang["pecahan_sub_bidang"],
        "category_detail_en": translate_text_basic(kod_bidang["pecahan_sub_bidang"]),
        "specification_document_name": specification_document_name,
        "specification_text": specification_text,
        "notice_id": extract_notice_id(notice_number, listing_row),
        "notice_url": notice_url,
        "query_text": build_query_text(CONFIG),
        "scraped_at": datetime.now(timezone.utc).isoformat(),
        "detail_control_id": listing_row["detail_control_id"],
    }

    if specification_bytes and not specification_text:
        pass

    if not items:
        items = [{"item_no": "1", "item_desc": supplier_item_desc, "item_uom": "", "item_quantity": ""}]

    records = []
    for item in items:
        item_desc = item.get("item_desc", "") or supplier_item_desc
        record = {
            **base_record,
            "item_no": item.get("item_no", ""),
            "item_desc": item_desc,
            "item_desc_en": translate_text_basic(item_desc),
            "item_uom": item.get("item_uom", ""),
            "item_quantity": item.get("item_quantity", ""),
        }
        record["dedup_key"] = build_dedup_key(record)
        records.append(record)

    return records


def scrape_results(page: Page, existing_rows: list[dict], progress: dict) -> list[dict]:
    all_rows = list(existing_rows)
    page_number = max(1, int(progress.get("resume_page", 1) or 1))

    while True:
        max_pages_label = CONFIG.max_pages if CONFIG.max_pages is not None else "all"
        debug_log(f"Scraping results page {page_number} of max {max_pages_label}")
        listing_rows = extract_listing_rows_from_current_page(page, page_number)
        completed_titles = completed_titles_for_page(progress, page_number)

        for listing_row in listing_rows:
            if clean(listing_row.get("title", "")) in completed_titles:
                debug_log(f"Skipping already completed row on page {page_number}: {listing_row['title']}")
                continue

            supplier_fields = open_supplier_detail(page, listing_row)
            detail_records = scrape_notice_detail(page, listing_row, supplier_fields)
            all_rows.extend(detail_records)
            mark_listing_row_completed(progress, listing_row)
            all_rows = deduplicate_output_rows(all_rows)
            save_to_excel(all_rows, CONFIG.output_xlsx)
            save_progress(CONFIG.progress_json, progress)
            safe_go_back_to_results(page, page_number)

        mark_page_completed(progress, page_number)
        save_progress(CONFIG.progress_json, progress)

        if CONFIG.max_pages is not None and page_number >= CONFIG.max_pages:
            debug_log("Reached max_pages limit. Not clicking next page.")
            break

        if not click_next_page(page):
            debug_log("Next page is disabled. Stopping pagination.")
            break

        page_number += 1

    return all_rows


def save_to_excel(rows: list[dict], output_xlsx: str) -> Path:
    rows = deduplicate_output_rows(rows)
    output_path = Path(output_xlsx).resolve()

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Tenders"
    worksheet.append(OUTPUT_FIELDS)

    for row in rows:
        worksheet.append([row.get(field, "") for field in OUTPUT_FIELDS])

    header_font = Font(bold=True)
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    body_alignment = Alignment(vertical="top", wrap_text=True)

    for cell in worksheet[1]:
        cell.font = header_font
        cell.alignment = header_alignment

    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = body_alignment

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    for column_index, field_name in enumerate(OUTPUT_FIELDS, start=1):
        max_length = len(field_name)
        for row in rows:
            cell_value = clean(str(row.get(field_name, "")))
            if len(cell_value) > max_length:
                max_length = min(len(cell_value), 120)
        adjusted_width = min(max(max_length + 2, 12), 60)
        worksheet.column_dimensions[get_column_letter(column_index)].width = adjusted_width

    workbook.save(output_path)
    return output_path


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--date-from", default=CONFIG.date_from, help="Publication start date (DD/MM/YYYY)")
    parser.add_argument("--date-to", default=CONFIG.date_to, help="Publication end date (DD/MM/YYYY); defaults to today")
    parser.add_argument("--max-pages", type=int, default=None, help="Limit the scrape to the first N result pages")
    parser.add_argument("--headless", action="store_true", help="Run without showing the browser window")
    parser.add_argument("--no-debug", action="store_true", help="Hide debug log lines")
    parser.add_argument("--bidang-code", type=str, default=None, help="Kod Bidang main code, e.g. 05")
    parser.add_argument("--sub-bidang-code", type=str, default=None, help="Kod Bidang sub code, e.g. 01")
    parser.add_argument("--pecahan-sub-bidang-code", type=str, default=None, help="Kod Bidang detail code, e.g. 02")
    parser.add_argument("--output-xlsx", type=str, default=None, help="Output Excel filename")
    parser.add_argument("--progress-json", type=str, default=None, help="Checkpoint progress filename")
    parser.add_argument("--restart", action="store_true", help="Start fresh from page 1 and overwrite previous output/progress")
    parser.add_argument("--resume", action="store_true", help="Resume from existing progress/output if present")
    args = parser.parse_args()

    try:
        start_date = datetime.strptime(args.date_from, "%d/%m/%Y")
        end_date = datetime.strptime(args.date_to, "%d/%m/%Y")
    except ValueError:
        parser.error("Dates must use DD/MM/YYYY")
    if start_date > end_date:
        parser.error("--date-from must not be after --date-to")

    runtime_config = SearchConfig(
        target_tab=CONFIG.target_tab,
        date_from=args.date_from,
        date_to=args.date_to,
        bidang_code=args.bidang_code if args.bidang_code is not None else CONFIG.bidang_code,
        sub_bidang_code=args.sub_bidang_code if args.sub_bidang_code is not None else CONFIG.sub_bidang_code,
        pecahan_sub_bidang_code=args.pecahan_sub_bidang_code if args.pecahan_sub_bidang_code is not None else CONFIG.pecahan_sub_bidang_code,
        headless=args.headless if args.headless else CONFIG.headless,
        slow_mo_ms=CONFIG.slow_mo_ms,
        timeout_ms=CONFIG.timeout_ms,
        output_xlsx=args.output_xlsx if args.output_xlsx is not None else CONFIG.output_xlsx,
        progress_json=args.progress_json if args.progress_json is not None else CONFIG.progress_json,
        restart_from_first_page=args.restart if args.restart else (False if args.resume else CONFIG.restart_from_first_page),
        debug=False if args.no_debug else CONFIG.debug,
        debug_dir=CONFIG.debug_dir,
        max_pages=args.max_pages if args.max_pages is not None else CONFIG.max_pages,
        max_rows_per_page=CONFIG.max_rows_per_page,
        country=CONFIG.country,
        country_cc=CONFIG.country_cc,
        currency=CONFIG.currency,
        source=CONFIG.source,
        status=CONFIG.status,
    )

    globals()["CONFIG"] = runtime_config

    if CONFIG.restart_from_first_page:
        existing_rows = []
        progress = {"resume_page": 1, "completed_titles": {}}
        reset_progress(CONFIG.progress_json)
        output_path = Path(CONFIG.output_xlsx).resolve()
        if output_path.exists():
            output_path.unlink()
    else:
        existing_rows = load_existing_rows_from_excel(CONFIG.output_xlsx)
        progress = load_progress(CONFIG.progress_json)

    rows = existing_rows
    startup_attempts = 2
    with sync_playwright() as playwright:
        for attempt in range(1, startup_attempts + 1):
            browser = None
            context = None
            try:
                browser = playwright.chromium.launch(headless=CONFIG.headless, slow_mo=CONFIG.slow_mo_ms)
                context = browser.new_context(viewport={"width": 1440, "height": 1200}, accept_downloads=True)
                page = context.new_page()
                page.set_default_timeout(CONFIG.timeout_ms)

                debug_log(f"Opening URL: {URL}")
                page.goto(URL, wait_until="domcontentloaded", timeout=CONFIG.timeout_ms)
                wait_for_ajax(page, pause_ms=2500)
                save_debug_artifacts(page, "page_loaded")

                switch_to_target_tab(page, CONFIG.target_tab)
                apply_kod_bidang_filter(page, CONFIG)
                apply_date_filter(page, CONFIG)
                click_search(page)

                resume_page = max(1, int(progress.get("resume_page", 1) or 1))
                if resume_page > 1:
                    debug_log(f"Resuming from page {resume_page}")
                    go_to_result_page(page, resume_page)

                rows = scrape_results(page, existing_rows, progress)
                save_debug_artifacts(page, "before_browser_close")
                break
            except PlaywrightError as exc:
                if attempt < startup_attempts and is_target_closed_error(exc):
                    debug_log(f"Browser target closed during startup/scrape (attempt {attempt}/{startup_attempts}). Retrying session.")
                    continue
                raise
            finally:
                if context is not None:
                    try:
                        context.close()
                    except Exception:
                        pass
                if browser is not None:
                    try:
                        browser.close()
                    except Exception:
                        pass

    output_path = save_to_excel(rows, CONFIG.output_xlsx)
    reset_progress(CONFIG.progress_json)
    print(f"Saved {len(rows)} rows to {output_path}")
    print(f"Tab used: {CONFIG.target_tab}")
    print(f"Tarikh Iklan: {CONFIG.date_from} -> {CONFIG.date_to}")
    print(f"Kategori Kod Bidang used: {(CONFIG.bidang_code or '') + (CONFIG.sub_bidang_code or '') + (CONFIG.pecahan_sub_bidang_code or '')}")
    print(f"Max pages scraped: {CONFIG.max_pages if CONFIG.max_pages is not None else 'all available pages'}")
    if CONFIG.debug:
        print(f"Debug files saved in: {Path(CONFIG.debug_dir).resolve()}")


if __name__ == "__main__":
    main()
