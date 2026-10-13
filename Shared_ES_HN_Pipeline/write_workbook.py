"""Write the ES/HN tender workbook from eshn_rows.pkl."""
import os
import sys
from datetime import date

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from build_rows import EXTRA, TEMPLATE  # noqa: E402
from tagger import KMAP, SEG  # noqa: E402

OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "ES_HN_IVD_Diabetes_Tenders_2019-2026.xlsx")
NOTES = os.path.join(HERE, "notes.json")
TRANSLATE = ["Tender title / description", "Product / test", "Procedure type", "Status", "Unit",
             "Contract type (comodato, por prueba, según demanda…)", "Contract duration & extensions"]
EN_SUFFIX = " (English)"
EN_FILL = PatternFill("solid", fgColor="375623")


def clean(s):
    import re
    return re.sub(r"\s+", " ", str(s or "")).strip()


def load_translations():
    """es -> en from the translation agents' output plus copy-through strings."""
    import glob
    import json
    tdir = os.path.join(HERE, "tr")
    m = json.load(open(os.path.join(tdir, "copy_through.json"), encoding="utf-8")) if os.path.exists(os.path.join(tdir, "copy_through.json")) else {}
    for f in sorted(glob.glob(os.path.join(tdir, "chunk_*.jsonl"))):
        out = f.replace("chunk_", "out_")
        if not os.path.exists(out):
            continue
        src = {json.loads(l)["id"]: json.loads(l)["es"] for l in open(f, encoding="utf-8") if l.strip()}
        for l in open(out, encoding="utf-8"):
            if l.strip():
                o = json.loads(l)
                if o.get("id") in src and str(o.get("en", "")).strip():
                    m[src[o["id"]]] = normalise_en(o["en"].strip())
    return m


def normalise_en(s):
    """Make leftover Spanish labels consistent across the translators' outputs."""
    import re
    for pat, rep in [(r"C[OÓ]DIGO\s+(DE\s+)?MERCANC[IÍ]A\s+ONU", "UN commodity code"), (r"\bC[OÓ]DIGO\s*:", "Code:"),
                     (r"\bC[ÓO]DIGO\b", "Code"), (r"\bREGISTRO\s+DNM\b", "DNM registration no."), (r"\bVENCIMIENTO\b", "Expiry"),
                     (r"\bMARCA\s*:", "Brand:"), (r"\bORIGEN\s*:", "Origin:"), (r"\bMODELO\s*:", "Model:")]:
        s = re.sub(pat, rep, s, flags=re.I)
    return s


def with_english(cols):
    out = []
    for c in cols:
        out.append(c)
        if c in TRANSLATE:
            out.append(c + EN_SUFFIX)
    return out


def add_english(rows, tmap, missing):
    for r in rows:
        for c in TRANSLATE:
            v = clean(r.get(c, ""))
            if v.lower() in ("nan", "none"):
                r[c], v = "", ""
            if v:
                en = tmap.get(v)
                if en is None:
                    missing.add(v)
                r[c + EN_SUFFIX] = en or ""
            else:
                r[c + EN_SUFFIX] = ""

HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(bold=True, color="FFFFFF", name="Arial", size=9)
BODY = Font(name="Arial", size=9)
EXTRA_FILL = PatternFill("solid", fgColor="D9E1F2")
EXTRA_FONT = Font(bold=True, color="1F3864", name="Arial", size=9)
TITLE = Font(bold=True, name="Arial", size=12)
BOLD = Font(bold=True, name="Arial", size=9)
THIN = Border(bottom=Side(style="thin", color="BFBFBF"))
WIDTH = {"Country": 11, "Source portal": 22, "Buying institution / hospital": 34, "Tender / procedure no.": 24, "Procedure type": 18,
         "Tender title / description": 50, "Indicator code": 9, "Segment name": 22, "Product / test": 55, "Brand / analyser": 16,
         "Awarded supplier (distributor)": 30, "Quantity": 10, "Unit": 10, "Unit price": 11, "Total awarded": 14, "Currency": 8,
         "Contract type (comodato, por prueba, según demanda…)": 18, "Contract duration & extensions": 16, "Publication date": 11,
         "Award date": 11, "Status": 14, "Funding (national / World Bank / IDB / Global Fund / PAHO)": 14, "Source link": 30,
         "Comments": 45, "Tag confidence": 9, "All matched codes": 16, "Tagged from": 8, "Buyer segment": 24, "Award year": 7,
         "Multi-product line": 8, "UNSPSC": 26, "Closing date": 11}


def data_sheet(wb, name, rows, cols):
    ws = wb.create_sheet(name)
    for j, c in enumerate(cols, 1):
        cell = ws.cell(row=1, column=j, value=c)
        extra = c in EXTRA or c == "Closing date"
        if c.endswith(EN_SUFFIX):
            cell.fill, cell.font = EN_FILL, HDR_FONT
        else:
            cell.fill, cell.font = (EXTRA_FILL, EXTRA_FONT) if extra else (HDR_FILL, HDR_FONT)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[get_column_letter(j)].width = WIDTH.get(c[:-len(EN_SUFFIX)] if c.endswith(EN_SUFFIX) else c, 14)
    for i, r in enumerate(rows, 2):
        for j, c in enumerate(cols, 1):
            v = r.get(c, "")
            if isinstance(v, float) and pd.isna(v):
                v = None
            if c == "Award year":
                v = int(v) if str(v or "").isdigit() else None
            cell = ws.cell(row=i, column=j, value=v)
            cell.font = BODY
            if c in ("Quantity",):
                cell.number_format = "#,##0.##"
            elif c in ("Unit price", "Total awarded"):
                cell.number_format = "#,##0.00"
            elif c == "Source link" and isinstance(v, str) and v.startswith("http"):
                cell.hyperlink = v
                cell.font = Font(name="Arial", size=9, color="0563C1", underline="single")
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{max(1, len(rows) + 1)}"
    ws.row_dimensions[1].height = 42
    return ws


def summary_sheet(wb, aw_rows, n_aw):
    ws = wb.create_sheet("Summary", 1)
    ws["A1"] = "Summary by indicator code (formulas over the Awarded tab – they update if you re-code rows)"
    ws["A1"].font = TITLE
    ws["A2"] = ("National columns count rows whose Funding = 'National', in each country's published currency (Honduras HNL, El Salvador USD). "
                "Donor-financed purchases (Global Fund, World Bank, IDB, UN) are counted only in the last column. Multi-product lot lines are included at lot value.")
    ws["A2"].font = Font(name="Arial", size=9, italic=True)
    hdr = ["Indicator code", "Segment code", "Segment name", "HN lines (national funding)", "HN total (HNL)", "ES lines (national funding)",
           "ES total (USD)", "Donor / other-funded lines (any source; may repeat – see Comments)"]
    for j, h in enumerate(hdr, 1):
        c = ws.cell(row=4, column=j, value=h)
        c.fill, c.font = HDR_FILL, HDR_FONT
        c.alignment = Alignment(wrap_text=True)
    used = {r["Indicator code"] for r in aw_rows}
    codes = [k["Indicator code"] for k in KMAP if k["Indicator code"] in used]
    codes.sort(key=lambda c: [int(x) for x in SEG[c].split(".")])
    last = n_aw + 1
    rng = lambda col: f"Awarded!${col}$2:${col}${last}"  # noqa: E731
    cols = with_english(TEMPLATE + EXTRA)
    L = lambda name: get_column_letter(cols.index(name) + 1)  # noqa: E731
    code_col, ctry_col, tot_col, cur_col = L("Indicator code"), L("Country"), L("Total awarded"), L("Currency")
    fund_col = L("Funding (national / World Bank / IDB / Global Fund / PAHO)")
    for i, c in enumerate(codes, 5):
        ws.cell(row=i, column=1, value=c)
        ws.cell(row=i, column=2, value=SEG[c])
        ws.cell(row=i, column=3, value=next(k["Segment name"] for k in KMAP if k["Indicator code"] == c))
        ws.cell(row=i, column=4, value=f'=COUNTIFS({rng(code_col)},$A{i},{rng(ctry_col)},"Honduras",{rng(fund_col)},"National")')
        ws.cell(row=i, column=5, value=f'=SUMIFS({rng(tot_col)},{rng(code_col)},$A{i},{rng(ctry_col)},"Honduras",{rng(cur_col)},"HNL",{rng(fund_col)},"National")')
        ws.cell(row=i, column=6, value=f'=COUNTIFS({rng(code_col)},$A{i},{rng(ctry_col)},"El Salvador",{rng(fund_col)},"National")')
        ws.cell(row=i, column=7, value=f'=SUMIFS({rng(tot_col)},{rng(code_col)},$A{i},{rng(ctry_col)},"El Salvador",{rng(cur_col)},"USD",{rng(fund_col)},"National")')
        ws.cell(row=i, column=8, value=f'=COUNTIFS({rng(code_col)},$A{i})-D{i}-F{i}')
        for j in range(1, 9):
            ws.cell(row=i, column=j).font = BODY
        for j in (5, 7):
            ws.cell(row=i, column=j).number_format = "#,##0"
    t = 5 + len(codes)
    ws.cell(row=t, column=1, value="Total").font = BOLD
    for j, col in zip(range(4, 9), "DEFGH"):
        ws.cell(row=t, column=j, value=f"=SUM({col}5:{col}{t - 1})").font = BOLD
        if col in "EG":
            ws.cell(row=t, column=j).number_format = "#,##0"
    # by year
    y0 = t + 3
    ws.cell(row=y0 - 1, column=1, value="By award year (national sources)").font = TITLE
    hdr2 = ["Award year", "HN lines", "HN total (HNL)", "ES lines", "ES total (USD)"]
    for j, h in enumerate(hdr2, 1):
        c = ws.cell(row=y0, column=j, value=h)
        c.fill, c.font = HDR_FILL, HDR_FONT
    yr_col = L("Award year")
    for k, y in enumerate(range(2019, 2027), 1):
        i = y0 + k
        ws.cell(row=i, column=1, value=y).font = BODY
        ws.cell(row=i, column=2, value=f'=COUNTIFS({rng(yr_col)},$A{i},{rng(ctry_col)},"Honduras",{rng(fund_col)},"National")')
        ws.cell(row=i, column=3, value=f'=SUMIFS({rng(tot_col)},{rng(yr_col)},$A{i},{rng(ctry_col)},"Honduras",{rng(cur_col)},"HNL",{rng(fund_col)},"National")')
        ws.cell(row=i, column=4, value=f'=COUNTIFS({rng(yr_col)},$A{i},{rng(ctry_col)},"El Salvador",{rng(fund_col)},"National")')
        ws.cell(row=i, column=5, value=f'=SUMIFS({rng(tot_col)},{rng(yr_col)},$A{i},{rng(ctry_col)},"El Salvador",{rng(cur_col)},"USD",{rng(fund_col)},"National")')
        for j in range(2, 6):
            ws.cell(row=i, column=j).font = BODY
        for j in (3, 5):
            ws.cell(row=i, column=j).number_format = "#,##0"
    for col, w in zip("ABCDEFGH", (10, 12, 34, 9, 16, 9, 16, 12)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A5"


def notes_sheet(wb, notes):
    ws = wb.create_sheet("Sources & notes", 0)
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 120
    r = 1
    ws.cell(row=r, column=1, value="El Salvador & Honduras – public tenders for IVD and diabetes monitoring, 2019–2026").font = TITLE
    r += 1
    ws.cell(row=r, column=1, value=f"Built {date.today().isoformat()} by Clearstate from public sources. Codes follow ES_HN_tender_keyword_map_by_segment.xlsx.").font = Font(name="Arial", size=9, italic=True)
    r += 2
    for section, items in notes:
        ws.cell(row=r, column=1, value=section).font = Font(bold=True, name="Arial", size=10, color="1F3864")
        r += 1
        for k, v in items:
            a = ws.cell(row=r, column=1, value=k)
            b = ws.cell(row=r, column=2, value=v)
            a.font, b.font = BOLD, BODY
            a.alignment = Alignment(vertical="top", wrap_text=True)
            b.alignment = Alignment(vertical="top", wrap_text=True)
            r += 1
        r += 1


if __name__ == "__main__":
    import json
    d = pd.read_pickle(os.path.join(HERE, "eshn_rows.pkl"))
    aw = sorted(d["awarded"], key=lambda r: (r["Country"], str(r.get("Award date") or ""), str(r.get("Tender / procedure no."))))
    op = sorted(d["open"], key=lambda r: (r["Country"], str(r.get("Publication date") or "")), reverse=False)
    notes = json.load(open(NOTES, encoding="utf-8")) if os.path.exists(NOTES) else []
    tmap, missing = load_translations(), set()
    add_english(aw, tmap, missing)
    add_english(op, tmap, missing)
    notes.insert(1, ["English translation", [
        ("Columns", "Every Spanish text column (tender title, product / test, procedure type, status, unit, contract type, duration) is followed by "
                    "an '(English)' column (green header). Originals are unchanged."),
        ("How", "Translated in full (not summarised) from the original text; brand names, catalogue/model/registration numbers, process numbers, "
                "quantities and units are kept as published. Codes, dates and text that was already English are copied as-is."),
        ("Coverage", f"{len(tmap):,} unique strings translated or copied; {len(missing):,} left blank"),
    ]])
    wb = Workbook()
    wb.remove(wb.active)
    notes_sheet(wb, notes)
    data_sheet(wb, "Awarded", aw, with_english(TEMPLATE + EXTRA))
    summary_sheet(wb, aw, len(aw))
    op_cols = [c for c in TEMPLATE if c not in ("Awarded supplier (distributor)", "Unit price", "Total awarded", "Award date")] + ["Closing date", "Tag confidence", "All matched codes", "Tagged from", "Buyer segment"]
    data_sheet(wb, "Open tenders", op, with_english(op_cols))
    wb.save(OUT)
    print("saved", OUT, len(aw), "awarded,", len(op), "open; untranslated strings:", len(missing))
