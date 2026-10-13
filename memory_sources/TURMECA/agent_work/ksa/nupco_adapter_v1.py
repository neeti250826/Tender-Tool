from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import time
import urllib.parse
import urllib.request
from dataclasses import dataclass, asdict
from datetime import datetime
from pathlib import Path
from typing import Any

from lxml import html
from openpyxl import load_workbook


ADAPTER_VERSION = "1.0.0"
LISTING_URL = "https://www.nupco.com/tenders/tenders-list/"
FINAL_RESULTS_PARAM = "e-filter-3085e1f-tender-status=final-results"
CANONICAL_COLUMNS = [
    "source", "country", "country_code", "publication_date", "title",
    "title_english", "description", "description_english", "buyer",
    "buyer_english", "classification", "status", "currency", "amount",
    "awarding_agency_name", "awarding_agency_name_english", "supplier_name",
    "supplier_name_english", "awarded_date", "contract_period",
    "contract_period_english", "item_no", "item_description",
    "item_description_english", "item_uom", "item_uom_english", "quantity",
    "notice_id", "notice_url", "_source_workbook", "_duplicate_rows_removed",
    "_dedupe_key", "medtech_filter_basis", "mapped_oem", "mapped_oem_basis",
    "mapped_oem_match_term", "likely_oem", "likely_oem_score",
    "oem_match_terms", "oem_match_basis", "medtech_confidence",
]


@dataclass
class Notice:
    notice_id: str
    status: str
    title: str
    notice_url: str
    submission_deadline: str = ""
    bid_opening: str = ""
    page: int = 0
    published: str = ""
    modified: str = ""
    final_result_url: str = ""
    preliminary_result_url: str = ""
    item_list_url: str = ""
    terms_url: str = ""


class Adapter:
    """Trusted NUPCO adapter for the current public WordPress tender pages.

    The public REST API returns 401, so the adapter uses stable official HTML,
    the explicit Final Results taxonomy filter, bounded pagination, and official
    attachment links. Candidate rows remain based on the immutable workbook;
    website evidence is used to validate status/coverage and identify gaps.
    """

    def __init__(self, url: str = LISTING_URL, output_root: str | Path | None = None,
                 workbook: str | Path | None = None):
        self.url = url or LISTING_URL
        self.output_root = Path(output_root or Path(__file__).resolve().parent)
        self.workbook = Path(workbook) if workbook else None
        self.raw_root = self.output_root / "raw_evidence"
        self.final_root = self.raw_root / "final_results"
        self.detail_root = self.raw_root / "details"
        self.listing_root = self.raw_root / "listings"
        for p in (self.raw_root, self.final_root, self.detail_root, self.listing_root):
            p.mkdir(parents=True, exist_ok=True)
        self.user_agent = "Mozilla/5.0 (compatible; NUPCO-evidence-adapter/1.0; +https://www.nupco.com/)"

    def _fetch(self, url: str, retries: int = 3) -> bytes:
        parts = urllib.parse.urlsplit(url)
        url = urllib.parse.urlunsplit((parts.scheme, parts.netloc, urllib.parse.quote(urllib.parse.unquote(parts.path), safe="/%"), parts.query, parts.fragment))
        last: Exception | None = None
        for attempt in range(retries):
            try:
                req = urllib.request.Request(url, headers={"User-Agent": self.user_agent, "Accept": "*/*"})
                with urllib.request.urlopen(req, timeout=25) as response:
                    return response.read()
            except Exception as exc:
                last = exc
                time.sleep(1.5 * (attempt + 1))
        raise RuntimeError(f"Fetch failed after {retries} attempts: {url}: {last}")

    @staticmethod
    def _clean(value: Any) -> str:
        if value is None:
            return ""
        if isinstance(value, datetime):
            return value.date().isoformat()
        return str(value).strip()

    @staticmethod
    def _sha256(path: Path) -> str:
        h = hashlib.sha256()
        with path.open("rb") as fh:
            for chunk in iter(lambda: fh.read(1024 * 1024), b""):
                h.update(chunk)
        return h.hexdigest()

    @staticmethod
    def _slug(notice_id: str) -> str:
        return re.sub(r"[^A-Za-z0-9._-]+", "_", notice_id).strip("_") or "unknown"

    @staticmethod
    def _iso_date(value: str) -> str:
        value = (value or "").strip()
        if not value:
            return ""
        for fmt in ("%a, %d/%m/%Y", "%d/%m/%Y", "%Y-%m-%d", "%Y-%m-%dT%H:%M:%S%z"):
            try:
                return datetime.strptime(value, fmt).date().isoformat()
            except ValueError:
                pass
        m = re.match(r"(\d{4}-\d{2}-\d{2})", value)
        return m.group(1) if m else value

    def _listing_url(self, page: int) -> str:
        base = self.url.rstrip("/") + "/"
        if page > 1:
            base += f"{page}/"
        return f"{base}?{FINAL_RESULTS_PARAM}"

    def _parse_listing(self, raw: bytes, page: int) -> tuple[list[Notice], int]:
        doc = html.fromstring(raw)
        max_page = 1
        anchors = doc.xpath('//div[contains(@class,"e-load-more-anchor")]')
        if anchors:
            try:
                max_page = int(anchors[0].get("data-max-page") or 1)
            except ValueError:
                max_page = 1
        notices: list[Notice] = []
        for link in doc.xpath('//a[contains(@href,"/tender/") and normalize-space()]'):
            url = link.get("href") or ""
            title = " ".join(link.text_content().split())
            if not url or title.lower() == "discover more":
                continue
            container = link
            for _ in range(8):
                parent = container.getparent()
                if parent is None:
                    break
                container = parent
                txt = " ".join(container.text_content().split())
                if "Tender ID" in txt and "Final Results" in txt:
                    break
            txt = " ".join(container.text_content().split())
            id_match = re.search(r"\b(?:NPT|NDP)\d{4}[-/]\d{2}\b", txt)
            if not id_match:
                continue
            notice_id = id_match.group(0).replace("-", "/")
            sub = re.search(r"Submission Deadline:\s*(?:[A-Za-z]{3},\s*)?(\d{2}/\d{2}/\d{4})", txt)
            opening = re.search(r"Bid Opening:\s*(?:[A-Za-z]{3},\s*)?(\d{2}/\d{2}/\d{4})", txt)
            notices.append(Notice(
                notice_id=notice_id,
                status="Final Results",
                title=title,
                notice_url=url,
                submission_deadline=self._iso_date(sub.group(1) if sub else ""),
                bid_opening=self._iso_date(opening.group(1) if opening else ""),
                page=page,
            ))
        by_id: dict[str, Notice] = {}
        for notice in notices:
            by_id.setdefault(notice.notice_id, notice)
        return list(by_id.values()), max_page

    def _parse_detail(self, notice: Notice, raw: bytes) -> Notice:
        doc = html.fromstring(raw)
        text = raw.decode("utf-8", "replace")
        pub = re.search(r'"datePublished":"([^"]+)"', text)
        mod = re.search(r'"dateModified":"([^"]+)"', text)
        notice.published = self._iso_date(pub.group(1) if pub else "")
        notice.modified = self._iso_date(mod.group(1) if mod else "")
        for link in doc.xpath('//a[@href]'):
            href = urllib.parse.urljoin(notice.notice_url, link.get("href") or "")
            label = " ".join(link.text_content().split()).lower()
            if not href:
                continue
            if "tender final result" in label or ("final" in label and "result" in label and ".pdf" in href.lower()):
                notice.final_result_url = href
            elif "preliminary" in label and "result" in label:
                notice.preliminary_result_url = href
            elif "item list" in label:
                notice.item_list_url = href
            elif "terms" in label and "condition" in label:
                notice.terms_url = href
        return notice

    def crawl(self) -> list[Notice]:
        notices: list[Notice] = []
        first = self._fetch(self._listing_url(1))
        first_path = self.listing_root / "final-results-page-1.html"
        first_path.write_bytes(first)
        page_notices, max_page = self._parse_listing(first, 1)
        notices.extend(page_notices)
        if max_page > 10:
            raise RuntimeError(f"Unexpected max page {max_page}; aborting bounded crawl")
        for page in range(2, max_page + 1):
            raw = self._fetch(self._listing_url(page))
            (self.listing_root / f"final-results-page-{page}.html").write_bytes(raw)
            part, _ = self._parse_listing(raw, page)
            notices.extend(part)
        deduped: dict[str, Notice] = {n.notice_id: n for n in notices}
        for index, notice in enumerate(deduped.values(), start=1):
            detail_path = self.detail_root / f"{self._slug(notice.notice_id)}.html"
            raw = detail_path.read_bytes() if detail_path.exists() else self._fetch(notice.notice_url)
            if not detail_path.exists():
                detail_path.write_bytes(raw)
            self._parse_detail(notice, raw)
            if notice.final_result_url:
                pdf_path = self.final_root / f"{self._slug(notice.notice_id)}-final-result.pdf"
                try:
                    if not pdf_path.exists():
                        pdf_path.write_bytes(self._fetch(notice.final_result_url))
                except Exception as exc:
                    (self.final_root / f"{self._slug(notice.notice_id)}-download-error.txt").write_text(str(exc), encoding="utf-8")
            if index % 10 == 0:
                time.sleep(0.5)
        return sorted(deduped.values(), key=lambda n: n.notice_id)

    def _canonical_rows_from_workbook(self) -> list[dict[str, Any]]:
        if not self.workbook:
            return []
        wb = load_workbook(self.workbook, read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        iterator = ws.iter_rows(values_only=True)
        headers = [self._clean(v) for v in next(iterator)]
        idx = {name: i for i, name in enumerate(headers)}
        rows: list[dict[str, Any]] = []
        counters: dict[str, int] = {}
        for values in iterator:
            status = self._clean(values[idx["status_en"]])
            if status != "Final Result":
                continue
            notice_id = self._clean(values[idx["notice_id"]]).replace("-", "/")
            counters[notice_id] = counters.get(notice_id, 0) + 1
            item_no = counters[notice_id]
            publication_date = self._iso_date(self._clean(values[idx["publication_date"]]))
            supplier = self._clean(values[idx["supplier_name"]])
            item_desc = self._clean(values[idx["item_description"]])
            amount = self._clean(values[idx["amount"]])
            row = {col: "" for col in CANONICAL_COLUMNS}
            row.update({
                "source": "ksa_nupco",
                "country": "Saudi Arabia",
                "country_code": "SA",
                "publication_date": publication_date,
                "title": self._clean(values[idx["title"]]),
                "title_english": self._clean(values[idx["title_en"]]),
                "description": self._clean(values[idx["description"]]),
                "description_english": self._clean(values[idx["description_en"]]),
                "buyer": self._clean(values[idx["buyer"]]),
                "buyer_english": self._clean(values[idx["buyer"]]),
                "classification": "medical / medtech (source workbook)",
                "status": "Final Result",
                "currency": self._clean(values[idx["currency"]]),
                "amount": amount,
                "awarding_agency_name": self._clean(values[idx["awarding_agency_name"]]),
                "awarding_agency_name_english": self._clean(values[idx["awarding_agency_name_en"]]),
                "supplier_name": supplier,
                "supplier_name_english": self._clean(values[idx["supplier_name_en"]]),
                "awarded_date": publication_date,
                "contract_period": self._clean(values[idx["contract_period"]]),
                "contract_period_english": self._clean(values[idx["contract_period"]]),
                "item_no": item_no,
                "item_description": item_desc,
                "item_description_english": self._clean(values[idx["item_description_en"]]),
                "item_uom": self._clean(values[idx["item_uom"]]),
                "item_uom_english": self._clean(values[idx["item_uom"]]),
                "quantity": self._clean(values[idx["item_quantity"]]),
                "notice_id": notice_id,
                "notice_url": self._clean(values[idx["notice_url"]]),
                "_source_workbook": str(self.workbook),
                "_duplicate_rows_removed": 0,
                "_dedupe_key": "|".join([notice_id, str(item_no), supplier, item_desc, amount]),
                "medtech_filter_basis": "existing NUPCO award workbook supplied by user; row retained only when status_en=Final Result",
                "medtech_confidence": "high",
            })
            rows.append(row)
        wb.close()
        return rows

    def extract(self, source: Any = None) -> list[dict[str, Any]]:
        return self._canonical_rows_from_workbook()

    def write_outputs(self) -> None:
        notices = self.crawl()
        notice_path = self.output_root / "official_final_result_notices.json"
        notice_path.write_text(json.dumps([asdict(n) for n in notices], indent=2, ensure_ascii=False), encoding="utf-8")

        rows = self.extract(None)
        candidate_path = self.output_root / "candidate_rows.csv"
        with candidate_path.open("w", newline="", encoding="utf-8-sig") as fh:
            writer = csv.DictWriter(fh, fieldnames=CANONICAL_COLUMNS)
            writer.writeheader()
            writer.writerows(rows)

        workbook_ids = sorted({r["notice_id"] for r in rows})
        official_ids = sorted(n.notice_id for n in notices if n.notice_id.startswith("NPT"))
        ndp_ids = sorted(n.notice_id for n in notices if n.notice_id.startswith("NDP"))
        first_result_ids: list[str] = []
        if self.workbook:
            wb = load_workbook(self.workbook, read_only=True, data_only=True)
            ws = wb[wb.sheetnames[0]]
            it = ws.iter_rows(values_only=True)
            hdr = [self._clean(v) for v in next(it)]
            ix = {name: i for i, name in enumerate(hdr)}
            first_result_ids = sorted({self._clean(v[ix["notice_id"]]).replace("-", "/") for v in it if self._clean(v[ix["status_en"]]) == "First Result"})
            wb.close()
        comparison = {
            "adapter_version": ADAPTER_VERSION,
            "as_of": datetime.now().astimezone().isoformat(),
            "official_listing_url": self._listing_url(1),
            "official_final_result_notice_count": len(official_ids),
            "workbook_final_result_notice_count": len(workbook_ids),
            "workbook_first_result_notice_ids": first_result_ids,
            "official_npt_missing_from_workbook_final_rows": sorted(set(official_ids) - set(workbook_ids)),
            "workbook_not_on_current_official_final_listing": sorted(set(workbook_ids) - set(official_ids)),
            "official_ndp_final_result_ids": ndp_ids,
            "ndp_policy": "assessed separately; not merged into candidate rows without review",
            "candidate_row_count": len(rows),
        }
        (self.output_root / "coverage_comparison.json").write_text(json.dumps(comparison, indent=2, ensure_ascii=False), encoding="utf-8")

        manifest: list[dict[str, Any]] = []
        for path in sorted(self.raw_root.rglob("*")):
            if path.is_file():
                manifest.append({
                    "path": str(path.resolve()),
                    "sha256": self._sha256(path),
                    "size_bytes": path.stat().st_size,
                })
        for path in (notice_path, candidate_path, self.output_root / "coverage_comparison.json"):
            manifest.append({"path": str(path.resolve()), "sha256": self._sha256(path), "size_bytes": path.stat().st_size})
        (self.output_root / "evidence_hashes.json").write_text(json.dumps(manifest, indent=2, ensure_ascii=False), encoding="utf-8")


def create_adapter(url: str) -> Adapter:
    return Adapter(url=url)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--url", default=LISTING_URL)
    parser.add_argument("--output", required=True)
    parser.add_argument("--workbook", required=True)
    args = parser.parse_args()
    Adapter(args.url, args.output, args.workbook).write_outputs()


if __name__ == "__main__":
    main()
