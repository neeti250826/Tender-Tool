"""
Dubai eSupply (past opportunities, buyers containing 'Health': Dubai Academic Health Corporation + Dubai Health Authority)
-> UAE_Dubai_eSupply_Volume.xlsx. VOLUME SIGNALS ONLY (Phase 4 rule): eSupply publishes requested items and
quantities but no award, winner or price, so no value is built.

Item lines are parsed from each opportunity's Description: "<n>: [<item code>] <ITEM TEXT> - <qty> <UOM>".
Category from the item text only (vocab layer). Year = publication year.
"""
import glob, os, re, sys, warnings
from collections import Counter, defaultdict

warnings.filterwarnings("ignore")
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))
from vocab_rules import Vocab  # noqa: E402

LINE = re.compile(r"^\s*(\d+)\s*:\s*(?:\[([^\]]+)\])?\s*(.+?)\s*-+\s*([\d.,]+)\s*([A-Za-z][A-Za-z ()/]{0,20})?\s*$")


def main():
    v = Vocab()
    opp, items = {}, []
    for f in sorted(glob.glob(os.path.join(HERE, "daily", "*.xlsx"))):
        ws = openpyxl.load_workbook(f).worksheets[0]
        rows = list(ws.iter_rows(values_only=True))
        h = next((i for i, r in enumerate(rows) if r and r[0] == "Project Code"), None)
        if h is None:
            continue
        H = list(rows[h])
        for r in rows[h + 1:]:
            if not r or not r[0] or r[0] in opp:
                continue
            d = dict(zip(H, r))
            opp[r[0]] = d
            for ln in str(d.get("Description") or "").splitlines():
                m = LINE.match(ln)
                if not m:
                    continue
                txt = m.group(3).strip(" -")
                pm = v.product(txt) or {}
                try:
                    qty = float(m.group(4).replace(",", ""))
                except ValueError:
                    qty = None
                pub = d.get("Publication Date")
                items.append({"Project Code": r[0], "Project Title": d.get("Project Title"), "Buyer": d.get("Buyer Organisation"),
                              "Publication Date": pub, "Year": pub.year if hasattr(pub, "year") else None, "Line": int(m.group(1)),
                              "Item Code": m.group(2), "Item": txt, "Qty": qty, "UOM": (m.group(5) or "").strip(),
                              "OU": pm.get("ou"), "Sub-OU": pm.get("sub_ou"), "Device Category": pm.get("product_category"),
                              "Tender Stage": "Pipeline (closed RFQ; award not published)", "Price Evidence": "None", "Manufacturer Evidence": "None"})
    mapped = [i for i in items if i["Device Category"]]
    agg = defaultdict(lambda: defaultdict(float))
    lines = defaultdict(Counter)
    for i in mapped:
        k = (i["OU"], i["Sub-OU"], i["Device Category"])
        agg[k][i["Year"]] += i["Qty"] or 0
        lines[k][i["Year"]] += 1
    cat_rows = [[*k, lines[k][2024], agg[k][2024], lines[k][2025], agg[k][2025], lines[k][2026], agg[k][2026]]
                for k in sorted(agg, key=lambda k: (k[0] or "", k[1] or "", k[2] or ""))]

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    hdr = PatternFill("solid", fgColor="1F4E78")

    def sheet(title, header, data, widths):
        ws = wb.create_sheet(title)
        ws.append(header)
        for c in ws[1]:
            c.font, c.fill, c.alignment = Font(bold=True, color="FFFFFF"), hdr, Alignment(wrap_text=True)
        for r in data:
            ws.append(r)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(header))}{len(data)+1}"
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    ws = wb.create_sheet("README")
    for line in [
        "Dubai eSupply past opportunities - Dubai Academic Health Corporation + Dubai Health Authority, publication 2024-01-01 to 2026-10-01",
        f"Opportunities: {len(opp)}; item lines parsed: {len(items)}; mapped to taxonomy: {len(mapped)}.",
        "VOLUME SIGNAL ONLY: requested quantities per RFQ; eSupply publishes no award, winner or price, so no value is built (Phase 4 rule).",
        "UOMs are as written by the buyer (EA, BOX, PKT ...); quantities are not converted to pieces. RFQs can be re-issued, so totals may double-count.",
        "Exported per publication day with eSupply's own 'Export List' (the list caps exports at 2,000 rows).",
    ]:
        ws.append([line])
    ws.column_dimensions["A"].width = 150
    cols = list(items[0].keys()) if items else ["(none)"]
    sheet("Mapped_Items", cols, [[i[c] for c in cols] for i in mapped], [16, 40, 30, 16, 6, 5, 14, 60, 8, 8, 26, 30, 36, 30, 8, 8])
    sheet("By_Category_Year", ["OU", "Sub-OU", "Device Category", "Lines 2024", "Qty 2024", "Lines 2025", "Qty 2025", "Lines 2026", "Qty 2026"],
          cat_rows, [28, 35, 40, 9, 11, 9, 11, 9, 11])
    unm = Counter(i["Item"][:100] for i in items if not i["Device Category"])
    sheet("Unmapped_Top_Items", ["Item (unmapped)", "Lines"], unm.most_common(3000), [100, 8])
    wb.save(os.path.join(HERE, "UAE_Dubai_eSupply_Volume.xlsx"))
    print(f"opportunities {len(opp)}; item lines {len(items)}; mapped {len(mapped)}; categories {len(cat_rows)}")


if __name__ == "__main__":
    main()
