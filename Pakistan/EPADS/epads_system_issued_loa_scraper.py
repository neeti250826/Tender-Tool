#!/usr/bin/env python3
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime, timezone
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Tuple
from urllib.parse import urljoin

import requests
from bs4 import BeautifulSoup, Tag
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover
    Workbook = None

ROOT_DIR = Path(__file__).resolve().parents[2]
if str(ROOT_DIR) not in sys.path:
    sys.path.insert(0, str(ROOT_DIR))

from procurement_utils import clean_text, parse_date_to_iso, stable_key, utc_now_iso  # noqa: E402

BASE_URL = "https://www.epads.gov.pk"
LIST_URL = f"{BASE_URL}/loi-issued"
DEFAULT_OUTPUT_STEM = "epads_system_issued_loas_all_present"

FIELDNAMES = [
    "source",
    "country",
    "country_code",
    "notice_id",
    "notice_url",
    "publication_date",
    "publication_date_raw",
    "closing_date",
    "closing_date_raw",
    "title",
    "buyer",
    "procurement_cell",
    "department",
    "classification",
    "procurement_type",
    "procedure",
    "procurement_method",
    "estimated_cost",
    "amount",
    "total_award_amount",
    "award_amount_difference",
    "currency",
    "status",
    "supplier_name",
    "supplier_email",
    "supplier_phone",
    "supplier_organization_type",
    "supplier_scope",
    "awarded_date",
    "awarded_date_raw",
    "loa_response_date",
    "loa_response_date_raw",
    "winner_sequence",
    "item_no",
    "item_description",
    "item_unspsc",
    "item_quantity",
    "item_unit_price",
    "item_awarded_value",
    "award_reason",
    "awarded_value_detail",
    "raw_item_text",
    "list_page",
    "list_row_number",
    "scraped_at_utc",
    "dedup_key",
]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Extract every currently listed Pakistan EPADS System-Issued LoA row, "
            "expanding multiple winning bidders and their awarded item rows."
        )
    )
    parser.add_argument("--output-dir", default="Pakistan/EPADS/output")
    parser.add_argument("--output-stem", default=DEFAULT_OUTPUT_STEM)
    parser.add_argument("--cache-dir", default="Pakistan/EPADS/output/loa_cache")
    parser.add_argument("--max-pages", type=int, default=0, help="0 means auto-discover all pages.")
    parser.add_argument("--max-details", type=int, default=0, help="0 means scrape all detail pages.")
    parser.add_argument("--workers", type=int, default=6)
    parser.add_argument("--timeout", type=int, default=90)
    parser.add_argument("--sleep", type=float, default=0.0, help="Optional delay after each HTTP fetch.")
    parser.add_argument("--refresh", action="store_true", help="Ignore cached HTML and refetch.")
    parser.add_argument("--list-only", action="store_true", help="Only collect and export the list rows.")
    return parser.parse_args()


def build_session() -> requests.Session:
    session = requests.Session()
    session.headers.update(
        {
            "User-Agent": (
                "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36"
            ),
            "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
        }
    )
    retry = Retry(
        total=4,
        read=4,
        connect=4,
        backoff_factor=1.0,
        status_forcelist=(429, 500, 502, 503, 504),
        allowed_methods=frozenset({"GET"}),
        raise_on_status=False,
    )
    adapter = HTTPAdapter(max_retries=retry, pool_connections=20, pool_maxsize=20)
    session.mount("https://", adapter)
    session.mount("http://", adapter)
    return session


def fetch_url(session: requests.Session, url: str, cache_path: Path, timeout: int, refresh: bool, sleep: float) -> str:
    if cache_path.exists() and not refresh:
        return cache_path.read_text(encoding="utf-8", errors="ignore")
    response = session.get(url, timeout=timeout)
    response.raise_for_status()
    cache_path.parent.mkdir(parents=True, exist_ok=True)
    cache_path.write_text(response.text, encoding="utf-8")
    if sleep:
        time.sleep(sleep)
    return response.text


def list_page_url(page_number: int) -> str:
    return LIST_URL if page_number <= 1 else f"{LIST_URL}?page={page_number}"


def parse_money(value: str) -> str:
    text = clean_text(value)
    match = re.search(r"-?\d[\d,]*(?:\.\d+)?", text)
    return match.group(0).replace(",", "") if match else ""


def extract_labeled_text(text: str, label: str) -> str:
    cleaned = clean_text(text)
    pattern = rf"{re.escape(label)}\s*:?\s*(.*?)(?=\s+[A-Z][A-Za-z ]+\s*:|$)"
    match = re.search(pattern, cleaned, flags=re.IGNORECASE)
    return clean_text(match.group(1)) if match else ""


def extract_first_date_raw(text: str) -> str:
    cleaned = clean_text(text)
    patterns = [
        r"(?:Monday|Tuesday|Wednesday|Thursday|Friday|Saturday|Sunday),\s+[A-Za-z]{3,9}\s+\d{1,2},\s+\d{4}(?:\s+\d{1,2}:\d{2}\s*(?:AM|PM)?)?",
        r"[A-Za-z]{3,9}\s+\d{1,2},\s+\d{4}(?:\s+\d{1,2}:\d{2}\s*(?:AM|PM)?)?",
        r"\d{4}-\d{2}-\d{2}",
        r"\d{1,2}/\d{1,2}/\d{4}",
    ]
    for pattern in patterns:
        match = re.search(pattern, cleaned, flags=re.IGNORECASE)
        if match:
            return clean_text(match.group(0))
    return ""


def normalize_epads_date(value: str) -> str:
    text = clean_text(value)
    if not text:
        return ""
    normalized = re.sub(
        r"^(Monday|Tuesday|Wednesday|Thursday|Friday|Saturday|Sunday),\s+",
        "",
        text,
        flags=re.IGNORECASE,
    )
    for fmt in (
        "%b %d, %Y %I:%M %p",
        "%B %d, %Y %I:%M %p",
        "%b %d, %Y",
        "%B %d, %Y",
        "%Y-%m-%d",
        "%d/%m/%Y",
        "%m/%d/%Y",
    ):
        try:
            return datetime.strptime(normalized, fmt).date().isoformat()
        except ValueError:
            continue
    return parse_date_to_iso(normalized)


def parse_date_label(text: str, label: str) -> Tuple[str, str]:
    value = extract_labeled_text(text, label)
    raw = extract_first_date_raw(value)
    return normalize_epads_date(raw), raw


def unique_detail_url_from_row(row: Tag) -> str:
    anchor = row.select_one('a[href*="/opportunities/federal/procurements/loi-issued/"]')
    return urljoin(BASE_URL, anchor.get("href", "")) if anchor else ""


def parse_list_rows(html: str, page_number: int) -> List[Dict[str, str]]:
    soup = BeautifulSoup(html, "html.parser")
    table = None
    for candidate in soup.select("table"):
        header_text = clean_text(candidate.get_text(" ", strip=True)).lower()
        if "procurement ref" in header_text and "procurement title" in header_text and "amount" in header_text:
            table = candidate
            break
    if table is None:
        return []

    rows: List[Dict[str, str]] = []
    for row in table.select("tbody tr"):
        cells = row.select("td")
        if len(cells) < 6:
            continue
        notice_url = unique_detail_url_from_row(row)
        if not notice_url:
            continue
        notice_id = clean_text(cells[1].get_text(" ", strip=True))
        title_anchor = cells[2].select_one('a[href*="/loi-issued/"]')
        title = clean_text(title_anchor.get_text(" ", strip=True)) if title_anchor else ""
        buyer_node = cells[2].select_one("[title]")
        buyer = clean_text(buyer_node.get("title", "")) if buyer_node else ""
        if not buyer:
            buyer = clean_text(cells[2].get_text(" ", strip=True).replace(title, "", 1))

        status_text = clean_text(cells[3].get_text(" ", strip=True))
        publication_date, publication_date_raw = parse_date_label(status_text, "Published On")
        closing_date, closing_date_raw = parse_date_label(status_text, "Closed On")
        amount_text = clean_text(cells[4].get_text(" ", strip=True))
        estimated_cost = parse_money(extract_labeled_text(amount_text, "Estimated Cost"))
        total_award_amount = parse_money(extract_labeled_text(amount_text, "Award Amount"))

        badge = cells[5].select_one(".badge")
        procurement_type = clean_text(badge.get_text(" ", strip=True)) if badge else ""
        procedure = ""
        for span in cells[5].select("span"):
            value = clean_text(span.get_text(" ", strip=True))
            if value and value != procurement_type:
                procedure = value
                break
        classification = " | ".join(part for part in [procurement_type, procedure] if part)

        rows.append(
            {
                "source": "EPADS System-Issued LoAs",
                "country": "Pakistan",
                "country_code": "PK",
                "notice_id": notice_id,
                "notice_url": notice_url,
                "publication_date": publication_date,
                "publication_date_raw": publication_date_raw,
                "closing_date": closing_date,
                "closing_date_raw": closing_date_raw,
                "title": title,
                "buyer": buyer,
                "classification": classification,
                "procurement_type": procurement_type,
                "procedure": procedure,
                "estimated_cost": estimated_cost,
                "amount": total_award_amount,
                "total_award_amount": total_award_amount,
                "currency": "PKR" if estimated_cost or total_award_amount else "",
                "status": "loi issued",
                "list_page": str(page_number),
                "list_row_number": clean_text(cells[0].get_text(" ", strip=True)),
            }
        )
    return rows


def discover_max_page(html: str) -> int:
    pages = [int(value) for value in re.findall(r"/loi-issued\?page=(\d+)", html)]
    return max(pages) if pages else 1


def card_values(soup: BeautifulSoup, title: str) -> List[str]:
    heading = soup.find(lambda tag: isinstance(tag, Tag) and tag.name == "h6" and clean_text(tag.get_text(" ", strip=True)).lower() == title.lower())
    if heading is None:
        return []
    container = heading.find_parent(["div"], class_=re.compile(r"\bbg-"))
    if container is None:
        container = heading.parent
    values = []
    for node in container.select(".proc-stat-value"):
        value = clean_text(node.get_text(" ", strip=True))
        if value:
            values.append(value)
    if not values:
        for node in container.find_all("h6"):
            value = clean_text(node.get_text(" ", strip=True))
            if value and value.lower() != title.lower():
                values.append(value)
    return values


def stat_label_values(soup: BeautifulSoup) -> Dict[str, str]:
    result: Dict[str, str] = {}
    for item in soup.select(".proc-stat-item"):
        label = item.select_one(".proc-stat-label")
        value = item.select_one(".proc-stat-value")
        if label and value:
            result[clean_text(label.get_text(" ", strip=True)).lower()] = clean_text(value.get_text(" ", strip=True))
    return result


def parse_item_name_and_unspsc(value: str) -> Tuple[str, str]:
    text = clean_text(value).replace("\xa0", " ")
    match = re.search(r"\bUNSPSC:\s*(.*)$", text, flags=re.IGNORECASE)
    unspsc = clean_text(match.group(1)) if match else ""
    item_name = clean_text(text[: match.start()] if match else text)
    return item_name, unspsc


def parse_winner_header(header: Tag) -> Dict[str, str]:
    header_text = clean_text(header.get_text(" ", strip=True))
    supplier_name = re.sub(r"\s*\(Winner\).*", "", header_text, flags=re.IGNORECASE).strip()
    supplier_name = re.sub(r"\s*Letter of Award.*", "", supplier_name, flags=re.IGNORECASE).strip()
    awarded_raw = ""
    response_raw = ""
    award_match = re.search(
        r"Letter of Award \(LoA\) Issuance Date\s*:\s*(.*?)(?=\s+Response Date\s*:|$)",
        header_text,
        flags=re.IGNORECASE,
    )
    if award_match:
        awarded_raw = extract_first_date_raw(award_match.group(1))
    response_match = re.search(r"Response Date\s*:\s*(.*)$", header_text, flags=re.IGNORECASE)
    if response_match:
        response_raw = extract_first_date_raw(response_match.group(1))
    return {
        "supplier_name": supplier_name,
        "awarded_date": normalize_epads_date(awarded_raw),
        "awarded_date_raw": awarded_raw,
        "loa_response_date": normalize_epads_date(response_raw),
        "loa_response_date_raw": response_raw,
    }


def parse_contact_values(item: Tag) -> Dict[str, str]:
    smalls = [clean_text(node.get_text(" ", strip=True)) for node in item.select(".accordion-body small")]
    smalls = [value for value in smalls if value]
    result = {
        "supplier_scope": smalls[0] if len(smalls) > 0 else "",
        "supplier_email": smalls[1] if len(smalls) > 1 else "",
        "supplier_phone": smalls[2] if len(smalls) > 2 else "",
        "supplier_organization_type": smalls[3] if len(smalls) > 3 else "",
    }
    if result["supplier_email"] and "@" not in result["supplier_email"]:
        result["supplier_email"] = ""
    return result


def parse_winner_items(item: Tag) -> List[Dict[str, str]]:
    tables = item.select(".ck-content-LoA table")
    if not tables:
        return []
    rows: List[Dict[str, str]] = []
    for table in tables:
        for tr in table.select("tbody tr"):
            cells = [clean_text(cell.get_text(" ", strip=True)) for cell in tr.select("td")]
            if len(cells) < 2:
                continue
            if not any(cells):
                continue
            if clean_text(cells[0]).lower().startswith("terms"):
                continue
            item_description, item_unspsc = parse_item_name_and_unspsc(cells[1])
            rows.append(
                {
                    "item_no": cells[0] if len(cells) > 0 else "",
                    "item_description": item_description,
                    "item_unspsc": item_unspsc,
                    "item_quantity": cells[2] if len(cells) > 2 else "",
                    "item_unit_price": parse_money(cells[3]) if len(cells) > 3 else "",
                    "item_awarded_value": parse_money(cells[4]) if len(cells) > 4 else "",
                    "award_reason": cells[5] if len(cells) > 5 else "",
                    "raw_item_text": " | ".join(cells),
                }
            )
    return rows


def parse_detail_rows(html: str, listing: Dict[str, str], scraped_at: str) -> List[Dict[str, str]]:
    soup = BeautifulSoup(html, "html.parser")
    full_title = clean_text(soup.select_one("h4").get_text(" ", strip=True)) if soup.select_one("h4") else listing.get("title", "")
    stat_values = stat_label_values(soup)
    publication_raw = stat_values.get("publication date", listing.get("publication_date_raw", ""))
    closing_raw = listing.get("closing_date_raw", "")
    estimated_cost = parse_money(stat_values.get("estimated cost", "")) or listing.get("estimated_cost", "")
    total_award_amount = parse_money(stat_values.get("award amount", "")) or listing.get("total_award_amount", "")
    difference = parse_money(stat_values.get("difference", ""))
    procurement_method_values = card_values(soup, "Procurement Method")
    procurement_method = " | ".join(procurement_method_values)
    procurement_cell = " | ".join(card_values(soup, "Procurement Cell"))
    department = " | ".join(card_values(soup, "Department"))
    procurement_type = " | ".join(card_values(soup, "Procurement Category")) or listing.get("procurement_type", "")

    base = dict(listing)
    base.update(
        {
            "title": full_title or listing.get("title", ""),
            "publication_date": normalize_epads_date(publication_raw) or listing.get("publication_date", ""),
            "publication_date_raw": publication_raw or listing.get("publication_date_raw", ""),
            "closing_date": normalize_epads_date(closing_raw) or listing.get("closing_date", ""),
            "closing_date_raw": closing_raw,
            "procurement_cell": procurement_cell,
            "department": department,
            "procurement_type": procurement_type,
            "procurement_method": procurement_method,
            "estimated_cost": estimated_cost,
            "amount": total_award_amount,
            "total_award_amount": total_award_amount,
            "award_amount_difference": difference,
            "currency": "PKR" if estimated_cost or total_award_amount else listing.get("currency", ""),
        }
    )
    if not base.get("buyer"):
        base["buyer"] = procurement_cell or department
    if not base.get("classification"):
        base["classification"] = " | ".join(part for part in [procurement_type, listing.get("procedure", "")] if part)

    winner_items = []
    for sequence, accordion_item in enumerate(soup.select(".accordion-item"), start=1):
        header = accordion_item.select_one(".accordion-header")
        if header is None:
            continue
        header_text = clean_text(header.get_text(" ", strip=True))
        if "(Winner)" not in header_text and "Letter of Award" not in header_text:
            continue
        winner = parse_winner_header(header)
        contacts = parse_contact_values(accordion_item)
        items = parse_winner_items(accordion_item)
        if not items:
            items = [{"item_no": "", "item_description": "", "item_unspsc": "", "item_quantity": "", "item_unit_price": "", "item_awarded_value": "", "award_reason": "", "raw_item_text": ""}]
        for item_index, award_item in enumerate(items, start=1):
            row = dict(base)
            row.update(winner)
            row.update(contacts)
            row.update(award_item)
            row["winner_sequence"] = str(sequence)
            row["awarded_value_detail"] = (
                f"Estimated Cost: {estimated_cost}; Total Award Amount: {total_award_amount}; "
                f"Item Awarded Value: {award_item.get('item_awarded_value', '')}"
            )
            row["scraped_at_utc"] = scraped_at
            row["dedup_key"] = stable_key(
                row.get("notice_id", ""),
                row.get("supplier_name", ""),
                str(sequence),
                row.get("item_no", ""),
                row.get("item_description", ""),
                row.get("item_awarded_value", ""),
            )
            winner_items.append(row)

    if winner_items:
        return winner_items

    fallback = dict(base)
    fallback["scraped_at_utc"] = scraped_at
    fallback["dedup_key"] = stable_key(fallback.get("notice_id", ""), fallback.get("notice_url", ""), "no_winner_rows")
    return [fallback]


def normalize_row(row: Dict[str, str]) -> Dict[str, str]:
    return {field: clean_text(row.get(field, "")) for field in FIELDNAMES}


def write_jsonl(path: Path, rows: Iterable[Dict[str, str]]) -> None:
    with path.open("w", encoding="utf-8") as handle:
        for row in rows:
            handle.write(json.dumps(normalize_row(row), ensure_ascii=False) + "\n")


def write_csv(path: Path, rows: List[Dict[str, str]]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=FIELDNAMES)
        writer.writeheader()
        for row in rows:
            writer.writerow(normalize_row(row))


def write_xlsx(path: Path, rows: List[Dict[str, str]]) -> None:
    if Workbook is None:
        return
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "system_issued_loas"
    sheet.append(FIELDNAMES)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
    for row in rows:
        sheet.append([normalize_row(row).get(field, "") for field in FIELDNAMES])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sample_limit = min(sheet.max_row, 500)
    for column_index, column_name in enumerate(FIELDNAMES, start=1):
        values = [str(sheet.cell(row=row_index, column=column_index).value or "") for row_index in range(1, sample_limit + 1)]
        width = min(max(max(len(value) for value in values) + 2, 12), 70)
        sheet.column_dimensions[get_column_letter(column_index)].width = width
    workbook.save(path)


def write_outputs(output_dir: Path, output_stem: str, rows: List[Dict[str, str]], list_rows: List[Dict[str, str]]) -> Dict[str, object]:
    output_dir.mkdir(parents=True, exist_ok=True)
    csv_path = output_dir / f"{output_stem}.csv"
    jsonl_path = output_dir / f"{output_stem}.jsonl"
    xlsx_path = output_dir / f"{output_stem}.xlsx"
    summary_path = output_dir / f"{output_stem}_summary.json"
    write_csv(csv_path, rows)
    write_jsonl(jsonl_path, rows)
    write_xlsx(xlsx_path, rows)

    summary = {
        "generated_at_utc": datetime.now(timezone.utc).replace(microsecond=0).isoformat(),
        "source_url": LIST_URL,
        "scope": "All currently listed Pakistan EPADS System-Issued LoAs; no keyword filtering.",
        "list_tenders": len(list_rows),
        "output_rows": len(rows),
        "unique_notice_ids": len({row.get("notice_id", "") for row in rows if row.get("notice_id", "")}),
        "unique_suppliers": len({row.get("supplier_name", "") for row in rows if row.get("supplier_name", "")}),
        "notices_with_multiple_output_rows": sum(
            1
            for notice_id in {row.get("notice_id", "") for row in rows}
            if notice_id and sum(1 for row in rows if row.get("notice_id", "") == notice_id) > 1
        ),
        "coverage": {
            field: f"{sum(1 for row in rows if clean_text(row.get(field, '')))}/{len(rows)}"
            for field in FIELDNAMES
        },
        "outputs": {
            "csv": str(csv_path),
            "jsonl": str(jsonl_path),
            "xlsx": str(xlsx_path),
        },
    }
    summary_path.write_text(json.dumps(summary, indent=2, ensure_ascii=False), encoding="utf-8")
    summary["outputs"]["summary"] = str(summary_path)
    return summary


def collect_list_rows(args: argparse.Namespace, session: requests.Session, cache_dir: Path) -> Tuple[List[Dict[str, str]], int]:
    page1_html = fetch_url(session, list_page_url(1), cache_dir / "list_page_001.html", args.timeout, args.refresh, args.sleep)
    max_page = discover_max_page(page1_html)
    if args.max_pages:
        max_page = min(max_page, args.max_pages)
    all_rows = parse_list_rows(page1_html, 1)
    print(json.dumps({"event": "list_page", "page": 1, "rows": len(all_rows), "max_page": max_page}), flush=True)
    for page_number in range(2, max_page + 1):
        html = fetch_url(
            session,
            list_page_url(page_number),
            cache_dir / f"list_page_{page_number:03d}.html",
            args.timeout,
            args.refresh,
            args.sleep,
        )
        page_rows = parse_list_rows(html, page_number)
        print(json.dumps({"event": "list_page", "page": page_number, "rows": len(page_rows)}), flush=True)
        if not page_rows:
            break
        all_rows.extend(page_rows)

    unique: Dict[str, Dict[str, str]] = {}
    for row in all_rows:
        key = row.get("notice_url", "")
        if key and key not in unique:
            unique[key] = row
    return list(unique.values()), max_page


def detail_cache_path(cache_dir: Path, notice_id: str, notice_url: str) -> Path:
    numeric_id = re.sub(r"\D+", "", notice_id) or notice_url.rstrip("/").split("/")[-1]
    return cache_dir / "details" / f"{numeric_id}.html"


def fetch_and_parse_detail(args: argparse.Namespace, listing: Dict[str, str], scraped_at: str) -> List[Dict[str, str]]:
    session = build_session()
    cache_path = detail_cache_path(Path(args.cache_dir), listing.get("notice_id", ""), listing.get("notice_url", ""))
    html = fetch_url(session, listing["notice_url"], cache_path, args.timeout, args.refresh, args.sleep)
    return parse_detail_rows(html, listing, scraped_at)


def main() -> None:
    args = parse_args()
    output_dir = Path(args.output_dir)
    cache_dir = Path(args.cache_dir)
    session = build_session()
    scraped_at = utc_now_iso()

    list_rows, max_page = collect_list_rows(args, session, cache_dir)
    if args.max_details:
        list_rows = list_rows[: args.max_details]

    if args.list_only:
        rows = []
        for row in list_rows:
            normalized = dict(row)
            normalized["scraped_at_utc"] = scraped_at
            normalized["dedup_key"] = stable_key(row.get("notice_id", ""), row.get("notice_url", ""))
            rows.append(normalized)
        summary = write_outputs(output_dir, f"{args.output_stem}_list_only", rows, list_rows)
        summary["pages_scanned"] = max_page
        print(json.dumps(summary, indent=2, ensure_ascii=False), flush=True)
        return

    output_rows: List[Dict[str, str]] = []
    failures: List[Dict[str, str]] = []
    with ThreadPoolExecutor(max_workers=max(args.workers, 1)) as executor:
        futures = {executor.submit(fetch_and_parse_detail, args, row, scraped_at): row for row in list_rows}
        completed = 0
        for future in as_completed(futures):
            listing = futures[future]
            completed += 1
            try:
                detail_rows = future.result()
                output_rows.extend(detail_rows)
                print(
                    json.dumps(
                        {
                            "event": "detail",
                            "completed": completed,
                            "total": len(list_rows),
                            "notice_id": listing.get("notice_id", ""),
                            "rows": len(detail_rows),
                        }
                    ),
                    flush=True,
                )
            except Exception as exc:  # pragma: no cover - live-site resilience
                failures.append({"notice_id": listing.get("notice_id", ""), "notice_url": listing.get("notice_url", ""), "error": str(exc)})
                fallback = dict(listing)
                fallback["scraped_at_utc"] = scraped_at
                fallback["dedup_key"] = stable_key(fallback.get("notice_id", ""), fallback.get("notice_url", ""), "detail_error")
                output_rows.append(fallback)
                print(
                    json.dumps(
                        {
                            "event": "detail_error",
                            "completed": completed,
                            "total": len(list_rows),
                            "notice_id": listing.get("notice_id", ""),
                            "error": str(exc),
                        }
                    ),
                    flush=True,
                )

    output_rows = sorted(output_rows, key=lambda row: (int(row.get("list_page", "0") or 0), int(row.get("list_row_number", "0") or 0), row.get("winner_sequence", ""), row.get("item_no", "")))
    summary = write_outputs(output_dir, args.output_stem, output_rows, list_rows)
    summary["pages_scanned"] = max_page
    summary["detail_failures"] = failures
    (output_dir / f"{args.output_stem}_summary.json").write_text(json.dumps(summary, indent=2, ensure_ascii=False), encoding="utf-8")
    print(json.dumps(summary, indent=2, ensure_ascii=False), flush=True)


if __name__ == "__main__":
    main()
