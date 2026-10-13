"""Refresh CSV, manifests and local progress after read-only final verification."""
import csv,hashlib,html,json
from pathlib import Path
from openpyxl import load_workbook
import supervise_canlat_oct_pipeline as pipeline

ROOT=pipeline.ROOT;OUT=ROOT/'excel_exports'

def main():
    final=OUT/'CompraMX Tender FY27 Q2 Aug-Oct 2026 mapped.xlsx'
    verification=json.loads((OUT/'final_workbook_verification.json').read_text(encoding='utf-8'))
    digest=hashlib.sha256(final.read_bytes()).hexdigest()
    if not verification['passed'] or verification['workbook_sha256']!=digest:raise RuntimeError('Final workbook verification has not passed for current file')
    book=load_workbook(final,read_only=True,data_only=True)
    with (OUT/'mapped_tender_data.csv').open('w',encoding='utf-8-sig',newline='') as stream:csv.writer(stream).writerows(book['Tenders'].iter_rows(values_only=True))
    book.close()
    qc=json.loads((OUT/'qc_findings.json').read_text(encoding='utf-8'))
    checks={'workbook_verified':True,'source_cell_comparisons':verification['source_cells_checked'],'formula_references_and_caches_checked':verification['formula_caches_checked'],'rows':77895,'columns':39,'fx_mxn_to_usd':0.05,'all_urls_attempted':True,'successful_urls':13241,'failed_urls':98,'english_missing_rows':11908,'october_complete':False,'source_snapshot_latest':'2026-10-10 15:52:08','source_coverage_certified_complete':False,'manufacturer_cells_cleared_after_review':16,'family_cells_cleared_after_review':1,'manufacturer_mapping_comprehensive':False,'final_workbook_sha256':digest,'independent_final_review':'pending'}
    (ROOT/'release_checks.json').write_text(json.dumps(checks,indent=2),encoding='utf-8')
    manifest={}
    for file in OUT.iterdir():
        if file.is_file() and file.suffix in {'.xlsx','.csv','.json'}:manifest[file.name]={'bytes':file.stat().st_size,'sha256':hashlib.sha256(file.read_bytes()).hexdigest()}
    (ROOT/'deliverables_manifest.json').write_text(json.dumps(manifest,indent=2),encoding='utf-8')
    pipeline.update('best_available_workbook_verified_pending_final_review')
    handoff=ROOT/'handoff.html';body=handoff.read_text(encoding='utf-8')
    note='<h2>Workbook checks</h2><p>The workbook contains 77,895 item rows and 39 CANLAT columns. All source cells, 545,265 formula references and cached results passed the final deterministic check. The approved MXN-to-USD multiplier is 0.05.</p><h2>Remaining limitations</h2><p>98 notice URLs could not be extracted after retry. 11,908 rows still lack English descriptions because translation requests were limited; original Spanish descriptions are retained. October coverage ends on 10 October 2026. Taxonomy and manufacturer fields use historical reviewed evidence and require relevance review; full market or source coverage is not certified.</p><p><a href="excel_exports/final_workbook_verification.json">Workbook verification</a> · <a href="review_full_candidate.json">Independent source review</a> · <a href="excel_exports/manufacturer_corrections.json">Reviewed mapping corrections</a> · <a href="excel_exports/zero_cache_repair_audit.json">Zero-result cache correction audit</a></p>'
    body=body.replace('</body>',note+'</body>');handoff.write_text(body,encoding='utf-8')
    print(json.dumps(checks),flush=True)

if __name__=='__main__':main()
