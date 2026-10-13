"""Find direct canonical manufacturer-name matches missing from the mapped tender file."""

from __future__ import annotations

import argparse
import json
import re
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook

from map_canlat_taxonomy import (
    PLACEHOLDER_LABELS,
    TOKEN_PATTERN,
    token_text,
)


COMPANY_GENERIC_TERMS = {
    "ADVANCED", "BIO", "CONCEPT", "DEVICES", "DRIVE", "GENERAL", "GLOBAL", "HEALTH",
    "HEMOSTASIS", "KOREA", "MED", "MEDICAL", "MODERN", "RAPID", "ROBOTICS", "SPINE",
    "SURGICAL", "SYSTEMS", "TECHNOLOGIES", "TECHNOLOGY",
}


def company_phrase(value: object):
    tokens = TOKEN_PATTERN.findall(token_text(value))
    meaningful = [token for token in tokens if token not in COMPANY_GENERIC_TERMS]
    if not meaningful:
        return None
    if len(tokens) == 1 and (len(tokens[0]) < 5 or tokens[0] in COMPANY_GENERIC_TERMS):
        return None
    if len(tokens) > 1 and not any(len(token) >= 3 for token in meaningful):
        return None
    return " ".join(tokens)


def normalized_terms(value: object) -> str:
    return " " + " ".join(TOKEN_PATTERN.findall(token_text(value))) + " "


def has_brand_context(value: object, phrase: str) -> bool:
    text = token_text(value).replace("&", " ")
    excluded_context = (
        "COMPATIBLE",
        "MANTENIMIENTO",
        "PARA SU USO EN EL EQUIPO",
        "SE SUGIERE",
        "SUGERIDA",
        "SERVICIO",
    )
    if any(term in text for term in excluded_context):
        return False
    phrase_pattern = re.escape(phrase).replace(r"\ ", r"\s+")
    marker_pattern = rf"\b(?:MARCA|MCA|BRAND|FABRICANTE|MANUFACTURER)\b(?:\s|:|/|-|\.){{0,10}}{phrase_pattern}\b"
    slash_pattern = rf"/\s*{phrase_pattern}(?=\s*/|\s*-|\s*[A-Z0-9]+\s*/ )"
    return bool(re.search(marker_pattern, text) or re.search(slash_pattern, text))


def target_columns(worksheet):
    headers = {str(value or "").replace("\n", " ").strip().casefold(): index for index, value in enumerate(next(worksheet.iter_rows(max_row=1, values_only=True)))}
    return {
        "spanish": headers["detailed description (esp)"],
        "english": headers["detailed description (eng)"],
        "manufacturer": headers["manufacturer name (from flatfile)"],
        "url": headers["url"],
    }


def company_map(path: Path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook["Canonical_Taxonomy"]
    header = {str(value or "").strip().casefold(): index for index, value in enumerate(next(worksheet.iter_rows(max_row=1, values_only=True)))}
    company_column = header["company"]
    mapping = defaultdict(Counter)
    for row in worksheet.iter_rows(min_row=2, values_only=True):
        company = str(row[company_column] or "").strip()
        if not company or token_text(company) in PLACEHOLDER_LABELS:
            continue
        identifier = company_phrase(company)
        if identifier:
            mapping[identifier][company] += 1
    workbook.close()
    return mapping


def stable_company(counter: Counter):
    if len(counter) != 1:
        return None
    return next(iter(counter))


def audit(target: Path, canonical: Path, sample_limit: int) -> dict:
    mapping = company_map(canonical)
    workbook = load_workbook(target, read_only=True, data_only=False)
    worksheet = workbook["Tenders"] if "Tenders" in workbook.sheetnames else workbook.active
    columns = target_columns(worksheet)
    matches = Counter()
    samples = {}

    for row_number, row in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
        if row[columns["manufacturer"]] not in (None, ""):
            continue
        source_text = str(row[columns["spanish"]] or "")
        text = normalized_terms(source_text)
        candidates = []
        for term in mapping:
            if f" {term} " not in text:
                continue
            if not has_brand_context(source_text, term):
                continue
            company = stable_company(mapping.get(term, Counter()))
            if company:
                candidates.append((len(term.split()), term, company))
        if not candidates:
            continue
        candidates.sort(reverse=True)
        width, term, company = candidates[0]
        if any(candidate[2] != company and candidate[0] == width for candidate in candidates):
            continue
        matches[(term, company)] += 1
        key = f"{term}|{company}"
        if key not in samples and len(samples) < sample_limit:
            samples[key] = {
                "row": row_number,
                "identifier": term,
                "manufacturer": company,
                "url": row[columns["url"]],
                "description": str(row[columns["spanish"]] or "")[:500],
            }
    workbook.close()
    return {
        "unique_company_identifiers": len(mapping),
        "unmapped_manufacturer_rows": sum(matches.values()),
        "identifier_company_pairs": [(term, company, count) for (term, company), count in matches.most_common(100)],
        "samples": list(samples.values()),
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--target", type=Path, required=True)
    parser.add_argument("--canonical", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--sample-limit", type=int, default=200)
    args = parser.parse_args()
    result = audit(args.target, args.canonical, args.sample_limit)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, indent=2, ensure_ascii=False), encoding="utf-8")
    print(json.dumps({key: value for key, value in result.items() if key != "samples"}, indent=2, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
