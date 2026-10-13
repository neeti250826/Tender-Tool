"""Supervise the separately launched extraction and finish its audited outputs."""
import ctypes,hashlib,html,json,logging,os,subprocess,sys,time
from datetime import datetime
from pathlib import Path

BASE=Path(__file__).resolve().parents[1]
ROOT=BASE/'output_fy27q2_20261012'
PYTHON=Path(r'C:\Users\Neeti\.cache\codex-runtimes\codex-primary-runtime\dependencies\python\python.exe')
NODE=Path(r'C:\Users\Neeti\.cache\codex-runtimes\codex-primary-runtime\dependencies\node\bin\node.exe')
SCRAPER_PYTHON=Path(r'C:\Users\Neeti\AppData\Local\Programs\Python\Python312\python.exe')
manifest=json.loads((ROOT/'run_manifest.json').read_text(encoding='utf-8'))
launch=json.loads((ROOT/'extraction_launch.json').read_text(encoding='utf-8-sig'))
store=ROOT/'extraction/by_url/FY27Q2'
retained={}

def process_alive(pid):
    kernel=ctypes.WinDLL('kernel32',use_last_error=True)
    kernel.OpenProcess.argtypes=[ctypes.c_uint,ctypes.c_bool,ctypes.c_uint];kernel.OpenProcess.restype=ctypes.c_void_p
    kernel.WaitForSingleObject.argtypes=[ctypes.c_void_p,ctypes.c_uint];kernel.WaitForSingleObject.restype=ctypes.c_uint
    kernel.CloseHandle.argtypes=[ctypes.c_void_p]
    handle=kernel.OpenProcess(0x00100000,False,pid)
    if not handle:return False
    try:return kernel.WaitForSingleObject(handle,0)==258
    finally:kernel.CloseHandle(handle)

def update(stage,error=None):
    for path in store.glob('*.json'):
        stat=path.stat()
        if retained.get(str(path),{}).get('mtime')==stat.st_mtime_ns:continue
        try:
            p=json.loads(path.read_text(encoding='utf-8'))
            retained[str(path)]={'mtime':stat.st_mtime_ns,'url':p['url'],'ok':bool(p.get('ok')),'contracts':len(p.get('contracts') or []),'items':sum(len(c.get('dialog_items') or []) for c in p.get('contracts') or [])}
        except Exception:continue
    values=list(retained.values());success=sum(p['ok'] for p in values)
    state={'stage':stage,'updated_at':datetime.now().astimezone().isoformat(),'started_at':launch['started_at'],'selected_urls':manifest['selected_urls'],'processed_urls':len(values),'successful_urls':success,'failed_urls':len(values)-success,'item_rows_so_far':sum(p['items'] for p in values if p['ok']),'remaining_urls':max(0,manifest['selected_urls']-len(values)),'october_complete':False,'source_snapshot_latest':manifest['snapshot_latest_publication'],'fx_mxn_to_usd':manifest['fx_mxn_to_usd'],'extraction_pid':launch['pid'],'error':error,'final_workbook':str(ROOT/'excel_exports/CompraMX Tender FY27 Q2 Aug-Oct 2026 mapped.xlsx')}
    temporary=ROOT/'pipeline_status.tmp.json';temporary.write_text(json.dumps(state,indent=2),encoding='utf-8');temporary.replace(ROOT/'pipeline_status.json')
    links=[]
    for relative,label in [('inputs/input_aug_oct_2026.xlsx','Prepared URL workbook'),('run_manifest.json','Source and scope manifest'),('pipeline_status.json','Current status'),('website_navigation_sop.docx','Extraction guide'),('excel_exports/notice_extraction_audit.csv','Notice coverage audit'),('excel_exports/CompraMX Tender FY27 Q2 Aug-Oct 2026 mapped.xlsx','Result workbook')]:
        if (ROOT/relative).exists():links.append(f'<a href="{html.escape(relative)}">{label}</a>')
    available_links=' · '.join(links)
    body=f'''<!doctype html><html><head><meta charset="utf-8"><meta http-equiv="refresh" content="30"><title>Mexico Aug–Oct extraction</title><style>body{{font:16px Arial;max-width:1000px;margin:50px auto;color:#17324d}}table{{border-collapse:collapse;width:100%}}td,th{{padding:12px;text-align:left;border-bottom:1px solid #ccd7e0}}th{{background:#eef3f7}}p{{line-height:1.5}}code{{overflow-wrap:anywhere}}a{{color:#175b8d}}</style></head><body><h1>Mexico tender extraction: August–October 2026</h1><p>Current stage: <b>{html.escape(stage)}</b>. Updated {html.escape(state['updated_at'])}.</p><p>The supplied source snapshot ends on {html.escape(manifest['snapshot_latest_publication'])}. October is partial. All publication-window notices and statuses are included for later CANLAT relevance review.</p><table><tr><th>Selected notice URLs</th><td>{state['selected_urls']:,}</td></tr><tr><th>Processed</th><td>{state['processed_urls']:,}</td></tr><tr><th>Parsed successfully</th><td>{state['successful_urls']:,}</td></tr><tr><th>Failed so far</th><td>{state['failed_urls']:,}</td></tr><tr><th>Extracted item rows so far</th><td>{state['item_rows_so_far']:,}</td></tr><tr><th>Remaining notice URLs</th><td>{state['remaining_urls']:,}</td></tr><tr><th>Approved MXN-to-USD multiplier</th><td>{state['fx_mxn_to_usd']}</td></tr></table><p>Each notice is saved independently. Failed notices are retried once after the first pass; successful captures are retained. Notices with no published contract items remain visible in the extraction audit.</p><p>{html.escape(error or '')}</p><p><a href="inputs/input_aug_oct_2026.xlsx">Prepared URL workbook</a> · <a href="run_manifest.json">Source and scope manifest</a> · <a href="pipeline_status.json">Current status</a> · <a href="excel_exports/notice_extraction_audit.csv">Notice coverage audit</a> · <a href="excel_exports/CompraMX Tender FY27 Q2 Aug-Oct 2026 mapped.xlsx">Result workbook when available</a></p><p>Do not treat this run as complete until the status and final QC report confirm all URLs were attempted. Blank taxonomy fields require review.</p></body></html>'''
    old_links='<a href="inputs/input_aug_oct_2026.xlsx">Prepared URL workbook</a> · <a href="run_manifest.json">Source and scope manifest</a> · <a href="pipeline_status.json">Current status</a> · <a href="excel_exports/notice_extraction_audit.csv">Notice coverage audit</a> · <a href="excel_exports/CompraMX Tender FY27 Q2 Aug-Oct 2026 mapped.xlsx">Result workbook when available</a>'
    body=body.replace(old_links,available_links)
    (ROOT/'handoff.html').write_text(body,encoding='utf-8')
    return state

def run(command,stage):
    update(stage)
    log=ROOT/'logs'/f'{stage}.log'
    with log.open('a',encoding='utf-8') as output:
        child=subprocess.Popen([str(x) for x in command],cwd=BASE,stdout=output,stderr=subprocess.STDOUT,env=dict(os.environ,PYTHONUTF8='1'))
        while child.poll() is None:update(stage);time.sleep(30)
        if child.returncode:raise RuntimeError(f'{stage} exited {child.returncode}. See {log}')

def main():
    chunked='--chunked-finish' in sys.argv
    finish_only='--finish-only' in sys.argv or chunked
    while not finish_only and process_alive(launch['pid']):update('extracting');time.sleep(30)
    first=update('resuming_saved_workbook_model' if finish_only else 'first_pass_finished')
    if not finish_only and (first['failed_urls'] or first['remaining_urls']):
        args=launch['arguments'];retrylog=ROOT/'logs/retry_scrape.log'
        with retrylog.open('w',encoding='utf-8') as output:
            child=subprocess.Popen([str(SCRAPER_PYTHON),*args],cwd=BASE,stdout=output,stderr=subprocess.STDOUT,env=dict(os.environ,PYTHONUTF8='1',PLAYWRIGHT_BROWSERS_PATH=r'C:\Users\Neeti\PycharmProjects\Clearstate\eprocure_v2\.playwright-mcp'))
            while child.poll() is None:update('retrying_failed_notices');time.sleep(30)
    # The extraction environment supplies requests for the historical translator;
    # workbook authoring still uses the bundled Node artifact runtime.
    if not finish_only:
        run([SCRAPER_PYTHON,BASE/'tools/prepare_canlat_oct_outputs.py','--root',ROOT],'translating_and_mapping')
    if chunked:
        model=json.loads((ROOT/'excel_exports/workbook_model.json').read_text(encoding='utf-8'))
        row_count=len(model['rows']);del model
        import zipfile
        for offset in range(0,row_count,8000):
            candidate=ROOT/f'excel_exports/chunks/chunk-{offset}.xlsx'
            if candidate.exists():
                try:
                    with zipfile.ZipFile(candidate) as book:
                        if book.testzip() is None:continue
                except (OSError,zipfile.BadZipFile):pass
            run([NODE,'--max-old-space-size=4096',ROOT/'build_workbooks.mjs','chunk',str(offset),'8000'],f'building_workbook_chunk_{offset}')
        run([PYTHON,BASE/'tools/assemble_canlat_artifact_chunks.py'],'assembling_workbook')
    else:
        run([NODE,'--max-old-space-size=12288',ROOT/'build_workbooks.mjs','final'],'building_workbook')
    # Export the recalculated values to CSV from the finished workbook, so the
    # CSV includes the sizing calculations rather than blank formula columns.
    from openpyxl import load_workbook
    import csv
    final_path=ROOT/'excel_exports/CompraMX Tender FY27 Q2 Aug-Oct 2026 mapped.xlsx'
    book=load_workbook(final_path,read_only=True,data_only=True)
    with (ROOT/'excel_exports/mapped_tender_data.csv').open('w',encoding='utf-8-sig',newline='') as stream:
        writer=csv.writer(stream)
        writer.writerows(book['Tenders'].iter_rows(values_only=True))
    book.close()
    qc=json.loads((ROOT/'excel_exports/qc_findings.json').read_text(encoding='utf-8'))
    checks={'extraction_complete':qc['extraction_complete'],'all_urls_attempted':qc['all_urls_attempted'],'failed_notices':qc['failed_notices'],'unprocessed_notices':qc['unprocessed_notices'],'english_missing_rows':qc['english_missing_rows'],'row_count':qc['item_rows'],'snapshot_is_partial_october':True,'independent_final_review':'pending'}
    (ROOT/'release_checks.json').write_text(json.dumps(checks,indent=2),encoding='utf-8')
    outputs={}
    for p in (ROOT/'excel_exports').iterdir():
        if p.is_file() and p.suffix in {'.xlsx','.csv','.json'}:outputs[p.name]={'bytes':p.stat().st_size,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()}
    (ROOT/'deliverables_manifest.json').write_text(json.dumps(outputs,indent=2),encoding='utf-8')
    update('outputs_ready_for_review' if qc['extraction_complete'] else 'source_limited_outputs_ready_for_review')

if __name__=='__main__':
    try:main()
    except Exception as error:
        update('pipeline_error',str(error));logging.exception('Pipeline failed');raise
