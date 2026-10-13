"""Audit strict family-level matches from a CANLAT canonical taxonomy workbook."""

from __future__ import annotations

import argparse
import json
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook

from map_canlat_taxonomy import (
    FAMILY_GENERIC_IDENTIFIERS,
    LABEL_FIELD_NAMES,
    PLACEHOLDER_LABELS,
    STOP_WORDS,
    family_identifier_forms,
    manufacturer_identifier_forms,
    ngrams,
    token_text,
    tokenize,
)


def is_safe_family(value: object) -> bool:
    tokens = [token for token in token_text(value).split() if token not in STOP_WORDS]
    useful = [token for token in tokens if token not in FAMILY_GENERIC_IDENTIFIERS]
    return bool(useful) and (any(len(token) >= 5 for token in useful) or len(useful) >= 2)


def find_target_columns(worksheet):
    headers = {str(value or "").strip().casefold(): index for index, value in enumerate(next(worksheet.iter_rows(max_row=1, values_only=True)))}
    return {
        "spanish": headers["detailed description (esp)"],
        "english": headers["detailed description (eng)"],
        "url": headers["url"],
    }


def canonical_family_map(path: Path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook["Canonical_Taxonomy"]
    header = {str(value or "").strip().casefold(): index for index, value in enumerate(next(worksheet.iter_rows(max_row=1, values_only=True)))}
    fields = ("sub-operating unit", "device category", "company", "family name")
    missing = [field for field in fields if field not in header]
    if missing:
        raise ValueError(f"Canonical taxonomy is missing columns: {', '.join(missing)}")

    mapping = defaultdict(Counter)
    company_forms = {}
    for row in worksheet.iter_rows(min_row=2, values_only=True):
        labels = tuple(str(row[header[field]] or "").strip() for field in fields)
        if not all(labels) or any(token_text(value) in PLACEHOLDER_LABELS for value in labels):
            continue
        family = labels[3]
        if not is_safe_family(family):
            continue
        for identifier in family_identifier_forms(family):
            mapping[identifier][labels] += 1
        company_forms[labels] = set(manufacturer_identifier_forms(labels[2]))
    workbook.close()
    return mapping, company_forms


def audit(target: Path, canonical: Path, sample_limit: int, require_company: bool) -> dict:
    family_map, company_forms = canonical_family_map(canonical)
    target_book = load_workbook(target, read_only=True, data_only=False)
    worksheet = target_book["Tenders"] if "Tenders" in target_book.sheetnames else target_book.active
    columns = find_target_columns(worksheet)
    matches = Counter()
    samples_by_identifier = {}

    for row_number, row in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
        text = " ".join(str(row[columns[key]] or "") for key in ("spanish", "english"))
        terms = set(ngrams(tokenize(text), minimum=1, maximum=6))
        candidates = []
        for term in terms:
            labels = family_map.get(term)
            if not labels or len(labels) != 1:
                continue
            label_tuple = next(iter(labels))
            if require_company and not (company_forms[label_tuple] & terms):
                continue
            candidates.append((len(term.split()), term, label_tuple))
        if not candidates:
            continue
        candidates.sort(reverse=True)
        best_width, best_term, best_labels = candidates[0]
        if any(candidate[2] != best_labels and candidate[0] == best_width for candidate in candidates):
            continue
        matches[best_term] += 1
        if best_term not in samples_by_identifier and len(samples_by_identifier) < sample_limit:
            samples_by_identifier[best_term] = (
                {
                    "row": row_number,
                    "identifier": best_term,
                    "labels": dict(zip(LABEL_FIELD_NAMES, best_labels)),
                    "url": row[columns["url"]],
                    "description": str(row[columns["spanish"]] or "")[:500],
                }
            )
    target_book.close()
    return {
        "unique_safe_family_identifiers": len(family_map),
        "matched_rows": sum(matches.values()),
        "unique_matched_identifiers": len(matches),
        "top_identifiers": matches.most_common(50),
        "samples": list(samples_by_identifier.values()),
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--target", type=Path, required=True)
    parser.add_argument("--canonical", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--sample-limit", type=int, default=120)
    parser.add_argument("--require-company", action="store_true", help="Only accept rows containing both the canonical company and family identifiers.")
    args = parser.parse_args()

    result = audit(args.target, args.canonical, args.sample_limit, args.require_company)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, indent=2, ensure_ascii=False), encoding="utf-8")
    print(json.dumps({key: value for key, value in result.items() if key != "samples"}, indent=2, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
