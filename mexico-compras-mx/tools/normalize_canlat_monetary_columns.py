"""Convert CompraMX monetary text in N, P, and R to numeric Excel values."""

from __future__ import annotations

import argparse
import json
import os
import re
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook


TENDER_SHEET = "Tenders"
MONETARY_COLUMNS = {
    14: "Unit price without taxes (MXN)",
    16: "IVA",
    18: "Total",
}
MONEY_PATTERN = re.compile(
    r"^\s*\$?\s*(?P<amount>(?:[0-9]{1,3}(?:,[0-9]{3})*|[0-9]+)(?:\.[0-9]+)?)\s*$"
)
MXN_NUMBER_FORMAT = '[$MXN]\\ #,##0.00;[Red]\\-[$MXN]\\ #,##0.00'


def parse_amount(value: object, row: int, column_name: str) -> float:
    if isinstance(value, bool):
        raise ValueError(f"{column_name}{row} contains a boolean, not an amount")
    if isinstance(value, (int, float, Decimal)):
        return float(value)
    if not isinstance(value, str):
        raise ValueError(f"{column_name}{row} has unsupported value {value!r}")

    match = MONEY_PATTERN.match(value)
    if not match:
        raise ValueError(f"{column_name}{row} has an unparseable amount {value!r}")
    return float(Decimal(match.group("amount").replace(",", "")))


def normalize_workbook(source: Path, output: Path) -> dict[str, int]:
    if output.exists():
        raise FileExistsError(f"Refusing to overwrite existing output: {output}")

    workbook = load_workbook(source, read_only=False, data_only=False)
    try:
        if TENDER_SHEET not in workbook.sheetnames:
            raise ValueError(f"Source workbook does not contain {TENDER_SHEET!r}")
        sheet = workbook[TENDER_SHEET]
        converted = {column_name: 0 for column_name in MONETARY_COLUMNS.values()}

        for column, column_name in MONETARY_COLUMNS.items():
            if sheet.cell(1, column).value != column_name:
                raise ValueError(
                    f"Expected {column_name!r} in column {column}, found {sheet.cell(1, column).value!r}"
                )

            for row in range(2, sheet.max_row + 1):
                cell = sheet.cell(row, column)
                if cell.value in (None, ""):
                    continue
                cell.value = parse_amount(cell.value, row, cell.coordinate[0])
                cell.number_format = MXN_NUMBER_FORMAT if column == 14 else "General"
                converted[column_name] += 1

        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
        output.parent.mkdir(parents=True, exist_ok=True)
        temporary = output.with_suffix(".in_progress.xlsx")
        workbook.save(temporary)
        os.replace(temporary, output)
        return converted
    finally:
        workbook.close()


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()

    converted = normalize_workbook(args.source.resolve(), args.output.resolve())
    print(json.dumps({"output": str(args.output.resolve()), "converted": converted}, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
