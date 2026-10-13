"""Create a CANLAT-formatted raw workbook with taxonomy columns removed."""

from __future__ import annotations

import argparse
import os
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


TENDER_SHEET = "Tenders"
MAPPING_START_COLUMN = 29  # AC
MAPPING_COLUMN_COUNT = 4   # AC:AF
EXPECTED_MAPPING_HEADERS = (
    "Sub OU",
    "Product Category\n(From flatfile)",
    "Manufacturer Name\n(From FlatFile)",
    "Family Name\n(From Flatfile)",
)


def require_blank_mapping_columns(sheet) -> None:
    headers = tuple(
        sheet.cell(1, column).value
        for column in range(MAPPING_START_COLUMN, MAPPING_START_COLUMN + MAPPING_COLUMN_COUNT)
    )
    if headers != EXPECTED_MAPPING_HEADERS:
        raise ValueError(f"Unexpected taxonomy headers in AC:AF: {headers!r}")

    populated_rows = []
    for row in range(2, sheet.max_row + 1):
        if any(
            sheet.cell(row, column).value not in (None, "")
            for column in range(MAPPING_START_COLUMN, MAPPING_START_COLUMN + MAPPING_COLUMN_COUNT)
        ):
            populated_rows.append(row)
            if len(populated_rows) == 5:
                break
    if populated_rows:
        raise ValueError(
            "The source contains mapping values in AC:AF; use the prepared pre-mapping workbook. "
            f"Example rows: {populated_rows}"
        )


def set_shifted_formulas(sheet) -> None:
    """Restore formulas whose references move when AC:AF are omitted."""
    for row in range(2, sheet.max_row + 1):
        sheet.cell(row, 29).value = f'=IFERROR((AVERAGE(V{row},W{row})*0.05)/AB{row},0)'
        sheet.cell(row, 30).value = f'=ROUND(IFERROR(AVERAGE(T{row},U{row})/AB{row},0),0)'
        sheet.cell(row, 31).value = f'=ROUND(AC{row}*$AI{row},0)'
        sheet.cell(row, 32).value = f'=ROUND(AD{row}*$AI{row},0)'
        sheet.cell(row, 33).value = f'=IFERROR(N{row}*0.05,"")'
        sheet.cell(row, 35).value = (
            f'=IFERROR(MAX(0,DATEDIF(MAX(Z{row},DATE(2026,5,1)),'
            f'MIN(AA{row},DATE(2026,7,31)),"m")+1),0)'
        )


def build_workbook(source: Path, output: Path) -> None:
    if output.exists():
        raise FileExistsError(f"Refusing to overwrite existing output: {output}")

    workbook = load_workbook(source, read_only=False, data_only=False)
    try:
        if TENDER_SHEET not in workbook.sheetnames:
            raise ValueError(f"Source workbook does not contain the {TENDER_SHEET!r} sheet")
        sheet = workbook[TENDER_SHEET]
        require_blank_mapping_columns(sheet)

        sheet.delete_cols(MAPPING_START_COLUMN, MAPPING_COLUMN_COUNT)
        set_shifted_formulas(sheet)
        sheet.auto_filter.ref = f"A1:{get_column_letter(sheet.max_column)}{sheet.max_row}"
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True

        output.parent.mkdir(parents=True, exist_ok=True)
        temporary = output.with_suffix(".in_progress.xlsx")
        workbook.save(temporary)
        os.replace(temporary, output)
    finally:
        workbook.close()


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, required=True, help="Prepared CANLAT workbook before taxonomy mapping")
    parser.add_argument("--output", type=Path, required=True, help="New workbook without the four taxonomy columns")
    args = parser.parse_args()

    build_workbook(args.source.resolve(), args.output.resolve())
    print(f"Created {args.output.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
