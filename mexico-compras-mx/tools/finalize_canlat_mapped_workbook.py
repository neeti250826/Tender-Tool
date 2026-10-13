"""Apply conservative canonical enrichment and add a reviewer README sheet."""

from __future__ import annotations

import argparse
import json
import re
import shutil
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from map_canlat_taxonomy import (
    FAMILY_GENERIC_IDENTIFIERS,
    PLACEHOLDER_LABELS,
    STOP_WORDS,
    TOKEN_PATTERN,
    family_identifier_forms,
    normalized_header,
    token_text,
)


LABEL_HEADERS = (
    "Sub OU",
    "Product Category",
    "Manufacturer Name",
    "Family Name",
)
COMPANY_GENERIC_TERMS = {
    "ADVANCED", "BIO", "CONCEPT", "DEVICES", "DRIVE", "GENERAL", "GLOBAL", "HEALTH",
    "HEMOSTASIS", "KOREA", "MED", "MEDICAL", "MODERN", "RAPID", "ROBOTICS", "SPINE",
    "SURGICAL", "SYSTEMS", "TECHNOLOGIES", "TECHNOLOGY",
}
FAMILY_MATCH_RULES = {
    "CARDIOHELP": "GETINGE",
    "SYSTEM 5000": "CONMED",
    "LEGEND": "MEDTRONIC",
    "FLOW SENSOR": "HAMILTON",
}


def normalized_terms(value: object) -> str:
    return " " + " ".join(TOKEN_PATTERN.findall(token_text(value))) + " "


def company_phrase(value: object):
    tokens = TOKEN_PATTERN.findall(token_text(value))
    meaningful = [token for token in tokens if token not in COMPANY_GENERIC_TERMS]
    if not meaningful:
        return None
    if len(tokens) == 1 and (len(tokens[0]) < 5 or tokens[0] in COMPANY_GENERIC_TERMS):
        return None
    if len(tokens) > 1 and not any(len(token) >= 3 for token in meaningful):
        return None
    return " ".join(tokens)


def has_brand_context(value: object, phrase: str) -> bool:
    text = token_text(value).replace("&", " ")
    excluded_context = (
        "COMPATIBLE",
        "MANTENIMIENTO",
        "PARA SU USO EN EL EQUIPO",
        "SE SUGIERE",
        "SUGERIDA",
        "SERVICIO",
    )
    if any(term in text for term in excluded_context):
        return False
    phrase_pattern = re.escape(phrase).replace(r"\ ", r"\s+")
    marker_pattern = rf"\b(?:MARCA|MCA|BRAND|FABRICANTE|MANUFACTURER)\b(?:\s|:|/|-|\.){{0,10}}{phrase_pattern}\b"
    slash_pattern = rf"/\s*{phrase_pattern}(?=\s*/|\s*-|\s*[A-Z0-9]+\s*/ )"
    return bool(re.search(marker_pattern, text) or re.search(slash_pattern, text))


def stable_value(counter: Counter):
    if len(counter) != 1:
        return None
    return next(iter(counter))


def is_safe_family(value: object) -> bool:
    tokens = [token for token in token_text(value).split() if token not in STOP_WORDS]
    useful = [token for token in tokens if token not in FAMILY_GENERIC_IDENTIFIERS]
    return bool(useful) and (any(len(token) >= 5 for token in useful) or len(useful) >= 2)


def get_columns(worksheet):
    header = next(worksheet.iter_rows(max_row=1, values_only=True))
    columns = {}
    for index, value in enumerate(header, start=1):
        normalized = normalized_header(value)
        if normalized == "detailed description (esp)":
            columns["spanish"] = index
        elif normalized in {"detailed description (eng)", "english"}:
            columns["english"] = index
        elif normalized == "sub ou":
            columns["sub_ou"] = index
        elif normalized.startswith("product category"):
            columns["product_category"] = index
        elif normalized.startswith("manufacturer name"):
            columns["manufacturer"] = index
        elif normalized.startswith("family name"):
            columns["family"] = index
    required = {"spanish", "english", "sub_ou", "product_category", "manufacturer", "family"}
    missing = sorted(required - set(columns))
    if missing:
        raise ValueError(f"Tenders sheet is missing columns: {', '.join(missing)}")
    return columns


def canonical_evidence(path: Path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook["Canonical_Taxonomy"]
    header = {normalized_header(value): index for index, value in enumerate(next(worksheet.iter_rows(max_row=1, values_only=True)))}
    required = ("sub-operating unit", "device category", "company", "family name")
    missing = [field for field in required if field not in header]
    if missing:
        raise ValueError(f"Canonical taxonomy is missing columns: {', '.join(missing)}")

    company_map = defaultdict(Counter)
    family_map = defaultdict(Counter)
    for row in worksheet.iter_rows(min_row=2, values_only=True):
        labels = tuple(str(row[header[field]] or "").strip() for field in required)
        if not all(labels) or any(token_text(value) in PLACEHOLDER_LABELS for value in labels):
            continue
        company_identifier = company_phrase(labels[2])
        if company_identifier:
            company_map[company_identifier][labels[2]] += 1
        if is_safe_family(labels[3]):
            for identifier in family_identifier_forms(labels[3]):
                if identifier in FAMILY_MATCH_RULES:
                    family_map[identifier][labels] += 1
    workbook.close()
    return company_map, family_map


def canonical_manufacturer(source_text: object, company_map):
    terms = normalized_terms(source_text)
    candidates = []
    for identifier, companies in company_map.items():
        if f" {identifier} " not in terms or not has_brand_context(source_text, identifier):
            continue
        company = stable_value(companies)
        if company:
            candidates.append((len(identifier.split()), identifier, company))
    if not candidates:
        return None
    candidates.sort(reverse=True)
    width, _, company = candidates[0]
    if any(candidate[2] != company and candidate[0] == width for candidate in candidates):
        return None
    return company


def canonical_family_labels(source_text: object, family_map):
    terms = normalized_terms(source_text)
    candidates = []
    for identifier, rule_company in FAMILY_MATCH_RULES.items():
        if f" {identifier} " not in terms or f" {rule_company} " not in terms:
            continue
        labels = stable_value(family_map.get(identifier, Counter()))
        if labels:
            candidates.append((len(identifier.split()), identifier, labels))
    if not candidates:
        return None
    candidates.sort(reverse=True)
    width, _, labels = candidates[0]
    if any(candidate[2] != labels and candidate[0] == width for candidate in candidates):
        return None
    return labels


def mapping_counts(worksheet, columns):
    return {
        name: sum(
            1
            for row in range(2, worksheet.max_row + 1)
            if worksheet.cell(row, columns[field]).value not in (None, "")
        )
        for name, field in zip(LABEL_HEADERS, ("sub_ou", "product_category", "manufacturer", "family"))
    }


def apply_canonical_enrichment(worksheet, columns, company_map, family_map):
    counts = Counter()
    for row_number in range(2, worksheet.max_row + 1):
        source_text = worksheet.cell(row_number, columns["spanish"]).value or ""
        existing = [
            worksheet.cell(row_number, columns[field]).value
            for field in ("sub_ou", "product_category", "manufacturer", "family")
        ]

        family_labels = canonical_family_labels(source_text, family_map)
        if family_labels and all(current in (None, "", expected) for current, expected in zip(existing, family_labels)):
            for field, expected in zip(("sub_ou", "product_category", "manufacturer", "family"), family_labels):
                cell = worksheet.cell(row_number, columns[field])
                if cell.value in (None, ""):
                    cell.value = expected
                    counts[f"{field}:audited_canonical_family"] += 1
            existing = [
                worksheet.cell(row_number, columns[field]).value
                for field in ("sub_ou", "product_category", "manufacturer", "family")
            ]

        if existing[2] in (None, ""):
            manufacturer = canonical_manufacturer(source_text, company_map)
            if manufacturer:
                worksheet.cell(row_number, columns["manufacturer"]).value = manufacturer
                counts["manufacturer:explicit_canonical_brand"] += 1
    return dict(counts)


def add_readme(workbook, tender_rows: int):
    if "README" in workbook.sheetnames:
        del workbook["README"]
    worksheet = workbook.create_sheet("README", 0)
    worksheet.sheet_view.showGridLines = False
    worksheet.sheet_properties.tabColor = "1F4E78"
    worksheet.freeze_panes = "A8"
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.page_margins.left = 0.3
    worksheet.page_margins.right = 0.3
    worksheet.page_margins.top = 0.45
    worksheet.page_margins.bottom = 0.45

    for column, width in {"A": 29, "B": 18, "C": 18, "D": 15, "E": 24, "F": 42}.items():
        worksheet.column_dimensions[column].width = width

    navy = "1F4E78"
    blue = "D9EAF7"
    yellow = "FFF2CC"
    light_blue = "EAF3F8"
    light_gray = "F2F2F2"
    white = "FFFFFF"
    thin_blue = Side(style="thin", color="9CC2E5")
    thin_gray = Side(style="thin", color="D9E2F3")
    section_fill = PatternFill("solid", fgColor=navy)
    table_fill = PatternFill("solid", fgColor=blue)
    note_fill = PatternFill("solid", fgColor=yellow)
    pale_fill = PatternFill("solid", fgColor=light_blue)
    gray_fill = PatternFill("solid", fgColor=light_gray)
    white_font = Font(name="Tahoma", size=10, bold=True, color=white)
    section_font = Font(name="Tahoma", size=10, bold=True, color=white)
    body_font = Font(name="Tahoma", size=10, color="000000")
    body_bold = Font(name="Tahoma", size=10, bold=True, color="000000")
    title_font = Font(name="Tahoma", size=16, bold=True, color=white)
    subtitle_font = Font(name="Tahoma", size=10, italic=True, color=white)
    wrapped = Alignment(horizontal="left", vertical="top", wrap_text=True)
    centered = Alignment(horizontal="center", vertical="center", wrap_text=True)

    worksheet.merge_cells("A1:F1")
    worksheet["A1"] = "CANLAT Mexico Tender Extract - Reviewer Guide"
    worksheet["A1"].fill = section_fill
    worksheet["A1"].font = title_font
    worksheet["A1"].alignment = Alignment(horizontal="left", vertical="center")
    worksheet.row_dimensions[1].height = 28

    worksheet.merge_cells("A2:F2")
    worksheet["A2"] = "FY27 Q1 | Tender period: May-July 2026 | Mapping and review notes"
    worksheet["A2"].fill = section_fill
    worksheet["A2"].font = subtitle_font
    worksheet["A2"].alignment = Alignment(horizontal="left", vertical="center")
    worksheet.row_dimensions[2].height = 20

    worksheet.merge_cells("A4:F4")
    worksheet["A4"] = "Purpose and Scope"
    worksheet["A4"].fill = section_fill
    worksheet["A4"].font = section_font
    worksheet["A4"].alignment = wrapped
    worksheet.merge_cells("A5:F6")
    worksheet["A5"] = (
        "This workbook contains CompraMX tender records prepared for CANLAT FY27 Q1. "
        "Rows outside May-July 2026 are marked separately in the Irrelevant/Relevant column. "
        "The taxonomy fields in columns AC:AF are populated only where the available evidence supports the mapping."
    )
    worksheet["A5"].fill = note_fill
    worksheet["A5"].font = body_font
    worksheet["A5"].alignment = wrapped
    worksheet.row_dimensions[5].height = 43

    worksheet.merge_cells("A8:F8")
    worksheet["A8"] = "Mapping Coverage"
    worksheet["A8"].fill = section_fill
    worksheet["A8"].font = section_font
    worksheet["A8"].alignment = wrapped
    worksheet["A9"] = "Total tender rows"
    worksheet["B9"] = f"=COUNTA('Tenders'!$A$2:$A${tender_rows + 1})"
    worksheet["A9"].font = body_bold
    worksheet["B9"].font = body_bold
    worksheet["A9"].fill = pale_fill
    worksheet["B9"].fill = pale_fill
    worksheet["B9"].number_format = "#,##0"
    for cell in worksheet["A9:F9"][0]:
        cell.border = Border(bottom=thin_blue)
        cell.alignment = wrapped

    coverage_headers = ("Field", "Mapped Rows", "Unmapped Rows", "Coverage", "Evidence Used", "Reviewer Action")
    for column, value in enumerate(coverage_headers, start=1):
        cell = worksheet.cell(11, column, value)
        cell.fill = table_fill
        cell.font = body_bold
        cell.alignment = centered
        cell.border = Border(top=thin_blue, bottom=thin_blue)
    mapping_columns = ("AC", "AD", "AE", "AF")
    evidence_text = (
        "Historic exact description/CUCoP; audited canonical family paths",
        "Historic exact description/CUCoP; audited canonical family paths",
        "Historic exact descriptions; explicit canonical brand evidence",
        "Historic exact descriptions; audited company + family paths",
    )
    reviewer_text = (
        "Review blanks against the product description and CUCoP.",
        "Review blanks against the product description and CUCoP.",
        "Review blanks only when the manufacturer is explicitly stated.",
        "Review blanks only when a specific product family is explicitly stated.",
    )
    for row, (field, column, evidence, action) in enumerate(zip(LABEL_HEADERS, mapping_columns, evidence_text, reviewer_text), start=12):
        worksheet.cell(row, 1, field)
        worksheet.cell(row, 2, f"=COUNTIF('Tenders'!${column}$2:${column}${tender_rows + 1},\"<>\")")
        worksheet.cell(row, 3, f"=$B$9-B{row}")
        worksheet.cell(row, 4, f"=IFERROR(B{row}/$B$9,0)")
        worksheet.cell(row, 5, evidence)
        worksheet.cell(row, 6, action)
        for column_index in range(1, 7):
            cell = worksheet.cell(row, column_index)
            cell.alignment = wrapped
            cell.border = Border(bottom=thin_gray)
            cell.font = body_font
        worksheet.cell(row, 1).font = body_bold
        worksheet.cell(row, 2).number_format = "#,##0"
        worksheet.cell(row, 3).number_format = "#,##0"
        worksheet.cell(row, 4).number_format = "0.0%"
        worksheet.row_dimensions[row].height = 32

    worksheet.merge_cells("A18:F18")
    worksheet["A18"] = "How the Mapping Works"
    worksheet["A18"].fill = section_fill
    worksheet["A18"].font = section_font
    worksheet["A18"].alignment = wrapped
    mapping_rules = (
        ("1. Exact historical evidence", "Matches the Spanish or English tender description to reviewed CANLAT historical tender records."),
        ("2. CUCoP evidence", "Uses a stable CUCoP code only for Sub OU and Product Category, never to infer Manufacturer or Family."),
        ("3. Explicit brand evidence", "Adds a manufacturer only where the tender explicitly declares a validated brand in the product context."),
        ("4. Audited family evidence", "Maps the complete AC:AF path only for the reviewed company-and-family combinations Cardiohelp, System 5000, Legend, and Flow Sensor."),
        ("5. Blank remains reviewable", "Blank taxonomy cells mean there was not enough reliable evidence for an automated mapping; they do not mean the tender is out of scope."),
    )
    for row, (rule, explanation) in enumerate(mapping_rules, start=19):
        worksheet.cell(row, 1, rule)
        worksheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=6)
        worksheet.cell(row, 2, explanation)
        worksheet.cell(row, 1).fill = gray_fill
        worksheet.cell(row, 1).font = body_bold
        worksheet.cell(row, 2).font = body_font
        worksheet.cell(row, 1).alignment = wrapped
        worksheet.cell(row, 2).alignment = wrapped
        for column in range(1, 7):
            worksheet.cell(row, column).border = Border(bottom=thin_gray)
        worksheet.row_dimensions[row].height = 30

    worksheet.merge_cells("A26:F26")
    worksheet["A26"] = "Workbook Layout"
    worksheet["A26"].fill = section_fill
    worksheet["A26"].font = section_font
    worksheet["A26"].alignment = wrapped
    layout_rows = (
        ("Tenders", "The working tender table. Do not change source descriptions, dates, or the formulated columns unless you are correcting source data."),
        ("AC:AF", "Team mapping fields: Sub OU, Product Category, Manufacturer Name, and Family Name. These remain editable for reviewer decisions."),
        ("AG:AM", "Formula-driven value, volume, ASP, tender-quarter, and relevant-month fields. Preserve the formulas when reviewing the file."),
        ("Irrelevant/Relevant", "Contains the timeframe decision. It is separate from taxonomy mapping and should show Out of time frame where applicable."),
    )
    for row, (area, explanation) in enumerate(layout_rows, start=27):
        worksheet.cell(row, 1, area)
        worksheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=6)
        worksheet.cell(row, 2, explanation)
        worksheet.cell(row, 1).fill = pale_fill
        worksheet.cell(row, 1).font = body_bold
        worksheet.cell(row, 2).font = body_font
        worksheet.cell(row, 1).alignment = wrapped
        worksheet.cell(row, 2).alignment = wrapped
        for column in range(1, 7):
            worksheet.cell(row, column).border = Border(bottom=thin_gray)
        worksheet.row_dimensions[row].height = 31

    worksheet.merge_cells("A33:F33")
    worksheet["A33"] = "Reference Sources Used"
    worksheet["A33"].fill = section_fill
    worksheet["A33"].font = section_font
    worksheet["A33"].alignment = wrapped
    sources = (
        "CANLAT Mexico Tenders File Q4FY26 - reviewed historical tender mappings.",
        "CANLAT Mexico Tenders Q1FY26 - reviewed historical tender mappings.",
        "CANLAT Canonical Taxonomy Normalization v1 (26 Jul 2026) - governed company and family taxonomy used for the controlled enrichment.",
    )
    for row, source in enumerate(sources, start=34):
        worksheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
        cell = worksheet.cell(row, 1, source)
        cell.alignment = wrapped
        cell.font = body_font
        cell.fill = note_fill if row == 36 else PatternFill(fill_type=None)
        cell.border = Border(bottom=thin_gray)
        worksheet.row_dimensions[row].height = 25

    worksheet.merge_cells("A39:F39")
    worksheet["A39"] = "Reviewer Reminder"
    worksheet["A39"].fill = section_fill
    worksheet["A39"].font = section_font
    worksheet["A39"].alignment = wrapped
    worksheet.merge_cells("A40:F41")
    worksheet["A40"] = (
        "When completing a blank mapping, use the tender description, model, and stated manufacturer as evidence. "
        "Avoid assigning a Family Name from generic product terms alone."
    )
    worksheet["A40"].fill = note_fill
    worksheet["A40"].font = body_font
    worksheet["A40"].alignment = wrapped
    worksheet.row_dimensions[40].height = 35


def finalize(workbook_path: Path, canonical_path: Path, audit_path: Path):
    temporary = workbook_path.with_suffix(".in_progress.xlsx")
    shutil.copy2(workbook_path, temporary)
    workbook = load_workbook(temporary, data_only=False)
    worksheet = workbook["Tenders"]
    columns = get_columns(worksheet)
    before = mapping_counts(worksheet, columns)
    company_map, family_map = canonical_evidence(canonical_path)
    enrichment = apply_canonical_enrichment(worksheet, columns, company_map, family_map)
    after = mapping_counts(worksheet, columns)
    add_readme(workbook, worksheet.max_row - 1)
    try:
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
        workbook.calculation.calcMode = "auto"
    except AttributeError:
        pass
    workbook.save(temporary)
    workbook.close()
    temporary.replace(workbook_path)

    audit = {
        "workbook": str(workbook_path),
        "canonical_taxonomy": str(canonical_path),
        "rows": worksheet.max_row - 1,
        "coverage_before": before,
        "coverage_after": after,
        "enrichment_counts": enrichment,
        "family_match_rules": FAMILY_MATCH_RULES,
        "notes": [
            "Existing values were retained and never overwritten.",
            "Canonical manufacturer values required explicit product-brand context.",
            "Complete taxonomy paths required a reviewed company-and-family match.",
        ],
    }
    audit_path.parent.mkdir(parents=True, exist_ok=True)
    audit_path.write_text(json.dumps(audit, indent=2, ensure_ascii=False), encoding="utf-8")
    return audit


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--canonical", type=Path, required=True)
    parser.add_argument("--audit-json", type=Path, required=True)
    args = parser.parse_args()
    print(json.dumps(finalize(args.workbook, args.canonical, args.audit_json), indent=2, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
