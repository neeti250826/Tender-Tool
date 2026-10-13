"""
Refine Company / Family mapping on the TURMECA working sheet and add reliability columns.

For every data row (Device Category is always set on the data tabs):
  Company
    1. keep an existing taxonomy company;
    2. else search the published manufacturer, then the make/brand text: alias table -> exact -> fuzzy>=95 (vocab_rules)
       -> normalised name (legal/generic words dropped, spaces/hyphens removed: BIEN AIR = Bien-Air, NSK DENTAL = NSK);
    3. else the supplier name, only when it is the OEM / its affiliate (same matcher);
    4. else mark: "Not in taxonomy" (manufacturer published, not one of the taxonomy companies),
       "Undefined" (only a distributor or nothing published), "Not related" (manufacturer cell holds status text e.g. HOLD).
  Family (only for a taxonomy company that is listed for the row's Device Category)
    strict token match (all distinctive family words) or compact-name match in model / brand / description text;
    else "Undefined" (families exist, none named) or "Not related" (no family listed for that company + category).
  New columns
    Expected Company      - unconfirmed lead: a company listed for this category whose name, or whose family name,
                            appears in the line text (never from "compatible with X" wording).
    Expected Family Name  - best partial family match, the only family listed for the company+category, or the candidates.
    Confidence Level      - High / Medium / Low for the whole OU -> Family chain (definitions in README_note).
    Mapping Status        - controlled reason value.
    Mapping Note          - row-level explanation.
Input : the sheet export (xlsx). Output: refine/<tab>__company.tsv, __family.tsv, __new.tsv (sheet row order) + stats.
"""
import json, os, re, sys
from collections import Counter, defaultdict

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from vocab_rules import Vocab, fold, FAM_STOP  # noqa: E402

SRC = sys.argv[1]
OUT = os.path.join(HERE, "refine")
os.makedirs(OUT, exist_ok=True)
TABS = ["KSA - NUPCO", "Turkey - EKAP", "Pakistan - KP MCC", "Pakistan - KP DGHS", "Pakistan - Federal EPMS", "UAE - Dubai eSupply"]
NEW = ["Expected Company", "Expected Family Name", "Confidence Level", "Mapping Status", "Mapping Note",
       "Identified Manufacturer (incl. outside taxonomy)", "Identified Brand / Model", "Identification Source"]

# parent roll-ups that are certain and whose parent is a taxonomy company
EXTRA_ALIAS = [(r"welch ?allyn|hill-?rom", "Baxter"), (r"\bmeril\b", "Meril Life Sciences")]
GENERIC = set("""medical medikal medizintechnik dental ag gmbh co company ltd limited llc inc corp corporation corporate usa pvt private
sa sas sp srl spa bv nv oy ab kg factory factories products product technologies technology technlogies healthcare health care
surgical international intl group industries industry industrial manufacturing mfg systems devices instruments the and of
sanayi ticaret ve anonim sirketi sti tibbi malzeme urunleri ithalat ihracat sdn bhd pte pty plc gmbh&co""".split())
STATUS = re.compile(r"^\s*(hold|cancel+ed|n/?a|none|nil|-+|tba|tbd|withdrawn|not awarded|no award|\?)\s*$", re.I)
COMPAT = re.compile(r"uyumlu|compatible|for use with|ile kullan|kullanilabil|suitable for|works with|fits", re.I)


REMARK = re.compile(r"market survey|offered price|trade price|higher|lower|required|not (offered|available|quoted)|sample|rejected|"
                    r"qualif|disqualif|importer|non[- ]?responsive|as per|n/?a\b|nil\b|^\s*-+\s*$", re.I)
GENERIC_WORDS = set("""sample adapter ilizarov fixator external bladeless optical retention fixation humidity warm drill drills titan trauma plate plates screw screws nail nails cuff manometer pressure specimen retrieval bag bags sampling line lines
titanium cage cages ring rings retractor protector wound tray trays handpiece instrument instruments forceps scissors clamp clamps suture sutures absorbable nonabsorbable non polypropylene polyester silk nylon steel polydioxanone polyglactin catgut
chromic monofilament multifilament braided coated sterile drape drapes tube tubes tracheostomy cuffed uncuffed reusable disposable single
use adult pediatric paediatric neonatal infant temporary permanent arterial venous sensor sensors probe probes cable cables circuit circuits
mask masks filter filters sidestream mainstream blade blades drill drills burr burrs balloon balloons ptca pta stent stents catheter catheters
lead leads warm warming air blanket blankets electrode electrodes needle needles kit kits set sets system systems accessories accessory
module modules monitor monitors pump pumps cannula cannulae clip clips stapler staplers reload reloads trocar trocars mesh valve valves
graft grafts sheath sheaths wire wires guidewire guidewires introducer dilator spo2 etco2 co2 oxygen humidifier ventilator ventilators
circuit hme bis loc nirs standard advanced premium basic new mini micro plus pro ultra xl xs one two dual triple""".split())


def stem(w):
    if len(w) <= 4:
        return w
    if w.endswith("ies"):
        return w[:-3] + "y"
    if w.endswith("sses"):
        return w[:-2]
    if re.search(r"(s|x|z|ch|sh)es$", w):
        return w[:-2]
    for suf in ("ing", "ed"):
        if w.endswith(suf) and len(w) - len(suf) >= 4:
            return w[: -len(suf)]
    if w.endswith("s") and not w.endswith("ss"):
        return w[:-1]
    return w


COMMON = set("""double rapid formed matrix insitu natural cutting straight profile laser anchor controller hemoclip cobalt chrome titan smart comfort select classic flow guide plus ultra micro mini single dual multi flex optima prime safe secure simple soft super total vital elite ideal unique apex alpha delta omega nova pro max tech care life line link vision view scope sense star true wave phoenix eagle atlas orion hydro aqua ever fine first focus form free gold green blue white black silver crystal diamond quick fast easy clear clean lock loop grip seal shield guard protect sterile surgical medical excel infinity integra primed creative general standard national international united global universal advanced premium medline medica
hold cancelled others other generic local local companies local company unknown various""".split())


def clean_name(n):
    n = re.sub(r"\s+", " ", str(n)).strip(" .,-")
    return n.title() if n.isupper() else n


def core(name):
    toks = [t for t in re.split(r"[^a-z0-9]+", fold(name)) if t and t not in GENERIC]
    return "".join(toks)


def compact(s):
    return re.sub(r"[^a-z0-9]+", "", fold(s))


class Refiner:
    def __init__(self):
        self.v = Vocab()
        rows = self.v.rows
        self.listed = defaultdict(set)          # category -> companies listed for it
        self.fams = defaultdict(list)           # (category, company) -> [family]
        for r in rows:
            if r[4]:
                self.listed[r[3]].add(r[4])
            if r[4] and r[5] and r[5] not in self.fams[(r[3], r[4])]:
                self.fams[(r[3], r[4])].append(r[5])
        self.core_idx = defaultdict(set)
        for c in self.v.companies:
            k = core(c)
            if len(k) >= 3:
                self.core_idx[k].add(c)
        self.extra = [(re.compile(x), c) for x, c in EXTRA_ALIAS]
        gen = {stem(w) for w in GENERIC_WORDS}
        tok_cos = defaultdict(set)
        for r in rows:
            if r[4] and r[5]:
                for t in re.split(r"[^a-z0-9]+", fold(r[5])):
                    if len(t) >= 3 and not t.isdigit():
                        tok_cos[stem(t)].add(r[4])
        gen |= {t for t, cs in tok_cos.items() if len(cs) >= 5}
        for r in rows:
            for x in (r[0], r[2], r[3]):
                gen |= {stem(t) for t in re.split(r"[^a-z0-9]+", fold(x)) if t}
        self.gen = gen
        self.co_toks = {c: {stem(t) for t in re.split(r"[^a-z0-9]+", fold(c)) if t} for c in self.v.companies}
        # family name -> (category, company) for expected-company search
        self.fam_by_cat = defaultdict(list)
        for (cat, co), fl in self.fams.items():
            for f in fl:
                dt = self.dtoks(f, co)
                if dt and len("".join(dt)) >= 5 and all(len(x) >= 4 for x in dt):
                    self.fam_by_cat[cat].append((f, co, dt))

    def dtoks(self, family, co):
        """distinctive tokens of a family name: not generic product words, not the company's own name, not 'others'."""
        if re.search(r"\bothers?\b|\bgeneric\b", family, re.I):
            return []
        return [t for t in re.split(r"[^a-z0-9]+", fold(family))
                if len(t) >= 3 and not t.isdigit() and stem(t) not in self.gen and stem(t) not in self.co_toks.get(co, ()) and t not in FAM_STOP]

    def has(self, toks, t, ct):
        wb = lambda x: re.search(r"(?<![a-z0-9])" + re.escape(x) + r"s?(?![a-z0-9])", t)
        if all(wb(x) for x in toks):
            return True
        return len(toks) > 1 and bool(wb("".join(toks)))

    def company(self, text):
        """-> (company, method) or (None, None)"""
        if not text or not str(text).strip():
            return None, None
        m = self.v.company(text)
        if m:
            return m["company"], m["method"]
        t = fold(text)
        for rx, c in self.extra:
            if rx.search(t):
                return c, "parent alias"
        k = core(text)
        if len(k) >= 3 and len(self.core_idx.get(k, ())) == 1:
            return next(iter(self.core_idx[k])), "normalised name"
        return None, None

    def family(self, cat, co, text):
        """-> (family, how, expected, candidates)"""
        cands = self.fams.get((cat, co)) or []
        if not cands:
            return None, None, None, []
        t, ct = fold(text), compact(text)
        best, score, hits = None, 0.0, []
        for name in cands:
            dt = self.dtoks(name, co)
            if not dt:
                continue
            if self.has(dt, t, ct):
                hits.append((len("".join(dt)), name))
                continue
            got = [x for x in dt if self.has([x], t, ct)]
            if got and len(got) / len(dt) > score and max(len(x) for x in got) >= 4:
                best, score = name, len(got) / len(dt)
        if hits:
            return max(hits)[1], "distinctive name in text", None, cands
        return None, None, (best if score >= 0.5 else None), cands

    # ---- beyond the taxonomy: any manufacturer / brand evidenced in the data itself
    def build_dictionary(self, names):
        """names = manufacturer names published anywhere in the data (NUPCO Manufacturer, KP make) + taxonomy companies."""
        self.all_fams = {f for fl in self.fams.values() for f in fl}
        d = {}
        for n in list(names) + list(self.v.companies):
            n = str(n or "").strip()
            if not n or STATUS.match(n) or REMARK.search(n) or len(n) > 60:
                continue
            words = [t for t in re.split(r"[^a-z0-9]+", fold(n)) if t and t not in GENERIC]
            key = " ".join(words)
            if len(key.replace(" ", "")) < 5 or key in COMMON:
                continue
            d.setdefault(key, self.v.company(n)["company"] if self.v.company(n) else clean_name(n))
        self.mdict = d
        keys = sorted(d, key=len, reverse=True)
        self.mrx = re.compile(r"(?<![a-z0-9])(" + "|".join(re.escape(k) for k in keys) + r")(?![a-z0-9])")
        # taxonomy brands usable as text evidence: distinctive family tokens unique to one company
        self.brand_rx = []
        for (cat, co), fl in self.fams.items():
            for f in fl:
                dt = self.dtoks(f, co)
                if dt and all(len(x) >= 5 for x in dt) and not all(x in COMMON for x in dt):
                    self.brand_rx.append((re.compile(r"(?<![a-z0-9])" + r"\W*".join(re.escape(x) for x in dt) + r"(?![a-z0-9])"), co, f, cat))

    def identify(self, man, brand, sup, desc, cat=""):
        """-> [identified manufacturer, identified brand/model, source] using only evidence in the row."""
        model = brand if brand and not REMARK.search(brand) else ""
        if man and not STATUS.match(man):
            c = self.v.company(man)
            return [c["company"] if c else clean_name(man), model, "Published manufacturer" + ("" if c else " (not a taxonomy company)")]
        if model and not man:
            c = self.v.company(model)
            if c:
                return [c["company"], model, "Published make/brand text"]
            if re.search(r"[a-z]{3}", model, re.I) and not re.fullmatch(r"[\w./-]*\d[\w./-]*", model):
                return [clean_name(model), "", "Published make/brand text (not a taxonomy company)"]
        t = fold(desc)
        if not COMPAT.search(t):
            host = re.search(r"sensor|probe|prob[u]?\b|cable|kablo|circuit|devre|module|modul|filter|filtre|blade|bicag|accessor|aksesuar", t)
            hint = " (may name the host equipment, not the maker)" if host else ""
            m = self.mrx.search(t)
            for rx, co, f, bcat in self.brand_rx:
                if bcat == cat and rx.search(t):        # a brand counts only in lines of its own category
                    if m and self.mdict[m.group(1)] != co:   # e.g. 'WEGO PROLEN': the maker is named, the brand word is generic
                        return [self.mdict[m.group(1)], f, "Manufacturer named in item text - unconfirmed; brand word used generically"]
                    generic = " - brand word often used generically, verify" if re.search(r"prolen|vicr[yi]l|surgi[cs]el|foley|ligasure", t) else ""
                    return [co, f, "Brand named in item text - unconfirmed" + generic + hint]
            if m:
                return [self.mdict[m.group(1)], model, "Manufacturer named in item text - unconfirmed" + hint]
        if sup:
            return ["", model, "Not identified - only the distributor is published"]
        return ["", model, "Not identified - no supplier or manufacturer published"]

    def expected_company(self, cat, text):
        if not text or COMPAT.search(fold(text)):
            return None, None
        t, ct = fold(text), compact(text)
        for f, co, dt in self.fam_by_cat.get(cat, []):
            if self.has(dt, t, ct):
                return co, f
        for co in self.listed.get(cat, ()):
            k = fold(co).strip()
            if len(k) >= 5 and re.search(r"(?<![a-z0-9])" + re.escape(k) + r"(?![a-z0-9])", t):
                return co, None
        return None, None


def main():
    R = Refiner()
    wb = openpyxl.load_workbook(SRC, read_only=True)
    names = set()
    for tab in TABS:
        rows = list(wb[tab].iter_rows(values_only=True))
        h = {c: i for i, c in enumerate(rows[0]) if c}
        for r in rows[1:]:
            if r and str(r[0] or "")[:3] in ("SA-", "TR-", "PK-", "AE-", "KW-"):
                names.add(r[h["Manufacturer"]])
                if tab.startswith("Pakistan"):
                    names.add(r[h["Disclosed Brand"]])
    R.build_dictionary(names)
    print("manufacturer dictionary:", len(R.mdict), "names;", len(R.brand_rx), "taxonomy brands")
    stats = {}
    for tab in TABS:
        rows = list(wb[tab].iter_rows(values_only=True))
        h = list(rows[0])
        ix = {c: i for i, c in enumerate(h) if c}
        data = [r for r in rows[1:] if r and str(r[0] or "")[:3] in ("SA-", "TR-", "PK-", "AE-", "KW-")]   # note rows excluded
        g = lambda r, c: (r[ix[c]] if c in ix and r[ix[c]] is not None else "")
        st = Counter()
        out_co, out_fa, out_new = [], [], []
        for r in data:
            cat, co0, fa0 = g(r, "Device Category"), g(r, "Company"), g(r, "Family Name")
            if fa0 not in R.all_fams:              # markers written by an earlier run are not families
                fa0 = ""
            if co0 not in R.v.companies:
                co0 = ""
            man, brand, sup = str(g(r, "Manufacturer")).strip(), str(g(r, "Disclosed Brand")).strip(), str(g(r, "Supplier/Distributor")).strip()
            desc = " ".join(str(g(r, c)) for c in ("English Description", "Original-language Description"))
            mev = g(r, "Manufacturer Evidence")
            text = " ".join([brand, man, desc])
            note, exp_co, exp_fa = [], "", ""
            ident = R.identify(man, brand, sup, desc, cat)
            if not cat:                            # excluded in the category audit: not a taxonomy device
                status = "Excluded - not a taxonomy device"
                st[status] += 1
                out_co.append("Not related")
                out_fa.append("Not related")
                out_new.append(["", "", "Not applicable", status, str(g(r, "Exclusion Reason")), *ident])
                continue
            # ---------------- company
            co, how = (co0, "existing") if co0 and co0 in R.v.companies else (None, None)
            src = "manufacturer"
            if not co and man and not STATUS.match(man):
                co, how = R.company(man)
            if brand and REMARK.search(brand):
                brand = ""
            if not co and brand and not man and not STATUS.match(brand):
                co, how = R.company(brand)
                src = "make/brand"
            if not co and sup and not man:
                co, how = R.company(sup)
                src = "supplier (OEM/affiliate)"
            explicit = bool(co) and (how in ("existing", "alias", "exact") and mev == "Explicit" or (how in ("alias", "exact", "parent alias") and src != "supplier (OEM/affiliate)"))
            if co:
                listed = co in R.listed.get(cat, ())
            if co and how != "existing":
                note.append(f"Company found by {how} on {src}")
            # ---------------- family / status
            fam = ""
            if co and not listed:
                status, conf = "Company not listed for this category", "Low"
                fam = "Not related"
                note.append(f"{co} is in the taxonomy but not listed for '{cat}' - check or propose new combination")
            elif co:
                f, fhow, fexp, cands = (fa0, "existing", None, R.fams.get((cat, co)) or []) if fa0 else R.family(cat, co, text)
                if f:
                    fam = f
                    status = "Mapped to family"
                    conf = "High" if (explicit and how in ("existing", "alias", "exact")) else "Medium"
                    if fhow == "distinctive name in text":
                        note.append("Family name found in model/description text")
                elif cands:
                    fam = "Undefined"
                    status, conf = "Mapped to company - family not identified", "Medium"
                    if fexp:
                        exp_fa = fexp
                        note.append("Expected family = partial name match in text")
                    elif [c for c in cands if re.search(r"\bothers?\b", c, re.I)] and len(cands) <= 1:
                        exp_fa = cands[0]
                        note.append("Only the taxonomy catch-all family exists for this company and category")
                    elif len(cands) == 1:
                        exp_fa = cands[0]
                        note.append("Only family listed for this company and category")
                    elif [c for c in cands if re.search(r"\bothers?\b", c, re.I)]:
                        exp_fa = [c for c in cands if re.search(r"\bothers?\b", c, re.I)][0]
                        note.append(f"No listed family named in text; taxonomy catch-all family suggested ({len(cands)} families listed)")
                    else:
                        exp_fa = f"Undefined - {len(cands)} candidates: " + "; ".join(cands[:5]) + ("; ..." if len(cands) > 5 else "")
                        note.append("Model/description names none of the listed families")
                else:
                    fam = "Not related"
                    status, conf = "Mapped to company - family not identified", "Medium"
                    note.append(f"No family listed in taxonomy for {co} / {cat}")
                if how not in ("existing", "alias", "exact") and conf == "High":
                    conf = "Medium"
            else:
                if man and STATUS.match(man):
                    co_mark, status = "Not related", "Not related - status text in manufacturer field"
                    note.append(f"Manufacturer cell reads '{man}' (award status, not a company)")
                elif man or brand:
                    co_mark, status = "Not in taxonomy", "Manufacturer not in taxonomy"
                    note.append(f"Published manufacturer/make '{man or brand}' is not one of the {len(R.v.companies)} taxonomy companies")
                elif sup:
                    co_mark, status = "Undefined", "Distributor only - manufacturer not published"
                    note.append(f"Only the distributor '{sup[:60]}' is published; it is not an OEM or OEM affiliate")
                else:
                    co_mark, status = "Undefined", "No manufacturer or supplier published"
                    note.append("Source publishes no winner / manufacturer (requested quantity only)" if tab.startswith("UAE") else "No manufacturer or supplier published")
                co, fam, conf = co_mark, "Undefined", "Low"
                ec, ef = R.expected_company(cat, desc + " " + brand) if co_mark == "Undefined" else (None, None)
                if ec:
                    exp_co, exp_fa = ec, ef or ""
                    note.append("Expected company from a brand/family name in the line text - unconfirmed")
                    if re.search(r"prolen|vicr[yi]l|surgicel|ligasure|hemovac|foley", fold(desc + " " + brand)):
                        note.append("brand word is often used generically in tenders (e.g. PROLEN = any polypropylene suture) - verify")
            st[status] += 1
            st["conf " + conf] += 1
            if exp_co:
                st["expected company"] += 1
            if co != co0 and co0 == "" and co not in ("Undefined", "Not in taxonomy", "Not related"):
                st["company newly filled"] += 1
            if fam != fa0 and fa0 == "" and fam not in ("Undefined", "Not related"):
                st["family newly filled"] += 1
            if co in R.v.companies:                # the identification columns repeat a taxonomy company when there is one
                fam_ok = fam if fam not in ("Undefined", "Not related") else ""
                model = brand if brand and not REMARK.search(brand) else ""
                if fam_ok:
                    ident = [co, fam_ok, "Taxonomy company + taxonomy family"]
                elif model:
                    ident = [co, model, "Taxonomy company + published model/catalogue (family not in taxonomy text)"]
                else:
                    ident = [co, "", "Taxonomy company; no model or family published"]
            out_co.append(co)
            out_fa.append(fam)
            out_new.append([exp_co, exp_fa, conf, status, "; ".join(note), *ident])
        name = re.sub(r"[^\w-]+", "_", tab)
        w = lambda f, hdr, rows_: open(os.path.join(OUT, f"{name}__{f}.tsv"), "w", encoding="utf-8", newline="").write(
            "\t".join(hdr) + "\n" + "".join("\t".join(str(x).replace("\t", " ").replace("\n", " ") for x in rr) + "\n" for rr in rows_))
        w("company", ["Company"], [[x] for x in out_co])
        w("family", ["Family Name"], [[x] for x in out_fa])
        w("new", NEW, out_new)
        first_new = ix["Expected Company"] + 1 if "Expected Company" in ix else max(i for c, i in ix.items()) + 2
        stats[tab] = {"rows": len(data), "company_col": ix["Company"] + 1, "family_col": ix["Family Name"] + 1,
                      "ou_col": ix["OU"] + 1, "incl_col": ix["Included/Review/Excluded"] + 1, "excl_col": ix["Exclusion Reason"] + 1,
                      "last_col": first_new - 1, **st}
        print(tab, json.dumps(stats[tab], ensure_ascii=False))
    json.dump(stats, open(os.path.join(OUT, "stats.json"), "w"), indent=1)


if __name__ == "__main__":
    main()
