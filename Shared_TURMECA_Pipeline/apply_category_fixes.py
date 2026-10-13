"""
Apply category_fixes.py to a TURMECA sheet export.
  SET     -> OU / Sub-OU / Device Category replaced with the corrected taxonomy triple
  EXCLUDE -> OU / Sub-OU / Device Category blanked, Included/Review/Excluded = Excluded, Exclusion Reason filled
Writes <out>.xlsx (same layout as the export) and refine/category_changes.csv (row-level log).
Usage: python apply_category_fixes.py <export.xlsx> <out.xlsx>
"""
import csv, os, sys
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from vocab_rules import Vocab  # noqa: E402
import category_fixes as cf  # noqa: E402

TABS = ["KSA - NUPCO", "Turkey - EKAP", "Pakistan - KP MCC", "Pakistan - KP DGHS", "Pakistan - Federal EPMS", "UAE - Dubai eSupply"]
SUFFIX = {"ICT": "Interventional Cardiology Therapies (ICT)", "PVH": "Peripheral Vascular Health"}


def main(src, out):
    v = Vocab()
    triples = {}
    for r in v.rows:
        triples.setdefault(r[3], set()).add((r[0], r[2], r[3]))

    def resolve(target):
        name, _, suf = target.partition("|")
        cands = sorted(triples.get(name, ()))
        if suf:
            cands = [c for c in cands if c[0] == SUFFIX[suf]]
        if len(cands) != 1:
            raise ValueError(f"ambiguous or unknown target {target}: {cands}")
        return cands[0]

    wb = openpyxl.load_workbook(src)
    log = []
    for tab in TABS:
        ws = wb[tab]
        hdr = [c.value for c in ws[1]]
        ix = {h: i + 1 for i, h in enumerate(hdr) if h}
        for ri in range(2, ws.max_row + 1):
            rid = ws.cell(ri, 1).value
            if not rid or not str(rid)[:3] in ("SA-", "TR-", "PK-", "AE-", "KW-"):
                continue
            g = lambda c: ws.cell(ri, ix[c]).value
            cat = g("Device Category")
            f = cf.fix(rid, cat, g("Original-language Description"), g("English Description"), v)
            if not f:
                continue
            old = (g("OU"), g("Sub-OU"), cat)
            if f[0] == "SET":
                new = resolve(f[1])
                if new == old:
                    continue
                for c, val in zip(("OU", "Sub-OU", "Device Category"), new):
                    ws.cell(ri, ix[c], val)
                note = f[2]
            else:
                new = ("", "", "")
                for c in ("OU", "Sub-OU", "Device Category"):
                    ws.cell(ri, ix[c]).value = None      # ws.cell(r, c, None) would leave the old value in place
                ws.cell(ri, ix["Included/Review/Excluded"], "Excluded")
                ws.cell(ri, ix["Exclusion Reason"], f[2])
                note = f[2]
            log.append([tab, rid, f[0], *old, *new, note, (g("English Description") or g("Original-language Description") or "")[:200]])
    wb.save(out)
    with open(os.path.join(HERE, "refine", "category_changes.csv"), "w", encoding="utf-8-sig", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["Tab", "Row ID", "Action", "Old OU", "Old Sub-OU", "Old Device Category", "New OU", "New Sub-OU", "New Device Category", "Reason", "Description"])
        w.writerows(log)
    from collections import Counter
    print(Counter((l[0], l[2]) for l in log))


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
