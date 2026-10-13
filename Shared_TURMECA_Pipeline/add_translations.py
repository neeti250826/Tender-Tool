"""Post-process TURMECA_Master_Mapped_Data.xlsx: for every text column holding Turkish/Arabic text, insert an
'<column> (EN)' column right after it with the English translation (translate.googleapis.com, cached).
Blank 'English Description' cells are filled from 'Original-language Description'. Run after build_master_mapped.py
(and before or after add_gap_notes.py - notes are below the data and are not touched)."""
import os, re, sys
from collections import Counter

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from translate import translate_many  # noqa: E402

F = os.path.join(HERE, "TURMECA_Master_Mapped_Data.xlsx")
TR = re.compile(r"[ğĞışŞİçÇöÖüÜ]")
AR = re.compile(r"[؀-ۿ]")
SKIP = {"Original-language Description", "English Description", "Row ID", "Source URL", "Data Source", "Caveat"}
EN_FILL = PatternFill("solid", fgColor="2E75B6")


# Turkish public-procurement glossary applied BEFORE machine translation (fixes e.g. KALEM = item, not 'pen')
GLOSS = [(r"(\d+\)?)\s*KALEM(L[İI]K)?\b", r"\1 ITEM"), (r"\bKALEM\b", "ITEM"), (r"\bKISIM\b", "LOT"),
         (r"\bKISIMLI\b", "MULTI-LOT")]
CO = [("LİMİTED ŞİRKETİ", "LIMITED COMPANY"), ("ANONİM ŞİRKETİ", "JOINT STOCK COMPANY"), ("TİCARET", "TRADE"), ("SANAYİ", "INDUSTRY"),
      ("TIBBİ MALZEME", "MEDICAL SUPPLIES"), ("TIBBİ CİHAZ", "MEDICAL DEVICES"), ("MEDİKAL", "MEDICAL"), ("SAĞLIK", "HEALTH"), ("İTHALAT", "IMPORT"),
      ("İHRACAT", "EXPORT"), ("PAZARLAMA", "MARKETING"), ("HİZMETLERİ", "SERVICES"), ("ÜRÜNLERİ", "PRODUCTS"), ("VE", "AND"), ("LTD. ŞTİ.", "LTD. CO."),
      ("LTD.ŞTİ.", "LTD. CO."), ("A.Ş.", "INC.")]
_FOLD = str.maketrans("ĞğIıŞşİÇçÖöÜü", "GgIiSsICcOoUu")


def pre(t, l):
    if l != "tr":
        return t
    for a, b in GLOSS:
        t = re.sub(a, b, t, flags=re.I)
    return t


def fallback(t):
    for a, b in CO:
        t = re.sub(r"(?<!\w)" + re.escape(a) + r"(?!\w)", b, t)
    return t.translate(_FOLD)


def translate_fixed(texts, l):
    """Glossary -> batch translate -> retry singles that came back unchanged -> glossary/ASCII fallback."""
    src = [pre(t, l) for t in texts]
    out = translate_many(src, src=l, verbose=False)
    for i, (s0, o) in enumerate(zip(src, out)):
        # proper names (Ümraniye, İNKA) legitimately keep Turkish letters: only fix results returned completely unchanged
        if not o or o.strip().upper() == s0.strip().upper():
            # bypass the cached (failed) result with a slightly different request, lower-cased (helps ALL-CAPS Turkish)
            o2 = translate_many([s0.lower() + " ."], src=l, verbose=False)[0].rstrip(" .")
            out[i] = o2.upper() if (o2 and o2.strip().upper() != s0.strip().upper()) else fallback(s0)
    return out


def lang(vals):
    t = sum(1 for v in vals if TR.search(v))
    a = sum(1 for v in vals if AR.search(v))
    n = max(1, len(vals))
    if a / n > 0.02:
        return "ar"
    if t / n > 0.02:
        return "tr"
    return None


def main():
    wb = openpyxl.load_workbook(F)
    report = []
    for ws in wb.worksheets:
        if ws.title in ("README", "Summary") or ws.cell(1, 1).value != "Row ID":
            continue
        hdr = [c.value for c in ws[1]]
        # data rows end at the first row whose Row ID is blank (gap notes sit below)
        last = 1
        for r in range(2, ws.max_row + 1):
            if ws.cell(r, 1).value:
                last = r
            else:
                break
        if last < 2:
            continue
        # 1) fill blank English Description from the original-language text
        if "English Description" in hdr and "Original-language Description" in hdr:
            ci_en, ci_or = hdr.index("English Description") + 1, hdr.index("Original-language Description") + 1
            blanks = [r for r in range(2, last + 1) if not ws.cell(r, ci_en).value and ws.cell(r, ci_or).value]
            if blanks:
                src = [str(ws.cell(r, ci_or).value) for r in blanks]
                l = lang(src) or "auto"
                for r, e in zip(blanks, translate_fixed(src, l if l != "auto" else "auto")):
                    ws.cell(r, ci_en, e)
                report.append(f"{ws.title}: filled {len(blanks)} blank English Description")
        # 2) add '(EN)' columns next to Turkish/Arabic text columns (right-to-left so indices stay valid)
        targets = []
        for j, h in enumerate(hdr, 1):
            if not h or h in SKIP or str(h).endswith("(EN)"):
                continue
            vals = [str(ws.cell(r, j).value) for r in range(2, last + 1) if isinstance(ws.cell(r, j).value, str)]
            l = lang(vals)
            if l:
                targets.append((j, h, l))
        for j, h, l in sorted(targets, reverse=True):
            vals = [ws.cell(r, j).value for r in range(2, last + 1)]
            texts = [str(v) for v in vals if isinstance(v, str) and (TR.search(v) or AR.search(v))]
            uniq = sorted(set(texts))
            tmap = dict(zip(uniq, translate_fixed(uniq, l))) if uniq else {}
            if j < len(hdr) and hdr[j] == f"{h} (EN)":        # refresh an existing EN column
                for r, v in zip(range(2, last + 1), vals):
                    if isinstance(v, str):
                        ws.cell(r, j + 1, tmap.get(v, v))
                report.append(f"{ws.title}: refreshed '{h} (EN)' - {len(uniq)} values")
                continue
            ws.insert_cols(j + 1)
            c = ws.cell(1, j + 1, f"{h} (EN)")
            c.font, c.fill, c.alignment = Font(bold=True, color="FFFFFF"), EN_FILL, Alignment(wrap_text=True, vertical="center")
            ws.column_dimensions[get_column_letter(j + 1)].width = max(18, ws.column_dimensions[get_column_letter(j)].width or 18)
            for r, v in zip(range(2, last + 1), vals):
                if isinstance(v, str):
                    ws.cell(r, j + 1, tmap.get(v, v))
            report.append(f"{ws.title}: added '{h} (EN)' [{l}] - {len(uniq)} distinct values translated")
        ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{last}"
    wb.save(F)
    print("\n".join(report))


if __name__ == "__main__":
    main()
