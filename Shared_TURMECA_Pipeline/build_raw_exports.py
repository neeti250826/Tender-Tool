"""
Raw (unmapped) extracted tender data, one workbook per country, one tab per website / source table.
No taxonomy columns (OU, Sub-OU, Device Category, Company, Family, Included/Review) - only what the portals publish,
plus the parsing-derived fields (winner flag, machine translations) marked as such in the README.
Output: turmeca/raw/<Country>_raw_extracted_data.xlsx
"""
import csv, glob, json, os, re, warnings
from datetime import datetime

warnings.filterwarnings("ignore")
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "raw")
os.makedirs(OUT, exist_ok=True)
DL = r"<DOWNLOADS>"
MAPCOLS = {"OU", "Sub-OU", "Device Category", "Company", "Company Method", "Family Name", "Mapping Method", "Included/Review/Excluded",
           "Device Category (from title)", "Segment", "Level3_Roche", "Level4_Roche", "Level5_Roche", "Segment_Roche", "source_ivd_segment",
           "source_cpv_status", "Tender Stage", "Price Evidence", "Manufacturer Evidence", "Medtronic relevance"}
HDR = PatternFill("solid", fgColor="1F4E78")


def read_csv(p):
    with open(p, encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


def read_xlsx(p, sheet=None):
    wb = openpyxl.load_workbook(p, read_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    h = [c for c in next(it)]
    return [dict(zip(h, r)) for r in it if r and any(x is not None for x in r)]


def num(v):
    if isinstance(v, (int, float)) or v is None:
        return v
    s = str(v).strip().replace(",", "")
    return float(s) if re.fullmatch(r"-?\d+(\.\d+)?", s) else v


class Book:
    def __init__(self, country, readme):
        self.wb = openpyxl.Workbook(write_only=True)
        self.country = country
        ws = self.wb.create_sheet("README")
        ws.append([f"{country} - raw extracted tender data (no taxonomy mapping)"])
        ws.append([f"Built {datetime.now():%Y-%m-%d} from the TURMECA extraction pipeline (<WORKSPACE>/turmeca). "
                   "Mapped versions: TURMECA_Master_Mapped_Data.xlsx."])
        for a, b in readme:
            ws.append([a, b])
        self.counts = []

    def sheet(self, name, rows, drop=(), order=None, derived=()):
        if not rows:
            return
        cols = order or [c for c in rows[0].keys() if c not in MAPCOLS and c not in drop]
        ws = self.wb.create_sheet(name[:31])
        ws.append(cols)
        for r in rows:
            ws.append([num(r.get(c)) if isinstance(r.get(c), str) and c in derived else r.get(c) for c in cols])
        self.counts.append((name, len(rows), len(cols)))

    def save(self):
        p = os.path.join(OUT, f"{self.country}_raw_extracted_data.xlsx")
        self.wb.save(p)
        # second pass for header styling + summary (write-only books cannot style after the fact)
        wb = openpyxl.load_workbook(p)
        for ws in wb.worksheets:
            if ws.title == "README":
                ws.column_dimensions["A"].width = 34
                ws.column_dimensions["B"].width = 140
                ws.append([])
                ws.append(["Tab", "Rows", "Columns"])
                for n, r, c in self.counts:
                    ws.append([n, r, c])
                continue
            for c in ws[1]:
                c.font, c.fill, c.alignment = Font(bold=True, color="FFFFFF"), HDR, Alignment(wrap_text=True, vertical="center")
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
            for i in range(1, ws.max_column + 1):
                ws.column_dimensions[get_column_letter(i)].width = 18
        wb.save(p)
        print(self.country, [(n, r) for n, r, _ in self.counts])


def ksa():
    b = Book("KSA", [
        ("NUPCO_Final_Results", "Award rows parsed from the 42 Final Result PDFs on nupco.com (41 with tables; NPT0033/25 parsed line-by-line, its descriptions "
                                "come from the item list). Columns as printed: NUPCO code, item no, description, vendor no/name (distributor), model/catalog, "
                                "manufacturer, country, group/remarks. Frameworks are multi-award: several rows per item. NUPCO publishes no prices."),
        ("NUPCO_Item_Lists", "Requested (initial) quantity per NUPCO code from each tender's Item List PDF/XLSX."),
        ("NUPCO_Tender_List", "All 119 tenders listed on nupco.com/tenders/tenders-list (ID, status, opening date, documents, URL)."),
        ("NUPCO_Previous_Extraction", "Earlier extraction (KSA_tender_awards_final.xlsx): 13,480 rows / 50 tenders; has supplier + quantity but no "
                                      "manufacturer; its 'amount' is one tender-level figure repeated per row."),
        ("Derived fields", "None - all columns are as published; 'page' = PDF page the row was read from."),
    ])
    fin = read_csv(os.path.join(HERE, "ksa", "nupco_awards_raw.csv")) + read_csv(os.path.join(HERE, "ksa", "nupco_awards_0033_25.csv"))
    dates = json.load(open(os.path.join(HERE, "ksa", "final_result_dates.json")))
    for r in fin:
        r["final_result_date"] = dates.get(r["tender_id"], {}).get("date")
    b.sheet("NUPCO_Final_Results", fin, derived=("page", "sn"))
    b.sheet("NUPCO_Item_Lists", read_csv(os.path.join(HERE, "ksa", "nupco_itemlist_qty.csv")), derived=("qty_requested", "lines"))
    tl = json.load(open(os.path.join(HERE, "sources", "nupco_parsed.json"), encoding="utf8"))
    b.sheet("NUPCO_Tender_List", [{"tender_id": t["id"], "status": t["type"], "title_ar": t["title"], "opening": t["opening"],
                                   "documents": ", ".join(t["docs"]), "url": t["href"],
                                   "document_links": " | ".join(f"{p['t']}: {p['h']}" for p in t["pdfs"])} for t in tl])
    b.sheet("NUPCO_Previous_Extraction", read_xlsx(os.path.join(DL, "KSA_tender_awards_final.xlsx")))
    b.save()


def turkey():
    b = Book("Turkey", [
        ("EKAP_Contract_Lots", "One row per contract x lot from EKAP v2 tender details (this project's pull + cached details from the earlier IVD pull): "
                               "tender (İKN), title, authority, city, tender date, OKAS codes, lot name (TR + machine EN), supplier, contract date/value, "
                               "approximate cost, termination. Contract value is per contract and may cover several lots. EKAP publishes no quantities or manufacturers."),
        ("EKAP_Tender_Lists", "Tender list hits per OKAS code x year from the search API (id, İKN, title, authority, dates)."),
        ("EKAP_Previous_Extraction", "Earlier QC-corrected extraction (Turkey_EKAP_Contracts_Ending_2024_2026_EN_Corrected.xlsx): same contracts with corrected "
                                     "English, contract start/end dates, low/high bids. IVD (Roche) segment columns removed."),
        ("Derived fields", "Lot Name (EN) = machine translation (translate.googleapis.com); Qty in Lot Name / Unit Price TRY = parsed from the lot text."),
    ])
    lots = read_csv(os.path.join(HERE, "ekap", "tr_ekap_lots_all.csv"))
    b.sheet("EKAP_Contract_Lots", lots, derived=("Contract Value TRY", "Line Value TRY", "Qty in Lot Name", "Unit Price TRY", "Approx. Cost TRY (tender/lot)", "Lots in Contract", "Award Year"))
    tl = []
    for f in sorted(glob.glob(os.path.join(HERE, "ekap", "lists", "*.json"))):
        d = json.load(open(f, encoding="utf8"))
        for x in d["list"]:
            tl.append({"okas_code": d["code"], "search_year": d["year"], "ekap_id": x.get("id"), "ikn": x.get("ikn"), "title": x.get("ihaleAdi"),
                       "tender_type": x.get("ihaleTipAciklama"), "procedure": x.get("ihaleUsulAciklama"), "status": x.get("ihaleDurumAciklama") or x.get("ihaleDurum"),
                       "authority": x.get("idareAdi"), "city": x.get("ilAdi") or x.get("il"), "tender_date": x.get("ihaleTarihi"),
                       "notice_date": x.get("ilanTarihi"), "url": f"https://ekapv2.kik.gov.tr/ekap/search/{x.get('id')}"})
    b.sheet("EKAP_Tender_Lists", tl)
    prev = []
    for sh in ("2024 Data", "2025 Data", "2026 Data"):
        for r in read_xlsx(os.path.join(DL, "Turkey_EKAP_Contracts_Ending_2024_2026_EN_Corrected.xlsx"), sh):
            r["contract_end_year_tab"] = sh[:4]
            prev.append(r)
    b.sheet("EKAP_Previous_Extraction", prev)
    b.save()


def pakistan():
    b = Book("Pakistan", [
        ("KP_MCC_BER_Bids", "Govt. MCC bid evaluation reports (healthkp.gov.pk, dghskp.gov.pk) FY2023-24..2026-27: one row per bid - firm, MCC ref, item text, "
                            "rate offered (PKR), lowest rate, final score. FY2025-26 is the draft comparative statement (approved BER is a scan)."),
        ("KP_DGHS_Equipment_BER_Bids", "DGHS 'Selection and Rate Contracting of Medical Equipment, Furniture, Instruments' BERs FY2024-25..2026-27: one row per bid "
                                       "with item (Model:/Make: parsed), technical score, quoted price, lowest price, total marks."),
        ("KPPRA_Evaluated_Tenders", "KPPRA evaluated goods tenders for KP health entities (list only; reports are technical-only)."),
        ("Federal_EPMS_Evaluations", "All 3,676 federal evaluation listings on epms.ppra.gov.pk (evaluation no, tender no, title, listing text, links)."),
        ("Federal_Medical_Evaluations", "The 99 medical-related evaluations with detail-page metadata (organisation, office, type, bids, lowest bidder, dates)."),
        ("Federal_Item_Bids", "Item-level bids parsed from EPADS-generated reports (item, UNSPSC, bidder, marks, offer = line total)."),
        ("Federal_Package_Bids", "Package-level reports (bidder, evaluated cost, most advantageous bidder); many OCR'd from scans - check amounts."),
        ("Federal_EPMS_Contracts", "EPMS public contracts register (694 contracts, all sectors)."),
        ("Derived fields", "winner / bids_on_item / block = derived by the parser (highest score wins); everything else as printed."),
    ])
    b.sheet("KP_MCC_BER_Bids", read_csv(os.path.join(HERE, "pakistan", "kp_mcc_bids.csv")), drop=("item_key", "description_enriched"), derived=("rate_pkr", "lowest_pkr", "final_score", "page", "bids_on_item"))
    b.sheet("KP_DGHS_Equipment_BER_Bids", read_csv(os.path.join(HERE, "pakistan", "kp_ber_bids.csv")), derived=("tech_score", "quoted_pkr", "lowest_pkr", "total_marks", "page", "bids_on_item"))
    b.sheet("KPPRA_Evaluated_Tenders", read_csv(os.path.join(HERE, "pakistan", "kppra", "kppra_evaluated.csv")))
    b.sheet("Federal_EPMS_Evaluations", read_csv(os.path.join(HERE, "pakistan", "federal", "epms_evaluations.csv")))
    b.sheet("Federal_Medical_Evaluations", read_csv(os.path.join(HERE, "pakistan", "federal", "federal_medical_evaluations.csv")))
    fx = os.path.join(HERE, "pakistan", "federal", "PK_Federal_Evaluations_Mapped.xlsx")
    b.sheet("Federal_Item_Bids", read_xlsx(fx, "Item_Bids"))
    b.sheet("Federal_Package_Bids", read_xlsx(fx, "Package_Bids"))
    b.sheet("Federal_EPMS_Contracts", read_csv(os.path.join(HERE, "pakistan", "federal", "epms_contracts.csv")))
    b.save()


def kuwait():
    b = Book("Kuwait", [
        ("CAPT_Winning_Bids", "All winning bids on capt.gov.kw/en/tenders/winning-bids (meeting dates Jan 2024 - Sep 2026, all ministries): meeting date, tender no, "
                              "subject (Arabic) + machine English, organisation, winner, total KD, item-level winners where CAPT shows a popup."),
        ("Derived fields", "Subject (EN, machine) and Total USD (approx) are derived; Health-related = keyword flag."),
    ])
    rows = read_xlsx(os.path.join(HERE, "kuwait", "KW_CAPT_Awards.xlsx"), "All_Awards")
    b.sheet("CAPT_Winning_Bids", rows, drop=("Use",))
    b.save()


LINE = re.compile(r"^\s*(\d+)\s*:\s*(?:\[([^\]]+)\])?\s*(.+?)\s*-+\s*([\d.,]+)\s*([A-Za-z][A-Za-z ()/]{0,20})?\s*$")


def uae():
    b = Book("UAE", [
        ("eSupply_Opportunities", "Dubai eSupply past opportunities (RFQs) of buyers containing 'Health' (Dubai Academic Health Corporation, Dubai Health Authority), "
                                  "published 2024-01-01..2026-10-01, exported day by day with eSupply's Export List. All export columns kept."),
        ("eSupply_Item_Lines", "Item lines parsed from each opportunity's Description ('n: [code] ITEM - qty UOM'). eSupply publishes no award, winner or price."),
        ("Derived fields", "Item line split (line no, item code, item text, qty, UOM) is parsed from the Description text."),
    ])
    opp, items, H = {}, [], None
    for f in sorted(glob.glob(os.path.join(HERE, "uae", "daily", "*.xlsx"))):
        ws = openpyxl.load_workbook(f).worksheets[0]
        rows = list(ws.iter_rows(values_only=True))
        h = next((i for i, r in enumerate(rows) if r and r[0] == "Project Code"), None)
        if h is None:
            continue
        hdr = list(rows[h])
        H = H or [x for x in hdr if x]
        for r in rows[h + 1:]:
            if not r or not r[0] or r[0] in opp:
                continue
            d = dict(zip(hdr, r))
            opp[r[0]] = d
            for ln in str(d.get("Description") or "").splitlines():
                m = LINE.match(ln)
                if m:
                    try:
                        q = float(m.group(4).replace(",", ""))
                    except ValueError:
                        q = None
                    items.append({"Project Code": r[0], "Project Title": d.get("Project Title"), "Buyer Organisation": d.get("Buyer Organisation"),
                                  "Publication Date": d.get("Publication Date"), "Closing date/time": d.get("Closing date/time"), "Line": int(m.group(1)),
                                  "Item Code": m.group(2), "Item": m.group(3).strip(" -"), "Qty": q, "UOM": (m.group(5) or "").strip()})
    uniq = []
    for c in H:
        if c not in uniq:
            uniq.append(c)
    b.sheet("eSupply_Opportunities", list(opp.values()), order=uniq)
    b.sheet("eSupply_Item_Lines", items)
    b.save()


if __name__ == "__main__":
    ksa(); turkey(); pakistan(); kuwait(); uae()
