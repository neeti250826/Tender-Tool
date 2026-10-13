"""Offline fix for '(EN)' columns in TURMECA_Master_Mapped_Data.xlsx when Google Translate is rate-limited:
use the cached full-text translation of the ORIGINAL Turkish value (translation_cache.sqlite, from the first run) and
correct the known procurement mistranslation 'KALEM' -> 'pen' (it means 'item(s)' in tenders)."""
import os, re, sqlite3
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
F = os.path.join(HERE, "TURMECA_Master_Mapped_Data.xlsx")
TRC = re.compile(r"[ğĞışŞİçÇöÖüÜ]")
TURKISH_LEFT = re.compile(r"\b(ALIMI|ALIM|TEMINI|TEMİNİ|ICIN|İÇİN|IHALESI|İHALESİ|MALZEMESI|MALZEMESİ|BIRIMLERI|HASTANE|SARF|TIBBI|TIBBİ|KALEM|KISIM|ISI|İŞİ)\b", re.I)


def fix_pen(src, en):
    if not re.search(r"kalem", src, re.I):
        return en
    en = re.sub(r"\bpens?\b", "items", en, flags=re.I)
    en = re.sub(r"\bstationery\b", "items", en, flags=re.I) if not re.search(r"k[ıi]rtasiye", src, re.I) else en
    return en


def main():
    con = sqlite3.connect(os.path.join(HERE, "translation_cache.sqlite"))
    cache = dict(con.execute("SELECT text, en FROM t WHERE src='tr' AND en <> ''").fetchall())
    wb = openpyxl.load_workbook(F)
    ws = wb["Turkey - EKAP"]
    hdr = [c.value for c in ws[1]]
    stats = {}
    for col in ("Tender Title", "Procurement Authority", "Hospital/Buyer", "Supplier/Distributor"):
        j, je = hdr.index(col) + 1, hdr.index(col + " (EN)") + 1
        fixed = left = 0
        for r in range(2, ws.max_row + 1):
            if not ws.cell(r, 1).value:
                break
            src, en = ws.cell(r, j).value, ws.cell(r, je).value
            if not isinstance(src, str):
                continue
            good = cache.get(src.strip()) or cache.get(src)
            bad = (not en) or en == src or TURKISH_LEFT.search(str(en)) or bool(TRC.search(src)) and str(en).upper() == src.upper().translate(str.maketrans("ĞIŞİÇÖÜ", "GISICOU"))
            if bad and good and good.strip().upper() != src.strip().upper():
                ws.cell(r, je, fix_pen(src, good))
                fixed += 1
            elif en and re.search(r"\bpens?\b", str(en), re.I):
                ws.cell(r, je, fix_pen(src, str(en)))
                fixed += 1
            if TURKISH_LEFT.search(str(ws.cell(r, je).value or "")):
                left += 1
        stats[col] = (fixed, left)
    wb.save(F)
    for c, (f, l) in stats.items():
        print(f"{c}: fixed {f} rows; rows still showing Turkish words {l}")


if __name__ == "__main__":
    main()
