"""
Add CY2025 'applicable period' columns to every data tab of the TURMECA master (Google Sheet export):
  Contract Start Date | Contract End Date | Applicable Duration (months) | Duration Source | Active in 2025? |
  Months Active in 2025 | 2025 Applicable Share | 2025 Applicable Volume | 2025 Applicable Value (LCY) |
  2025 Applicable Value USD | 2025 Basis / Note
Assessment period = calendar year 2025. Share = days of the contract that fall in 2025 / total contract days.
Duration sources: EKAP contract start/end (previous QC file), KP financial-year rate contracts (Jul-Jun),
eSupply one-off RFQs (publication date). NUPCO framework terms are not published -> Unknown.
Outputs: cy2025/<tab>.tsv (for pasting into the Google Sheet) and cy2025/TURMECA_Master_Mapped_Data_CY2025.xlsx
"""
import csv, os, re, sys
from datetime import date, datetime, timedelta

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "TURMECA_Master_Mapped_Data.xlsx")
OUTD = os.path.join(HERE, "cy2025")
os.makedirs(OUTD, exist_ok=True)
PREV_TR = r"<DOWNLOADS>/Turkey_EKAP_Contracts_Ending_2024_2026_EN_Corrected.xlsx"
Y0, Y1 = date(2025, 1, 1), date(2025, 12, 31)
MONTH = 365.25 / 12
NEW = ["Contract Start Date", "Contract End Date", "Applicable Duration (months)", "Duration Source", "Active in 2025?",
       "Months Active in 2025", "2025 Applicable Share", "2025 Applicable Volume", "2025 Applicable Value (LCY)",
       "2025 Applicable Value USD", "2025 Basis / Note"]


def d(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    m = re.search(r"(\d{4})-(\d{2})-(\d{2})", str(v or ""))
    if m:
        return date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    m = re.search(r"(\d{2})\.(\d{2})\.(\d{4})", str(v or ""))
    if m:
        return date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
    return None


def num(v):
    try:
        return float(v) if v not in (None, "") else None
    except (TypeError, ValueError):
        return None


def period(start, end):
    """-> (duration_months, active, months_in_2025, share)"""
    total = (end - start).days + 1
    if total <= 0:
        return None, "Unknown", None, None
    lo, hi = max(start, Y0), min(end, Y1)
    ov = (hi - lo).days + 1 if hi >= lo else 0
    return round(total / MONTH, 1), ("Yes" if ov > 0 else "No"), round(ov / MONTH, 1), round(ov / total, 4)


def fill(r, start, end, dur_src, note, one_off=False):
    vol, lcy, usd = num(r.get("Volume")), num(r.get("Total Award Value")), (num(r.get("Value USD")) or num(r.get("Est. Value USD (ref-price based)")))
    line = str(r.get("Price Evidence") or "").startswith("Line price") or r.get("Country") == "KSA"
    if start and end:
        dur, active, months, share = period(start, end)
        if one_off:
            dur, months = "n/a (one-off)", "n/a"
        return [start, end, dur, dur_src, active, months, share,
                round(vol * share, 2) if (vol is not None and share is not None) else None,
                round(lcy * share, 2) if (lcy is not None and share is not None and line) else None,
                round(usd * share, 2) if (usd is not None and share is not None and (line or usd and r.get("Country") == "KSA")) else None,
                note if line else (note + "; value left blank - contract value covers several lots")]
    # duration unknown
    active = "Unknown"
    if start:
        active = "Yes" if (Y0 <= start <= Y1) else ("Unknown" if start < Y0 else "No")
    return [start, None, "Unknown", dur_src, active, None, None, None, None, None, note]


def turkey_dates():
    """Row ID -> (start, end) for Turkey rows: this project's rows via tr_ekap_lots_all.csv index, previous-only rows via notice/item."""
    prev = {}
    wb = openpyxl.load_workbook(PREV_TR, read_only=True)
    for sh in ("2024 Data", "2025 Data", "2026 Data"):
        it = wb[sh].iter_rows(values_only=True)
        h = next(it)
        for row in it:
            if not row or not any(row):
                continue
            p = dict(zip(h, row))
            k = (str(p.get("notice_id")), str(p.get("item_description") or "").strip().upper(), str(p.get("supplier_name") or "").strip().upper())
            prev[k] = (d(p.get("start_date")), d(p.get("end_date")))
            prev[("PREV", str(p.get("notice_id")), str(p.get("item_no")))] = (d(p.get("start_date")), d(p.get("end_date")))
    with open(os.path.join(HERE, "ekap", "tr_ekap_lots_all.csv"), encoding="utf-8-sig") as f:
        lots = list(csv.DictReader(f))
    out = {}
    for i, a in enumerate(lots):
        k = (a["Tender ID"], a["Lot Name (TR)"].strip().upper(), a["Supplier/Distributor"].strip().upper())
        if k in prev:
            out[f"TR-EKAP-{a['Tender ID']}-{i}"] = prev[k]
    return out, prev


def main():
    wb = openpyxl.load_workbook(SRC)
    tr_idx, prev = turkey_dates()
    summary = []
    for ws in wb.worksheets:
        if ws.cell(1, 1).value != "Row ID":
            continue
        hdr = [c.value for c in ws[1]]
        if "Active in 2025?" in hdr:
            continue
        c0 = len(hdr) + 1
        for j, h in enumerate(NEW):
            c = ws.cell(1, c0 + j, h)
            c.font, c.fill, c.alignment = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="548235"), Alignment(wrap_text=True, vertical="center")
        tsv, stats = [], {"rows": 0, "Yes": 0, "No": 0, "Unknown": 0, "with_share": 0}
        for ri in range(2, ws.max_row + 1):
            rid = ws.cell(ri, 1).value
            if not rid:
                break
            r = {h: ws.cell(ri, j + 1).value for j, h in enumerate(hdr)}
            country, portal = r["Country"], r["Portal"]
            if country == "Turkey":
                se = tr_idx.get(rid) or (prev.get(("PREV", str(r["Tender ID"]), str(r["Lot/Item No."]))) if str(rid).startswith("TR-EKAP-PREV") else None)
                start, end = se if se else (d(r["Award/Contract Date"]), None)
                vals = fill(r, start, end, "EKAP contract start/end dates (QC-corrected extraction)" if se else "EKAP contract date only; end date not available",
                            "Contract days in CY2025 / total contract days" if se else "End date unknown - 2025 share not computable")
            elif country == "KSA":
                rd = d(r["Result Date"])
                vals = fill(r, rd, None, "Not published (NUPCO open framework; term not stated in final result)",
                            "Active flag = final result published in 2025 (award-year basis). Framework term and call-off timing not published, so months/share unknown; "
                            "previous extraction lists only a 90/120-day delivery period per order")
            elif country == "Pakistan" and "KP" in str(portal):
                m = re.search(r"FY(\d{4})-(\d{2})", str(r["Tender ID"]))
                y = int(m.group(1)) if m else None
                start, end = (date(y, 7, 1), date(y + 1, 6, 30)) if y else (None, None)
                vals = fill(r, start, end, "KP financial-year rate contract (1 Jul - 30 Jun)",
                            "Rate contract; BER gives unit price only (no quantity/total), so applicable volume/value stay blank")
            elif country == "Pakistan":
                rd = d(r["Result Date"])
                vals = fill(r, rd, None, "Not published (federal evaluation; contract period not stated)", "Evaluation published " + (str(rd) if rd else "n/a"))
            elif country == "UAE":
                pd_ = d(r["Publication Date"])
                vals = fill(r, pd_, pd_, "One-off RFQ (publication date)", "One-off purchase request: full requested quantity attributed to the publication year", one_off=True)
            else:
                ad = d(r["Award/Contract Date"])
                vals = fill(r, ad, None, "Not published", "Award meeting date only")
            for j, v in enumerate(vals):
                c = ws.cell(ri, c0 + j, v)
                if isinstance(v, date):
                    c.number_format = "yyyy-mm-dd"
            tsv.append(["" if v is None else (v.isoformat() if isinstance(v, date) else v) for v in vals])
            stats["rows"] += 1
            stats[vals[4]] += 1
            stats["with_share"] += vals[6] is not None
        with open(os.path.join(OUTD, re.sub(r"[^\w-]+", "_", ws.title) + ".tsv"), "w", encoding="utf-8", newline="") as f:
            f.write("\t".join(NEW) + "\n")
            for row in tsv:
                f.write("\t".join(str(v).replace("\t", " ").replace("\n", " ") for v in row) + "\n")
        summary.append((ws.title, get_column_letter(c0), stats))
        for j in range(len(NEW)):
            ws.column_dimensions[get_column_letter(c0 + j)].width = 16
    wb.save(os.path.join(OUTD, "TURMECA_Master_Mapped_Data_CY2025.xlsx"))
    for t, col, s in summary:
        print(f"{t:26} first new col {col:4} {s}")


if __name__ == "__main__":
    main()
