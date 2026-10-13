"""Write the TURMECA mapping prompt (TURMECA_Mapping_Prompt_as_built.md) to a formatted Word document."""
import os, re
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "TURMECA_Mapping_Prompt_as_built.md")
OUT = os.path.join(HERE, "TURMECA_Mapping_Prompt.docx")
TAX = "https://docs.google.com/spreadsheets/d/<TAXONOMY_SHEET_ID>/edit?gid=1456151223#gid=1456151223"
NAVY = RGBColor(0x1F, 0x4E, 0x78)

prompt = re.search(r"```\n(.*?)```", open(SRC, encoding="utf-8").read(), re.S).group(1).rstrip("\n")


def hyperlink(par, url, text):
    rid = par.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    h = OxmlElement("w:hyperlink"); h.set(qn("r:id"), rid)
    r = OxmlElement("w:r"); rp = OxmlElement("w:rPr")
    c = OxmlElement("w:color"); c.set(qn("w:val"), "0563C1"); rp.append(c)
    u = OxmlElement("w:u"); u.set(qn("w:val"), "single"); rp.append(u)
    r.append(rp); t = OxmlElement("w:t"); t.text = text; t.set(qn("xml:space"), "preserve"); r.append(t); h.append(r)
    par._p.append(h)


def add_text(par, text):
    """Plain text, with the taxonomy URL turned into a clickable link."""
    if TAX in text:
        a, b = text.split(TAX, 1)
        if a:
            par.add_run(a)
        hyperlink(par, TAX, TAX)
        if b:
            par.add_run(b)
    else:
        par.add_run(text)


doc = Document()
st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10.5)
for s in ("Heading 1", "Heading 2"):
    doc.styles[s].font.color.rgb = NAVY

t = doc.add_heading("TURMECA Tender Data - Taxonomy Mapping Prompt", 0)
p = doc.add_paragraph("Use this prompt to map TURMECA tender data (KSA, Turkey, Pakistan, Kuwait, UAE) or similar tender extracts "
                      "to the approved Clearstate taxonomy. Replace the bracketed fields before use. A plain-text copy for pasting "
                      "into an AI tool is at the end of this document.")
p.runs[0].italic = True

# ---- formatted version: section headers -> Heading 2, "- " lines -> bullets, deeper-indented lines -> continuation
lines = prompt.split("\n")
doc.add_heading("Prompt", 1)
para = None
for ln in lines:
    if not ln.strip():
        para = None
        continue
    head = re.match(r"(INPUTS|OUTPUT COLUMNS|OUTPUT COLUMNS TO ADD|GENERAL RULES|STEP \d[A-Z]? - .*|QUALITY CHECKS BEFORE RETURNING)$", ln.strip())
    if head and not ln.startswith(" "):
        txt = ln.strip()
        txt = {"INPUTS": "Inputs", "OUTPUT COLUMNS TO ADD": "Output columns to add", "OUTPUT COLUMNS": "Output columns", "GENERAL RULES": "General rules",
               "QUALITY CHECKS BEFORE RETURNING": "Quality checks before returning"}.get(txt, txt)
        if txt.startswith("STEP"):
            n, rest = re.match(r"STEP (\d[A-Z]?) - (.*)", txt).groups()
            rest = rest if "(" in rest else rest.title().replace("Ou", "OU").replace("Sub-OU", "Sub-OU")
            txt = f"Step {n} - " + re.sub(r"^([A-Z ,\-/]+)(\s*\(|$)", lambda m: m.group(1).title().replace("Ou", "OU") + m.group(2), rest)
        doc.add_heading(txt.replace(" And ", " and "), 2)
        para = None
    elif re.match(r"\s*(Included|Review|Excluded|High|Medium|Low)\s*=", ln) or re.match(r'\s*"(Not in taxonomy|Undefined|Not related)"\s*=', ln) or re.match(r"\s+([1-9]|[a-d])\. ", ln):
        para = doc.add_paragraph(style="List Bullet 2"); add_text(para, ln.strip())
    elif re.match(r"(Fill|Add): ", ln):
        para = doc.add_paragraph(style="List Bullet"); add_text(para, ln)
    elif ln.startswith("- "):
        para = doc.add_paragraph(style="List Bullet"); add_text(para, ln[2:])
    elif ln.startswith(" ") and para is not None:
        add_text(para, " " + ln.strip())
    elif para is not None and para.text.endswith("|") or (para is not None and "|" in ln and "|" in para.text):
        add_text(para, " " + ln.strip())
    else:
        para = doc.add_paragraph(); add_text(para, ln)

# ---- copy-paste version
doc.add_page_break()
doc.add_heading("Copy-paste version", 1)
for ln in lines:
    q = doc.add_paragraph(); q.paragraph_format.space_after = Pt(0)
    r = q.add_run(ln if ln else " "); r.font.name = "Consolas"; r.font.size = Pt(9)
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")

doc.save(OUT)
print(OUT)
