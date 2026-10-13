"""Final pass on Turkey '(EN)' columns: translate remaining Turkish procurement words with a fixed glossary
(offline; Google Translate was rate-limited). Brand names are kept, Turkish letters are ASCII-folded."""
import os, re
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
F = os.path.join(HERE, "TURMECA_Master_Mapped_Data.xlsx")

# longest phrases first; patterns run on an ASCII-folded, upper-cased string
G = [
    ("IL AMBULANS SERVISI BASHEKIMLIGINE", "PROVINCIAL AMBULANCE SERVICE"), ("ILI BAGLI SAGLIK TESISLERI", "PROVINCE AFFILIATED HEALTH FACILITIES"),
    ("ISM BAGLI SAGLIK TESISLERI", "PROVINCIAL HEALTH DIRECTORATE AFFILIATED HEALTH FACILITIES"), ("MUDURLUGUMUZE BAGLI SAGLIK TESISLERI", "HEALTH FACILITIES AFFILIATED WITH OUR DIRECTORATE"),
    ("MUDURLUGUMUZE BAGLI", "AFFILIATED WITH OUR DIRECTORATE"), ("ORGAN NAKLI KLINIGI", "ORGAN TRANSPLANT CLINIC"), ("DIS KLINIGI", "DENTAL CLINIC"),
    ("HASTANE GENELI", "HOSPITAL-WIDE"), ("DEVLET HASTANESI", "STATE HOSPITAL"), ("AGIZ-DIS SAGLIGI", "ORAL AND DENTAL HEALTH"),
    ("OZEL SAGLIK HIZMETLERI", "PRIVATE HEALTH SERVICES"), ("TASFIYE HALINDE", "IN LIQUIDATION"), ("ARAC VE GERECLERI", "TOOLS AND EQUIPMENT"),
    ("LIMITED SIRKETI", "LIMITED COMPANY"), ("ANONIM SIRKETI", "JOINT STOCK COMPANY"), ("IC VE DIS TICARET", "DOMESTIC AND FOREIGN TRADE"),
    ("DIS TICARET", "FOREIGN TRADE"), ("YEDEK PARCA", "SPARE PARTS"), ("PLAN PROJE", "PLANNING PROJECT"), ("ECZA DEPOSU", "PHARMACEUTICAL WAREHOUSE"),
    ("DIS DEPOSU", "DENTAL DEPOT"), ("SARF MALZEMELERI", "CONSUMABLES"), ("SARF MALZEMESI", "CONSUMABLES"), ("SARF MALZEME", "CONSUMABLES"),
    ("TIBBI CIHAZLARI", "MEDICAL DEVICES"), ("TIBBI CIHAZLAR", "MEDICAL DEVICES"), ("TIBBI CIHAZ", "MEDICAL DEVICE"), ("TIBBI URUNLER", "MEDICAL PRODUCTS"),
    ("TIBBI MALZEMELERI", "MEDICAL SUPPLIES"), ("TIBBI MALZEMELER", "MEDICAL SUPPLIES"), ("TIBBI MALZEME", "MEDICAL SUPPLIES"), ("TIBBI GERECLER", "MEDICAL EQUIPMENT"),
    ("TIBBI GAZ", "MEDICAL GAS"), ("MEDIKAL GAZ", "MEDICAL GAS"),
    (r"(\d+)\s*KALEM", r"\1 ITEMS"), ("KALEM", "ITEMS"), (r"(\d+)\s*KISIM", r"\1 LOTS"), ("KISIM", "LOT"),
    ("MALZEMELERI", "MATERIALS"), ("MALZEMESI", "MATERIALS"), ("MALZEME", "MATERIALS"), ("ALIMI", "PROCUREMENT"), ("ALIM", "PROCUREMENT"),
    ("IHALESI", "TENDER"), ("ISI", "WORK"), ("TEMINI", "SUPPLY"), ("TIBBI", "MEDICAL"), ("MEDIKAL", "MEDICAL"), ("CIHAZLARI", "DEVICES"), ("CIHAZLAR", "DEVICES"),
    ("CIHAZ", "DEVICE"), ("URUNLERI", "PRODUCTS"), ("URUNLER", "PRODUCTS"), ("GERECLER", "EQUIPMENT"), ("LABORATUVAR", "LABORATORY"),
    ("LABOROTUVAR", "LABORATORY"), ("LABORATUAR", "LABORATORY"), ("BIYOMEDIKAL", "BIOMEDICAL"), ("BIYOMEDIK", "BIOMEDICAL"), ("KARDIYOLOJI", "CARDIOLOGY"),
    ("ANJIYOGRAFI", "ANGIOGRAPHY"), ("ANJIYO", "ANGIO"), ("ANJIO", "ANGIO"), ("GIRISIMSEL", "INTERVENTIONAL"), ("PERFUZYON", "PERFUSION"),
    ("KVC", "CARDIOVASCULAR SURGERY"), ("AMELIYATHANE", "OPERATING ROOM"), ("BIRIMINE", "UNIT"), ("UNITESI", "UNIT"), ("PATOLOJI", "PATHOLOGY"),
    ("TUKETIM", "CONSUMPTION"), ("KIMYASAL", "CHEMICAL"), ("YILLARI", "YEARS"), ("YILI", "YEAR"), ("AYLIK", "MONTHLY"), ("ICIN", "FOR"), ("GOZ", "EYE"),
    ("DIS", "DENTAL"), ("ILAC", "PHARMACEUTICAL"), ("ECZA", "PHARMACEUTICAL"), ("DEPOSU", "DEPOT"), ("SAN\\.", "IND."), ("SANAYI", "INDUSTRY"),
    ("TICARET", "TRADE"), ("TIC\\.", "TRADE"), ("IMALAT", "MANUFACTURING"), ("INSAAT", "CONSTRUCTION"), ("INS\\.", "CONSTR."), ("GIDA", "FOOD"),
    ("TURIZM", "TOURISM"), ("TEKSTIL", "TEXTILE"), ("TEMIZLIK", "CLEANING"), ("ITHALAT", "IMPORT"), ("IHRACAT", "EXPORT"), ("PAZARLAMA", "MARKETING"),
    ("HIZMETLERI", "SERVICES"), ("SAGLIK", "HEALTH"), ("MIMARLIK", "ARCHITECTURE"), ("ENERJI", "ENERGY"), ("TEKNOLOJIK", "TECHNOLOGICAL"),
    ("GAZ", "GAS"), ("SISTEMLERI", "SYSTEMS"), ("OZEL", "PRIVATE"), ("VE", "AND"), ("AD", "DEPT."),
]
G = [(re.compile(r"(?<![A-Z0-9])" + a + r"(?![A-Z0-9])"), b) for a, b in G]
FOLD = str.maketrans("ĞğIıŞşİiÇçÖöÜü", "GGIISSIICCOOUU")
LEFT = re.compile(r"(?<![A-Z])(KALEM|ALIMI|SARF|MALZEME\w*|TIBBI|TİBBİ|TIBBİ|CİHAZ\w*|URUNLER|ÜRÜNLER|İHALESİ|IHALESI|LABORATUVAR)(?![A-Z])")


def glossary(s):
    t = re.sub(r"\s+", " ", str(s).upper().translate(FOLD)).strip()
    for rx, b in G:
        t = rx.sub(b, t)
    return re.sub(r"\s+", " ", t).strip()


def main():
    wb = openpyxl.load_workbook(F)
    ws = wb["Turkey - EKAP"]
    hdr = [c.value for c in ws[1]]
    for col in ("Tender Title", "Supplier/Distributor", "Procurement Authority", "Hospital/Buyer"):
        j, je = hdr.index(col) + 1, hdr.index(col + " (EN)") + 1
        n = 0
        for r in range(2, ws.max_row + 1):
            if not ws.cell(r, 1).value:
                break
            en = ws.cell(r, je).value
            if isinstance(en, str) and LEFT.search(en.upper()):
                ws.cell(r, je, glossary(en))
                n += 1
        print(f"{col}: glossary-fixed {n} rows")
    wb.save(F)


if __name__ == "__main__":
    for t in ["35 KALEM SARF MALZEME ALIMI", "MERİL TIBBİ CİHAZLAR İMALAT VE TİCARET ANONİM ŞİRKETİ",
              "PURCHASE OF CONSUMABLES FOR 23 KALEM EYE AND ORGAN TRANSPLANTATION CLINIC", "6 AYLIK KVC-ANJIO TIBBI SARF MALZEME ALIMI",
              "ÖZGATA TIBBİ MALZEMELER İLAÇ SAN. GIDA TURİZM VE İNŞ. TRADE LIMITED CO"]:
        print(" ", t, "=>", glossary(t))
    main()
