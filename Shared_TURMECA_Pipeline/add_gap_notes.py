"""Post-process TURMECA_Master_Mapped_Data.xlsx: explain empty/thin website tabs and add tabs for websites with no data.
Run after build_master_mapped.py."""
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

F = os.path.join(os.path.dirname(os.path.abspath(__file__)), "TURMECA_Master_Mapped_Data.xlsx")
NOTE_FILL = PatternFill("solid", fgColor="FFF2CC")

EXISTING = {
    "Kuwait - CAPT": [
        "Status: data collected, but nothing could be mapped to the taxonomy (0 rows).",
        "What we got: 808 CAPT winning-bid awards (Jan 2024 - Sep 2026), incl. 70 health-related, with tender no., organisation, winner and total KD (see kuwait/KW_CAPT_Awards.xlsx).",
        "Why no mapped rows: CAPT only publishes Kuwait's large central tenders. The Ministry of Health awards on CAPT are services and works "
        "(guarding, cleaning, hotel apartments for medical staff, medical gases, vehicles, maintenance) - no medical devices.",
        "MOH buys devices through its own direct purchases / practices, which are not published on CAPT; item-level award details sit in CAPT's paid "
        "electronic archive (250 KWD, registered companies only).",
        "Other limits: no item lines or unit prices on CAPT; Jan-May 2024 has almost no awards published; site is behind Cloudflare (accessed via the user's browser after accepting the site's terms).",
    ],
    "Pakistan - Federal EPMS": [
        "Status: only 3 mapped rows (silk sutures, Islamabad District Health Office FY2026-27).",
        "What we got: EPMS lists 3,676 federal evaluation reports; 99 medical-related reports downloaded (scanned ones OCR'd) - see pakistan/federal/PK_Federal_Evaluations_Mapped.xlsx.",
        "Why so little mapped: federal medical procurement is mostly (a) drugs, or (b) capital-equipment packages for one new hospital (Federal Govt Polyclinic PGMI-II: "
        "ventilators, monitors, electrosurgical, anaesthesia, ECG, laparoscope, OT tables) whose package titles do not correspond to a taxonomy product category.",
        "Package reports give only bidders, evaluated package cost and the winning bidder - no item lines, quantities, unit prices or manufacturers; many are scans, so OCR'd amounts are incomplete.",
        "EPMS public contracts register has only 694 contracts in total, almost none medical; the legacy ppra.org.pk 'contracts over Rs 50M' page now redirects to EPMS. EPADS itself is login-only.",
    ],
    "UAE - Dubai eSupply": [
        "Note: rows here are requested quantities only (volume signal). eSupply publishes no award, winner or price, so no value or company can be assigned.",
    ],
}

NEW_TABS = {
    "KSA - Etimad": [
        "Website: tenders.etimad.sa (Saudi government e-procurement)",
        "Status: no data extracted.",
        "Why: award information is tender-level only (no item lines, no manufacturer). Fetching award details at volume triggers Etimad's anti-bot "
        "protection and then redirects to a login, so a reliable bulk pull was not possible. NUPCO (the KSA health procurement agency) was used instead.",
    ],
    "Pakistan - KPPRA": [
        "Website: kppra.gov.pk (Khyber Pakhtunkhwa procurement authority)",
        "Status: searched, no usable data.",
        "Why: only 54 evaluated goods tenders exist across all 216 KP health entities; the few device/consumable ones (e.g. MTI Ayub Teaching Hospital 2025) "
        "are technical-only evaluations (document checklists and marks) with no prices, items or manufacturers. KP data came from DGHS/Health Dept evaluation reports instead.",
    ],
    "Pakistan - Sindh PPRA": [
        "Website: pprasindh.gov.pk (Sindh procurement, incl. BER CSV export)",
        "Status: not accessible.",
        "Why: the site refuses connections from outside Pakistan (empty reply / helpdesk error page), also through the VPN's 'Pakistan' location, which "
        "physically exits in Singapore. Needs a genuine Pakistani IP or a machine in Pakistan.",
    ],
    "Pakistan - Punjab DGHS": [
        "Website: dghs.punjab.gov.pk (Punjab health tenders / evaluation reports)",
        "Status: not accessible.",
        "Why: same as Sindh - the server returns an empty reply to non-Pakistani connections, including the VPN's Singapore-based 'Pakistan' exit. Needs a genuine Pakistani IP.",
    ],
    "UAE - Abu Dhabi ADGPG": [
        "Website: adgpg.gov.ae (Abu Dhabi Government Procurement Gate)",
        "Status: not extracted.",
        "Why: only open tender notices are published (medical category ~214 tenders) - no winners, quantities or prices. Phase 4 scope was Dubai eSupply quantities only.",
    ],
    "UAE - Federal MoF DPP": [
        "Website: mof.gov.ae Digital Procurement Platform (federal UAE)",
        "Status: not extracted.",
        "Why: the public pages show only current open opportunities (RFQ no., entity, title, dates); awards and item details are not public.",
    ],
}


def main():
    wb = openpyxl.load_workbook(F)
    for name, lines in EXISTING.items():
        ws = wb[name]
        start = ws.max_row + 2
        if ws.max_row <= 2 and str(ws.cell(2, 1).value or "").startswith("(no rows"):
            ws.delete_rows(2)
            start = 3
        for i, t in enumerate(lines):
            c = ws.cell(start + i, 1, t)
            c.font = Font(bold=(i == 0))
            c.fill = NOTE_FILL
            c.alignment = Alignment(wrap_text=False)
    for name, lines in NEW_TABS.items():
        if name in wb.sheetnames:
            del wb[name]
        ws = wb.create_sheet(name[:31])
        for i, t in enumerate(lines, 1):
            c = ws.cell(i, 1, t)
            c.font = Font(bold=(i <= 2))
            c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions["A"].width = 160
    wb.save(F)
    print("notes added to", list(EXISTING), "| new tabs", list(NEW_TABS))


if __name__ == "__main__":
    main()
