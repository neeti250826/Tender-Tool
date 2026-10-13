"""
Build the TURMECA tender master workbook + refreshed Mapping_Taxonomy.xlsx.

Inputs (turmeca/sources/, harvested 2026-10-01 with Playwright):
  Taxonomy_v2.xlsx   Google Sheet <TAXONOMY_SHEET_ID> exported as xlsx
                     (tab "Combined Manufacturer List_Late": OU / Sub-OU / Product Category / Company / Family)
  nupco_parsed.json  all 119 tenders on nupco.com/tenders/tenders-list/1..5 (ID, status, dates, documents)
  cpv331.csv         every CPV 2008 code starting 331, from the EU Publications Office SPARQL endpoint
  okas_tr.json       OKAS code -> [EKAP Turkish name, item count] from 9,551 EKAP v2 tender detail JSONs

Outputs:
  <DOWNLOADS>/Mapping_Taxonomy.xlsx
  <WORKSPACE>/turmeca/TURMECA_Tender_Master.xlsx
"""
import csv, json, os, shutil, collections, datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "sources")
OUT_MASTER = os.path.join(HERE, "TURMECA_Tender_Master.xlsx")
OUT_TAX = r"<DOWNLOADS>/Mapping_Taxonomy.xlsx"
TODAY = "2026-10-01"

# ---------------------------------------------------------------- styling
NAVY = "1F4E78"
HDR = PatternFill("solid", fgColor=NAVY)
BAND = PatternFill("solid", fgColor="D9E1F2")
NEWFILL = PatternFill("solid", fgColor="C9DAF8")
REL_FILL = {"Core": "C6EFCE", "Partial": "FFEB9C", "Catch-all": "FCE4D6", "Review": "E7E6E6", "Out": "F2F2F2"}
WHITE_B = Font(bold=True, color="FFFFFF")
BOLD = Font(bold=True)
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")


def header(ws, row, cols, widths=None):
    for i, c in enumerate(cols, 1):
        x = ws.cell(row, i, c)
        x.font, x.fill, x.alignment, x.border = WHITE_B, HDR, Alignment(wrap_text=True, vertical="center"), BOX
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row + 1, 1)


def put_rows(ws, start, rows, rel_col=None):
    for r, vals in enumerate(rows, start):
        for c, v in enumerate(vals, 1):
            x = ws.cell(r, c, v)
            x.alignment, x.border = WRAP, BOX
        if rel_col:
            key = str(ws.cell(r, rel_col).value or "").split(" ")[0]
            if key in REL_FILL:
                ws.cell(r, rel_col).fill = PatternFill("solid", fgColor=REL_FILL[key])
    ws.auto_filter.ref = f"A{start-1}:{get_column_letter(ws.max_column)}{max(start, start+len(rows)-1)}"


# ---------------------------------------------------------------- taxonomy
ACM = "Acute Care & Monitoring"
CVS = "CardioVascular Surgery"
CST = "Cranial and Spinal Technologies"
ENT = "Ear Nose Throat"
EPT = "Electrophysiology Therapies (EPT)"
ENDO = "Endoscopy"
ICT = "Interventional Cardiology Therapies (ICT)"
NM = "Neuromodulation"
NV = "Neurovascular"
PH = "Pelvic Health"
PVH = "Peripheral Vascular Health"
SI = "Surgical Innovations"
SR = "Surgical Robotics"
OUS = [ACM, CVS, CST, ENT, EPT, ENDO, ICT, NM, NV, PH, PVH, SI, SR]

tw = openpyxl.load_workbook(os.path.join(SRC, "Taxonomy_v2.xlsx")).worksheets[0]
TAX = []  # OU, Sub-OU MDT, Sub-OU Clearstate, Product Category, Company, Family, Newly identified
for r in range(6, tw.max_row + 1):
    vals = [tw.cell(r, c).value for c in range(2, 8)]
    if not any(vals):
        continue
    new = tw.cell(r, 2).fill.fgColor.rgb == "FFC9DAF8"
    TAX.append([(str(v).strip() if v is not None else None) for v in vals] + ["Yes" if new else ""])
assert {t[0] for t in TAX} == set(OUS), {t[0] for t in TAX} ^ set(OUS)
SUBOUS = collections.OrderedDict()
for t in TAX:
    SUBOUS.setdefault(t[0], [])
    if t[2] not in SUBOUS[t[0]]:
        SUBOUS[t[0]].append(t[2])
SUB_TO_OU = collections.defaultdict(set)
for ou, subs in SUBOUS.items():
    for s in subs:
        SUB_TO_OU[s].add(ou)


def write_taxonomy():
    if os.path.exists(OUT_TAX):
        bak = OUT_TAX.replace(".xlsx", "_old_1Qg7_backup.xlsx")
        if not os.path.exists(bak):
            shutil.copy2(OUT_TAX, bak)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Taxonomy"
    header(ws, 1, ["Operating Unit", "Sub-OU (MDT provided)", "Sub-OU (Clearstate mapped)", "Product Category",
                   "Company", "Family", "Newly identified - TURMECA"], [30, 34, 34, 44, 28, 50, 14])
    for i, t in enumerate(TAX, 2):
        for c, v in enumerate(t, 1):
            ws.cell(i, c, v)
        if t[6]:
            for c in range(1, 8):
                ws.cell(i, c).fill = NEWFILL
    ws.auto_filter.ref = f"A1:G{len(TAX)+1}"

    def dim(name, cols, rows, widths):
        d = wb.create_sheet(name)
        header(d, 1, cols, widths)
        for i, r in enumerate(rows, 2):
            for c, v in enumerate(r, 1):
                d.cell(i, c, v)

    def uniq(seq):
        seen, out = set(), []
        for x in seq:
            if x not in seen:
                seen.add(x); out.append(x)
        return out

    dim("dim_overview", ["ou"], [[o] for o in uniq(t[0] for t in TAX)], [40])
    dim("dim_subview", ["sub_ou", "ou", "sub_ou_mdt_provided"],
        uniq((t[2], t[0], t[1]) for t in TAX), [50, 40, 50])
    dim("dim_product", ["product_category", "sub_ou", "ou"], uniq((t[3], t[2], t[0]) for t in TAX), [55, 45, 40])
    dim("dim_company", ["company"], [[c] for c in sorted(uniq(t[4] for t in TAX), key=str.lower)], [40])
    # dim_procedure keeps the column layout mapping/taxonomy.py reads (0=term, 1=canonical, 2=group, 10=category)
    proc = []
    for t in TAX:
        proc.append([t[5], t[5], t[5], None, None, None, None, None, "taxonomy_v2 Family", None, t[3],
                     None, None, None, t[4], t[2], t[0]])
    dim("dim_procedure", ["procedure", "Mapped Canonical procedure", "Mapped procedure group", "procedure_specific",
                          "Mapped procedure technique", "is_technique_variant", "is_archived_metric",
                          "mapping_confidence", "mapping_method", "notes", "Mapped procedure category",
                          "procedure_hierarchy_level", "procedure_parent", "procedure_sort_order",
                          "company", "sub_ou", "ou"], proc, [40, 20, 20] + [8] * 7 + [40, 8, 8, 8, 28, 40, 34])
    rd = wb.create_sheet("README", 0)
    lines = [
        ["Mapping_Taxonomy.xlsx", ""],
        ["Source", "Google Sheet <TAXONOMY_SHEET_ID>, tab 'Combined Manufacturer List_Late' "
                   "(Clearstate_Combined Manufacturer Product List), exported " + TODAY + " via Playwright."],
        ["Rows", f"{len(TAX)} rows; {len(SUBOUS)} OUs; {len({t[2] for t in TAX})} Sub-OUs (Clearstate mapped); "
                 f"{len({t[3] for t in TAX})} product categories; {len({t[4] for t in TAX})} companies; "
                 f"{len({t[5] for t in TAX})} families; {sum(1 for t in TAX if t[6])} rows flagged "
                 "'Newly identified - TURMECA' (blue in the source)."],
        ["Taxonomy", "Full flat list as in the source (OU > Sub-OU > Product Category > Company > Family)."],
        ["dim_* tabs", "Derived lookups, in the same tab names/columns mapping/taxonomy.py reads. "
                       "Family here is the product family/model line from the new sheet (previous file used "
                       "procedure groups), so dim_procedure carries Family in cols A-C and Product Category in col K."],
        ["Previous file", "The July 2026 export of sheet 1Qg7REHhlORq... was kept as Mapping_Taxonomy_old_1Qg7_backup.xlsx."],
    ]
    for i, (a, b) in enumerate(lines, 1):
        rd.cell(i, 1, a).font = BOLD
        rd.cell(i, 2, b).alignment = WRAP
    rd.column_dimensions["A"].width = 18
    rd.column_dimensions["B"].width = 120
    wb.save(OUT_TAX)


# ---------------------------------------------------------------- CPV / OKAS 331 scope
# relevance: Core | Partial | Catch-all | Out
# value: (relevance, [(OU, "Sub-OU; Sub-OU"), ...], note)
P = "Partial"
C = "Core"
K = "Catch-all"
CPV = {
    "33100000": (K, [(o, "all") for o in OUS], "Top of the 331 tree. Pull only with Turkish keyword screen."),
    "33110000": (P, [(CST, "Large Capital")], "Imaging parent; only intra-op imaging (O-arm class) is in scope."),
    "33111000": (P, [(CST, "Large Capital")], "Mobile/intra-op X-ray (O-arm) only."),
    "33111400": (P, [(CST, "Large Capital"), (SI, "Lung Health & Visualization")], "Fluoroscopy: O-arm, fluoroscopic navigation."),
    "33111710": (P, [(ICT, "PCI Non Stents"), (PVH, "Peripheral Vascular"), (NV, "Acute Ischemic Stroke; Hemorrhagic Stroke")],
                 "Angiography supplies: guiding/diagnostic catheters, guidewires, sheaths, microcatheters."),
    "33111730": (C, [(ICT, "PCI Non Stents"), (PVH, "Peripheral Vascular")], "Angioplasty supplies: PTCA/PTA balloons, DCB."),
    "33111740": (P, [(ICT, "PCI Non Stents"), (PVH, "Peripheral Vascular")], "Angioplasty devices (atherectomy systems, inflation devices)."),
    "33115000": (P, [(CST, "Large Capital")], "Tomography: intra-op cone-beam CT (O-arm) only."),
    "33115100": (P, [(CST, "Large Capital")], "CT scanners: intra-op CT only; diagnostic CT out."),
    "33120000": (K, [(EPT, "all"), (ACM, "Cerebral Compromise; Respiratory Compromise"), (CST, "Small Capital"), (ENT, "ENT")],
                 "Recording systems parent."),
    "33121000": (P, [(EPT, "Cardiovascular Diagnostics and Services")], "Long-term ambulatory recording: insertable cardiac monitors (ICM)."),
    "33121100": (P, [(ACM, "Cerebral Compromise")], "EEG class: LoC (BIS-type) monitors/sensors land here."),
    "33121300": (P, [(CST, "Small Capital"), (ENT, "ENT")], "EMG: intra-op nerve monitoring (NIM)."),
    "33121500": (P, [(EPT, "CAS - Other")], "ECG: EP recording systems."),
    "33123000": (K, [(EPT, "all"), (ICT, "all"), (CVS, "all")], "Cardiovascular devices parent."),
    "33123200": (P, [(EPT, "CAS - Other")], "Electrocardiography devices: EP recording / mapping systems."),
    "33123210": (P, [(EPT, "Cardiovascular Diagnostics and Services; Patient Management")], "Cardiac monitoring: ICM, remote monitoring."),
    "33123220": (P, [(ICT, "PCI Non Stents"), (EPT, "CAS - Other")], "Cardio-angiography devices."),
    "33125000": (P, [(PH, "Pelvic Health")], "Urology exploration: SNM/PTNM evaluation systems."),
    "33140000": (K, [(o, "all") for o in OUS], "Medical consumables parent."),
    "33141000": (K, [(o, "all") for o in OUS], "Largest EKAP class (108k items in prior pull). Keyword screen mandatory."),
    "33141100": (P, [(SI, "Wound Management; Instruments and Access")], "Dressings + clip/suture/ligature parent."),
    "33141120": (C, [(SI, "Wound Management; Instruments and Access; Stapling")], "Clips, sutures, ligatures."),
    "33141121": (C, [(SI, "Wound Management")], "Surgical sutures (conventional, barbed, antibacterial)."),
    "33141122": (C, [(SI, "Stapling; Hernia")], "Surgical staples, reloads, skin staplers, hernia tackers."),
    "33141125": (C, [(SI, "Wound Management")], "Suture material."),
    "33141126": (P, [(SI, "Instruments and Access")], "Ligatures / ligating clips."),
    "33141127": (C, [(SI, "Wound Management")], "Absorbable haemostats."),
    "33141128": (P, [(SI, "Wound Management")], "Suture needles."),
    "33141200": (P, [(ICT, "all"), (PVH, "all"), (NV, "all"), (EPT, "CAS - Cryo; CAS - Other; CAS - PFA"), (CVS, "ECT and Revasc; Aortic")],
                 "Catheters parent: shared with urinary/IV catheters."),
    "33141210": (C, [(ICT, "PCI Non Stents; Catheter Based Therapy"), (PVH, "Peripheral Vascular"),
                     (NV, "Acute Ischemic Stroke; Hemorrhagic Stroke"), (CVS, "Aortic"), (CST, "Total Spinal"), (EPT, "CAS - Cryo")],
                 "Balloon catheters: PTCA/PTA/DCB, BGC, occlusion balloons, stent-graft balloons, kyphoplasty, cryoballoon."),
    "33141220": (P, [(CVS, "ECT and Revasc")], "Cannulae: CPB/ECLS cannulae (exclude IV cannulae)."),
    "33141230": (P, [(ENT, "ENT"), (ICT, "all"), (PVH, "all")], "Dilators: sinus balloon dilation, vessel dilators/introducers."),
    "33141240": (P, [(ICT, "PCI Non Stents; Catheter Based Therapy"), (PVH, "Peripheral Vascular"), (NV, "all"),
                     (EPT, "CAS - Other"), (CVS, "Aortic")], "Catheter accessories: guidewires, guide extensions, sheaths."),
    "33141320": (P, [(ENDO, "GI - Hepatopancreatobiliary (HPB)"), (NM, "Neuromodulation - Pain"), (EPT, "CAS - Other")],
                 "Medical needles parent; only the biopsy/procedural sub-classes matter."),
    "33141323": (P, [(ENDO, "GI - Hepatopancreatobiliary (HPB)"), (SI, "Lung Health & Visualization")], "Biopsy needles: EUS-FNA/FNB, lung biopsy tools."),
    "33141326": (P, [(NM, "Neuromodulation - Pain"), (EPT, "CAS - Other"), (CST, "Total Spinal")],
                 "Radiology procedural needles: bone access, transseptal needles, ablation probes."),
    "33141328": (P, [(NM, "Neuromodulation - Pain")], "Epidural needles (SCS lead introducers)."),
    "33141620": (K, [(SI, "all"), (ACM, "all"), (CST, "CSF & CC"), (ENDO, "all")], "Medical kits: procedure packs; keyword screen."),
    "33141640": (P, [(CST, "CSF & CC")], "Drains: external ventricular/lumbar drainage (EDMS)."),
    "33141641": (P, [(ACM, "Temperature"), (ENDO, "GI - Esophageal")], "Probes: temperature probes, pH/manometry catheters."),
    "33141642": (P, [(CST, "CSF & CC")], "Drain accessories: EDMS."),
    "33141700": (P, [(CST, "Trauma; Total Spinal")], "Orthopaedic supplies parent."),
    "33141770": (C, [(CST, "Trauma")], "Fracture appliances, pins and plates."),
    "33150000": (P, [(NM, "all"), (PH, "Pelvic Health")], "Parent of the electrotherapy/stimulator branch only."),
    "33157000": (P, [(ACM, "Breathing Systems; Airway Management")], "Gas-therapy & respiratory devices parent."),
    "33157100": (P, [(ACM, "Breathing Systems")], "Medical gas masks."),
    "33157110": (P, [(ACM, "Breathing Systems")], "Oxygen masks."),
    "33157200": (P, [(ACM, "Breathing Systems")], "Oxygen kits (cannulae, tubing)."),
    "33157400": (P, [(ACM, "Breathing Systems; Airway Management")], "Breathing devices: circuits, HME filters, closed suction. Ventilators themselves are out."),
    "33157700": (P, [(ACM, "Breathing Systems")], "Blow bottles: incentive spirometers."),
    "33157800": (P, [(ACM, "Breathing Systems")], "Oxygen administration units: nebulisers, aerosol delivery."),
    "33157810": (P, [(ACM, "Breathing Systems")], "Oxygen therapy units."),
    "33158000": (P, [(NM, "all"), (PH, "Pelvic Health")], "Electrical/electromagnetic treatment parent."),
    "33158200": (P, [(NM, "all"), (PH, "Pelvic Health")], "Electrotherapy devices."),
    "33158210": (C, [(NM, "Neuromodulation - Pain; Neuromodulation - Brain Modulation"), (PH, "Pelvic Health"), (ENDO, "GI - Esophageal")],
                 "Stimulators: SCS, DBS, SNM, PTNM, gastric electrical stimulation."),
    "33160000": (K, [(SI, "all"), (CST, "all"), (SR, "all"), (ENT, "ENT")], "Operating techniques parent."),
    "33161000": (C, [(SI, "Electrosurgery; Vessel Sealing; Hardware - COT"), (CST, "Small Capital")],
                 "Electrosurgical units: generators, pencils, return pads, vessel sealing, Aquamantys-type energy."),
    "33162000": (K, [(SI, "all"), (CST, "Large Capital; Small Capital"), (SR, "Robotic Surgical Technologies")], "OR devices & instruments parent."),
    "33162100": (P, [(SI, "Hardware - COT; Electrosurgery"), (SR, "Robotic Surgical Technologies"), (CST, "Large Capital")],
                 "OR devices: robots, navigation, smoke evacuation."),
    "33162200": (P, [(SI, "Instruments and Access"), (CST, "Small Capital"), (ENT, "ENT")], "OR instruments."),
    "33164000": (C, [(SI, "Instruments and Access; Stapling"), (SR, "Robotic Surgical Technologies")], "Coelioscopy/laparoscopy: trocars, clip appliers, retrieval bags."),
    "33165000": (P, [(EPT, "CAS - Cryo"), (CVS, "Surgical Ablation & Valve Therapies - Surgical Ablation (SA)")], "Cryosurgery/cryoablation."),
    "33168000": (C, [(ENDO, "all"), (SI, "GYN Health; Lung Health & Visualization"), (ENT, "ENT")],
                 "Endoscopy/endosurgery devices: capsule, motility, hysteroscopic tissue removal, lung navigation."),
    "33168100": (P, [(SI, "GYN Health"), (ENT, "ENT"), (ENDO, "GI - Chronic and Colorectal")], "Endoscopes (hysteroscopes, sinuscopes, capsule)."),
    "33169000": (P, [(SI, "Instruments and Access"), (CST, "Small Capital"), (ENT, "ENT")], "Surgical instruments parent: PSI drills, microdebriders."),
    "33169100": (P, [(PVH, "EndoVenous")], "Surgical laser: endovenous laser (EVLA)."),
    "33169500": (C, [(CST, "Large Capital"), (ENT, "ENT"), (SI, "Lung Health & Visualization")], "Surgical tracking/navigation systems."),
    "33170000": (K, [(ACM, "all")], "Anaesthesia & resuscitation parent."),
    "33171000": (P, [(ACM, "Airway Management; Breathing Systems")], "Instruments for anaesthesia & resuscitation."),
    "33171100": (P, [(ACM, "Airway Management")], "Instruments for anaesthesia: ET tubes, stylets, videolaryngoscopes."),
    "33171110": (P, [(ACM, "Breathing Systems")], "Anaesthesia masks."),
    "33171200": (P, [(ACM, "Airway Management")], "Instruments for resuscitation."),
    "33171210": (P, [(ACM, "Breathing Systems")], "Resuscitation masks."),
    "33172000": (P, [(ACM, "Respiratory Compromise; Cerebral Compromise")], "Anaesthesia & resuscitation devices parent."),
    "33172100": (P, [(ACM, "Respiratory Compromise; Cerebral Compromise")], "Anaesthesia devices: capnography/LoC modules."),
    "33172200": (P, [(ACM, "Airway Management")], "Resuscitation devices."),
    "33180000": (K, [(EPT, "all"), (CVS, "all"), (ICT, "all"), (CST, "all"), (NM, "all")], "Functional support parent."),
    "33182000": (C, [(EPT, "all"), (CVS, "all"), (ICT, "Catheter Based Therapy")], "Cardiac support devices parent."),
    "33182100": (C, [(EPT, "Defibrillation Solutions")], "Defibrillators: ICD, CRT-D, EV-ICD. Exclude external AED/manual defibs."),
    "33182200": (C, [(EPT, "Cardiac Pacing Therapies; Defibrillation Solutions")], "Cardiac stimulation: IPG, CRT-P, leadless."),
    "33182210": (C, [(EPT, "Cardiac Pacing Therapies")], "Pacemakers."),
    "33182220": (C, [(CVS, "Surgical Ablation & Valve Therapies - Surgical Valve Therapies (SVT)"), (ICT, "Catheter Based Therapy")],
                 "Cardiac valves: mechanical/tissue valves, TAVI/TMVR."),
    "33182240": (C, [(EPT, "Cardiac Pacing Therapies; Defibrillation Solutions")], "Pacemaker parts/accessories: leads, envelopes."),
    "33182300": (C, [(CVS, "ECT and Revasc; Surgical Ablation & Valve Therapies - Surgical Ablation (SA)")],
                 "Cardiac surgery devices: positioners, surgical ablation."),
    "33183000": (C, [(CST, "Total Spinal; Trauma")], "Orthopaedic support devices parent."),
    "33183100": (C, [(CST, "Total Spinal; Trauma")], "Orthopaedic implants: pedicle screws, cages, nails."),
    "33183200": (P, [(CST, "Total Spinal")], "Orthopaedic prostheses: cervical artificial discs (joints out)."),
    "33183300": (C, [(CST, "Trauma; Total Spinal")], "Osteosynthesis: plates, screws, fixation."),
    "33184000": (K, [(CVS, "all"), (PVH, "all"), (ICT, "all"), (SI, "Hernia"), (NM, "all"), (CST, "CSF & CC")], "Artificial body parts parent."),
    "33184100": (P, [(SI, "Hernia"), (CST, "CSF & CC; Total Spinal"), (NM, "all"), (PH, "Pelvic Health")],
                 "Surgical implants: hernia mesh, shunts, neurostim implants."),
    "33184200": (C, [(CVS, "Aortic"), (PVH, "Peripheral Vascular; EndoVenous")], "Vascular prostheses: EVAR/TEVAR stent grafts, peripheral & venous stents."),
    "33184300": (C, [(CVS, "Surgical Ablation & Valve Therapies - Surgical Valve Therapies (SVT)"), (ICT, "Catheter Based Therapy")],
                 "Artificial heart parts: annuloplasty rings, TEER."),
    "33184500": (C, [(ICT, "PCI Stents")], "Coronary endoprostheses: DES."),
    "33185300": (P, [(ENT, "ENT")], "Otolaryngology implants: nasal splints, vocal-fold augmentation."),
    "33186000": (C, [(CVS, "ECT and Revasc")], "Extracorporeal circulation: heart-lung machines, ECLS."),
    "33186100": (C, [(CVS, "ECT and Revasc")], "Oxygenators."),
    "33186200": (P, [(ACM, "Temperature"), (CVS, "ECT and Revasc")], "Blood/fluid warming: patient warming lands here too."),
    "33190000": (K, [(o, "all") for o in OUS], "Miscellaneous medical devices parent."),
    "33194110": (P, [(NM, "Neuromodulation - Pain")], "Infusion pumps: implantable drug pumps (TDD) only."),
    "33195000": (C, [(ACM, "Respiratory Compromise; Cerebral Compromise")], "Patient monitoring parent."),
    "33195100": (C, [(ACM, "Respiratory Compromise; Cerebral Compromise")], "Monitors: pulse oximetry, capnography, NIRS, LoC."),
    "33195110": (C, [(ACM, "Respiratory Compromise")], "Respiratory monitors: capnography."),
    "33195200": (P, [(ACM, "Respiratory Compromise")], "Central monitoring stations."),
    "33197000": (P, [(EPT, "Patient Management"), (ENDO, "GI - Chronic and Colorectal")], "Medical computer equipment: remote monitoring, AI colonoscopy."),
}


def out_reason(code):
    c = code
    if c.startswith("3313") or c.startswith("331418") or c in ("33192400", "33192410", "33126000"):
        return "Dental"
    if c.startswith("33124") or c in ("33127000", "33159000", "33192500", "33141900") or c.startswith("331415"):
        return "IVD / laboratory / blood products"
    if c.startswith("3311") or c.startswith("33182400") or c == "33128000":
        return "Diagnostic imaging capital"
    if c.startswith("33181"):
        return "Renal / dialysis"
    if c.startswith("33191") or c.startswith("33198"):
        return "Sterilisation / hospital paper"
    if c.startswith("33192") or c.startswith("33193") or c.startswith("33196"):
        return "Furniture / mobility aids"
    if c.startswith("3315"):
        return "Radiotherapy / physiotherapy"
    if c.startswith("331411") or c.startswith("331414") or c.startswith("331413") or c.startswith("331416"):
        return "General consumables (dressings, syringes, needles, bags)"
    if c.startswith("33185") or c.startswith("33184"):
        return "Implant class without a Medtronic product (hearing, breast, eye, larynx)"
    if c == "33182230":
        return "Ventricular assist (Medtronic exited HVAD in 2021)"
    return "No product in the Medtronic taxonomy"


CPV_ROWS = []
with open(os.path.join(SRC, "cpv331.csv"), encoding="utf-8") as f:
    for r in csv.DictReader(f):
        CPV_ROWS.append((r["code"], r["en"]))
CPV_ROWS.sort()
CPV_CODES = {c for c, _ in CPV_ROWS}
missing = set(CPV) - CPV_CODES
assert not missing, missing
OKAS = json.load(open(os.path.join(SRC, "okas_tr.json"), encoding="utf-8"))

# ---------------------------------------------------------------- Turkish keyword screen per Sub-OU (for EKAP lot names)
TR_KW = {
    (ACM, "Airway Management"): "endotrakeal tüp; entübasyon tüpü; trakeostomi kanülü; stile; videolaringoskop",
    (ACM, "Breathing Systems"): "solunum devresi; anestezi devresi; HME filtre; bakteri viral filtre; kapalı aspirasyon; nebülizatör; oksijen maskesi; triflo / spirometre",
    (ACM, "Cerebral Compromise"): "BIS sensörü; bilinç düzeyi; anestezi derinliği; NIRS; serebral oksimetre",
    (ACM, "Respiratory Compromise"): "pulse oksimetre; SpO2 probu / sensörü; kapnografi; EtCO2; nellcor",
    (ACM, "Temperature"): "hasta ısıtma; ısıtma battaniyesi; sıcak hava ısıtma; ısı probu",
    (CVS, "Aortic"): "stent greft; aort stent greft; EVAR; TEVAR; endoanchor",
    (CVS, "ECT and Revasc"): "oksijenatör; kalp akciğer makinesi; perfüzyon seti; kanül (arteriyel / venöz / kardiyopleji); ECMO; otransfüzyon; stabilizatör",
    (CVS, "Surgical Ablation & Valve Therapies - Surgical Ablation (SA)"): "cerrahi ablasyon; kriyoablasyon probu; RF ablasyon (cerrahi)",
    (CVS, "Surgical Ablation & Valve Therapies - Surgical Valve Therapies (SVT)"): "kalp kapağı; mekanik kalp kapağı; biyolojik kapak; anüloplasti halkası",
    (CST, "Total Spinal"): "pedikül vidası; transpediküler; kafes / cage; interbody; servikal plak; kifoplasti; vertebroplasti; kemik grefti; DBM; servikal disk protezi",
    (CST, "CSF & CC"): "şant; hidrosefali; ventriküloperitoneal; eksternal ventriküler drenaj; lomber drenaj; kranial fiksasyon",
    (CST, "Large Capital"): "nöronavigasyon; navigasyon sistemi; O-arm; intraoperatif görüntüleme; spinal robot",
    (CST, "Small Capital"): "nöromonitörizasyon; sinir monitörü; motorlu cerrahi el aleti; yüksek devirli drill; craniotome; bipolar enerji",
    (CST, "Trauma"): "intramedüller çivi; kilitli çivi; plak vida; kanüllü vida; eksternal fiksatör",
    (ENT, "ENT"): "sinüs balonu; mikrodebrider; KBB drill; burun tamponu; nazal splint; sinir monitörü (KBB); KBB navigasyon",
    (EPT, "Cardiac Pacing Therapies"): "kalp pili; pacemaker; tek / çift odacıklı; CRT-P; leadless; elektrot / lead",
    (EPT, "Cardiovascular Diagnostics and Services"): "implante edilebilir kardiyak monitör; loop recorder; ICM",
    (EPT, "CAS - Cryo"): "kriyobalon; kriyoablasyon kateteri; kriyo konsol",
    (EPT, "CAS - Other"): "ablasyon kateteri; elektrofizyoloji kateteri; tanısal EP kateteri; haritalama; transseptal iğne",
    (EPT, "CAS - PFA"): "pulsed field ablasyon; PFA kateteri",
    (EPT, "Defibrillation Solutions"): "ICD; implante edilebilir defibrilatör; CRT-D; antibakteriyel zarf",
    (EPT, "Patient Management"): "uzaktan izlem; remote monitoring",
    (ENDO, "GI - Chronic and Colorectal"): "kapsül endoskopi; endoskopik hemostaz; klip (endoskopik)",
    (ENDO, "GI - Esophageal"): "özofagus manometri; pH metre; empedans; Barrett ablasyon; gastrik stimülatör",
    (ENDO, "GI - Hepatopancreatobiliary (HPB)"): "EUS iğnesi; FNA / FNB iğnesi; mikrodalga ablasyon; RF ablasyon iğnesi",
    (ICT, "Catheter Based Therapy"): "TAVI; transkateter aort kapak; perkütan kapak; mitral klip / TEER",
    (ICT, "PCI Non Stents"): "PTCA balon; koroner balon; ilaç kaplı balon; kılavuz kateter; kılavuz tel; guide extension",
    (ICT, "PCI Stents"): "koroner stent; ilaç salınımlı stent; DES",
    (ICT, "Renal Denervation (RDN)"): "renal denervasyon",
    (NM, "GI - Hepatopancreatobiliary (HPB)"): "mikrodalga ablasyon jeneratörü; RF ablasyon sistemi",
    (NM, "Neuromodulation - Brain Modulation"): "derin beyin stimülasyonu; DBS; beyin pili; mikroelektrot kayıt",
    (NM, "Neuromodulation - Pain"): "spinal kord stimülatörü; nöromodülasyon; ağrı pompası; intratekal pompa; RF ablasyon (kemik); kemik çimentosu",
    (NV, "Acute Ischemic Stroke"): "stent retriever; trombektomi; aspirasyon kateteri; balon kılavuz kateter",
    (NV, "Hemorrhagic Stroke"): "anevrizma koili; akım yönlendirici / flow diverter; mikrokateter; sıvı embolizan; intrasakküler",
    (PH, "Pelvic Health"): "sakral nöromodülasyon; tibial sinir stimülasyonu; PTNS",
    (PVH, "EndoVenous"): "varis; endovenöz lazer; radyofrekans ablasyon (venöz); siyanoakrilat / venöz yapıştırıcı; venöz stent; embolizasyon koili",
    (PVH, "Peripheral Vascular"): "PTA balonu; periferik balon; periferik stent; aterektomi; ilaç kaplı balon; emboli koruma; CTO kateteri",
    (SI, "Electrosurgery"): "elektrokoter; koter kalemi; nötr elektrot / hasta plağı; argon; duman tahliye",
    (SI, "GYN Health"): "histeroskop; histeroskopik doku çıkarma; sıvı yönetim sistemi",
    (SI, "Hardware - COT"): "elektrocerrahi ünitesi; jeneratör",
    (SI, "Hernia"): "fıtık yaması; mesh; prolen yama; tacker / zımba (fıtık)",
    (SI, "Instruments and Access"): "trokar; klip aplikatör; endoklip; doku torbası; aspirasyon irrigasyon; yara retraktörü",
    (SI, "Lung Health & Visualization"): "akciğer navigasyon; elektromanyetik navigasyon bronkoskopi",
    (SI, "Stapling"): "stapler; endostapler; kartuş; sirküler stapler; deri stapleri",
    (SI, "Vessel Sealing"): "damar mühürleme; ligasure; ultrasonik makas / sonicision; bipolar damar mühürleme",
    (SI, "Wound Management"): "sütür; cerrahi iplik; emilebilir sütür; barbed sütür; hemostatik; doku yapıştırıcı",
    (SR, "Robotic Surgical Technologies"): "robotik cerrahi; robot kolu; robotik alet; steril örtü (robot)",
}
for (ou, sub) in TR_KW:
    assert sub in SUBOUS[ou], (ou, sub)

# ---------------------------------------------------------------- NUPCO tender families
NUPCO_EN = {
    "NDP0846/26": "Medical supply items for health clusters", "NDP0848/26": "Drug items for health clusters",
    "NDP0849/26": "Medical supply items for health clusters", "NDP0852/26": "Medical devices - State Security Presidency (lab equipment)",
    "NDP0873/26": "Drug items for Ministry of Health", "NDP0875/26": "ENT supplies",
    "NDP0881-26": "INUPCO - drug items for e-marketplace", "NDP0883/26": "SRM - medical supplies for e-marketplace",
    "NDP0885/26": "Clinical nutrition programme devices - KSAU-HS Riyadh", "NDP0886/26": "INUPCO - drug items for e-marketplace",
    "NDP0887/26": "INUPCO - drug items for e-marketplace", "NDP0888/26": "Telemedicine devices - Madinah Health Cluster",
    "NDP0889/26": "Ophthalmology & dermatology supplies", "NDP0890/26": "Physiotherapy supplies",
    "NDP0891/26": "Sterilisation & endoscopy supplies", "NDP0892/26": "General medical supplies",
    "NDP0893/26": "INUPCO - drug items for e-marketplace", "NDP0895/26": "Urgent general medical supplies",
    "NDP0896/26": "Urgent general medical supplies", "NDP0897/26": "INUPCO - medical supplies for e-marketplace",
    "NDP0903/26": "Laboratory testing via reference laboratories",
    "NPT0001/23": "Lab devices with reagents & consumables, fixed price per confirmed result (GPPRR) - phase 2",
    "NPT0001/26": "Open framework - dialysis & artificial kidney supplies",
    "NPT0002/25": "Open framework - mortuary & morgue refrigerators",
    "NPT0002/26": "Open framework - respiratory therapy & anaesthesia supplies",
    "NPT0003/26": "Open framework - custom rehabilitation & prosthetics",
    "NPT0004/23": "Addendum - framework supply & installation of medical devices",
    "NPT0004/26": "Framework - prosthetic limb supplies",
    "NPT0005/26": "Open framework - orthopaedic & spine supplies",
    "NPT0006/26": "Open framework - specialty surgery supplies",
    "NPT0007/26": "Open framework - IV solutions & parenteral nutrition",
    "NPT0008/25": "Open framework - chemicals & pharmaceutical preparations",
    "NPT0008/26": "Open framework - therapeutic nutrition items",
    "NPT0009/25": "Open framework - laser devices & accessories",
    "NPT0009/26": "Open framework - narcotic & controlled drugs",
    "NPT0010/23": "Framework - supply & installation of medical imaging devices",
    "NPT0010/26": "Framework - infertility & IVF supplies",
    "NPT0011/26": "Open framework - Wasfaty service drugs",
    "NPT0012/26": "Open framework - pharmacy department",
    "NPT0013/23": "Framework - specialty laboratory devices",
    "NPT0013/26": "Open framework - radiology department",
    "NPT0014/23": "Framework - medical rehabilitation devices",
    "NPT0014/26": "Framework - histopathology supplies & reagents",
    "NPT0015/25": "Open framework - vitamins & minerals",
    "NPT0015/26": "Open framework - medical rehabilitation department",
    "NPT0016/25": "Open framework - radioisotopes & contrast media",
    "NPT0016/26": "Open framework - surgical sutures, laparoscopic surgery & robotic surgery supplies",
    "NPT0017/25": "Open framework - hormones & GI drugs",
    "NPT0018/25": "Open framework - blood products & derivatives",
    "NPT0018/26": "Open framework - physiotherapy & occupational therapy supplies",
    "NPT0019/23": "Framework supply & installation of medical devices - addendum 2",
    "NPT0019/25": "Open framework - dental department",
    "NPT0019/26": "Open framework - sterilisation supplies",
    "NPT0020/23": "Framework - general laboratory devices",
    "NPT0020/25": "Neuroscience supplies",
    "NPT0020/26": "Open framework - specialty laboratories department",
    "NPT0021/24": "Open framework - intensive care & monitoring department",
    "NPT0021/25": "Medical rehabilitation services",
    "NPT0021/26": "Open framework - general laboratory department",
    "NPT0022/24": "Open framework - general departments, calibration & simulation",
    "NPT0022/25": "Open framework - respiratory & sensory health (drugs)",
    "NPT0022/26": "Open framework - microbiology & molecular lab supplies",
    "NPT0023/25": "Open framework - refrigerators & freezers",
    "NPT0023/26": "Open framework - genetics / molecular / cytogenetics lab supplies",
    "NPT0024/25": "Open framework - ENT department",
    "NPT0024/26": "Open framework - general medical supplies (nursing & wound care)",
    "NPT0025/26": "Open framework - chemistry, immunochemistry & point-of-care supplies",
    "NPT0026/24": "Open framework - operating rooms & surgical drills",
    "NPT0026/24-1": "Update to the OR & surgical drills framework - phase 1",
    "NPT0026/26": "Open framework - immunology, HLA & flow cytometry supplies",
    "NPT0027/24": "Open framework - medical furniture",
    "NPT0027/25": "Open framework - internal medicine & chronic disease items (drugs)",
    "NPT0027/26": "Open framework - general lab supplies & reagents",
    "NPT0028/25": "Open framework - sera & vaccines",
    "NPT0028/26": "Open framework - maxillofacial surgery & dental implant supplies",
    "NPT0029/25": "Open framework - sterilisation & infection control department",
    "NPT0029/26": "Open framework - GI endoscopy supplies",
    "NPT0030/25": "Open framework - ophthalmology department",
    "NPT0030/26": "Open framework - toxicology / biochemical genetics / newborn screening lab supplies",
    "NPT0031/25": "Open framework - disinfectant supplies",
    "NPT0031/26": "Open framework - blood bank / stem cells / haematology / coagulation supplies",
    "NPT0032/26": "Open framework - general dental supplies",
    "NPT0033/24": "Open framework - ambulance outfitting & supply",
    "NPT0033/25": "ENT supplies",
    "NPT0034/24": "Open framework - home care department",
    "NPT0034/25": "Open framework - infectious disease drugs",
    "NPT0035/25": "General medical supplies (nursing & wound care)",
    "NPT0036/26": "Open framework - robotic surgery on pay-per-use",
    "NPT0037/24": "Drugs addendum",
    "NPT0037/25": "General surgical supplies (urology & surgery) - supplementary",
    "NPT0038/25": "Open framework - interventional radiology & cardiac catheterisation supplies",
    "NPT0039/23": "Lab devices with reagents & consumables, fixed price per confirmed result (GPPRR) - phase 3",
    "NPT0041/25": "Open framework - critical care & neuro/psychiatric (drugs)",
    "NPT0042/23": "Home dialysis service",
    "NPT0042/25": "Open framework - cardiac & vascular surgery supplies",
    "NPT0042/26": "Open framework - long-term care services",
    "NPT0043/25": "Dental supplies (lab, implant, maxillofacial) - supplementary",
    "NPT0043/26": "Medical waste management services",
    "NPT0044/25": "General supplies (sutures, GI endoscopy, robotic surgery) - supplementary",
    "NPT0045/25": "Open framework - immunology & rare disease drugs",
    "NPT0046/24": "Open framework - infusion pumps & consumables",
    "NPT0047/24": "Open framework - diabetes department & consumables",
    "NPT0047/25": "Open framework - oncology drugs",
    "NPT0048/24": "Manufacturer pre-qualification - value-based agreement",
    "NPT0049/24": "RFI - outsourcing of medical radiology services",
    "NPT0049/25": "Open framework - ophthalmology supplies",
    "NPT0050/24": "Lab devices with reagents & consumables, fixed price per confirmed result (GPPRR) - phase 1",
    "NPT0050/25": "Wasfaty drugs addendum",
    "NPT0052/24": "Open framework - medical devices addendum 1",
    "NPT0054/25": "Home healthcare services",
    "NPT0055/24": "Long-stay patient healthcare services",
    "NPT0055/25": "Medical imaging (radiology) services",
    "NPT0056/24": "Open framework - medical devices addendum 2",
    "NPT0059/25": "Non-emergency medical transport services",
    "NPT0060/25": "Radiotherapy services",
    "NPT0062/25": "Catering services in health facilities",
    "NPT0063/25": "Medical waste management services",
    "NPT0064/25": "Doping-detection laboratory supply & fit-out",
    "": "Organ & tissue transport service",
}
GEN = "General supplies - screen items"
NUPCO_SCOPE = {  # id -> (relevance, [(OU, sub-OUs)], note)
    "NPT0016/26": (C, [(SI, "Wound Management; Stapling; Instruments and Access; Hernia; Vessel Sealing"), (SR, "Robotic Surgical Technologies")], "Successor of NPT0044/25."),
    "NPT0044/25": (C, [(SI, "Wound Management; Stapling; Instruments and Access"), (ENDO, "all"), (SR, "Robotic Surgical Technologies")], "Final results on site."),
    "NPT0038/25": (C, [(ICT, "PCI Non Stents; PCI Stents; Catheter Based Therapy"), (PVH, "Peripheral Vascular; EndoVenous"),
                       (NV, "Acute Ischemic Stroke; Hemorrhagic Stroke"), (EPT, "CAS - Other; CAS - Cryo; CAS - PFA"), (CVS, "Aortic")], "Cath lab / IR framework."),
    "NPT0042/25": (C, [(CVS, "all"), (PVH, "Peripheral Vascular")], "Under study - no award yet."),
    "NPT0033/25": (C, [(ENT, "ENT")], ""),
    "NPT0024/25": (C, [(ENT, "ENT")], "ENT department (capital + consumables)."),
    "NDP0875/26": (C, [(ENT, "ENT")], "Direct purchase."),
    "NPT0020/25": (C, [(CST, "CSF & CC; Small Capital; Large Capital"), (NM, "Neuromodulation - Pain; Neuromodulation - Brain Modulation"), (NV, "all")],
                   "Item list: EEG/IONM electrodes, stimulators, shunts, drills, navigation."),
    "NPT0026/24": (C, [(CST, "Small Capital; Large Capital"), (ENT, "ENT"), (SI, "Hardware - COT")], "OR & surgical drills."),
    "NPT0026/24-1": (C, [(CST, "Small Capital; Large Capital"), (ENT, "ENT"), (SI, "Hardware - COT")], "Update to NPT0026/24."),
    "NPT0005/26": (C, [(CST, "Total Spinal; Trauma")], ""),
    "NPT0021/24": (C, [(ACM, "Respiratory Compromise; Cerebral Compromise; Temperature")], "ICU & monitoring."),
    "NPT0002/26": (C, [(ACM, "Airway Management; Breathing Systems")], ""),
    "NPT0029/26": (C, [(ENDO, "all")], ""),
    "NPT0036/26": (C, [(SR, "Robotic Surgical Technologies")], "Pay-per-use - capital may not appear as a line price."),
    "NPT0006/26": (P, [(SI, "Stapling; Instruments and Access; Electrosurgery; Wound Management")], "Item list: packs, trocars, staplers, electrodes."),
    "NPT0037/25": (P, [(SI, "all"), (PH, "Pelvic Health")], "Urology & surgery."),
    "NPT0024/26": (P, [(SI, "Wound Management"), (ACM, "Breathing Systems; Temperature")], GEN),
    "NPT0035/25": (P, [(SI, "Wound Management"), (ACM, "Breathing Systems")], GEN),
    "NPT0034/24": (P, [(ACM, "Breathing Systems; Respiratory Compromise")], "Home care: ventilators, circuits, masks, suction."),
    "NPT0052/24": (P, [(ACM, "Respiratory Compromise"), (CST, "Large Capital")], "Device addendum: monitors, ventilators, electrodes."),
    "NPT0056/24": (P, [(ACM, "Respiratory Compromise"), (CVS, "ECT and Revasc")], "Device addendum 2 (incl. BiVAD, blood-gas monitor)."),
    "NPT0004/23": (P, [(ACM, "all"), (CST, "Large Capital")], "2023 device framework - before the 2024-26 window."),
    "NPT0019/23": (P, [(ACM, "all"), (CST, "Large Capital")], "2023 device framework - before the 2024-26 window."),
    "NDP0891/26": (P, [(SI, "Instruments and Access; Wound Management"), (ENDO, "all")], "Item list: trocars, sutures, gloves."),
    "NDP0846/26": (P, [(SI, "all"), (ACM, "all")], GEN), "NDP0849/26": (P, [(SI, "all"), (ACM, "all")], GEN),
    "NDP0883/26": (P, [(SI, "all"), (ACM, "all")], GEN), "NDP0892/26": (P, [(SI, "all"), (ACM, "all")], GEN),
    "NDP0895/26": (P, [(SI, "all"), (ACM, "all")], GEN), "NDP0896/26": (P, [(SI, "all"), (ACM, "all")], GEN),
    "NDP0897/26": (P, [(SI, "all"), (ACM, "all")], GEN),
    "NPT0048/24": ("Review", [], "Manufacturer pre-qualification, no product lines. Check whether Medtronic qualified."),
}
OUT_CAT = {
    "Pharmaceuticals": ["NDP0848/26", "NDP0873/26", "NDP0881-26", "NDP0886/26", "NDP0887/26", "NDP0893/26", "NPT0007/26", "NPT0008/25",
                        "NPT0009/26", "NPT0011/26", "NPT0012/26", "NPT0015/25", "NPT0016/25", "NPT0017/25", "NPT0018/25", "NPT0022/25",
                        "NPT0027/25", "NPT0028/25", "NPT0034/25", "NPT0037/24", "NPT0041/25", "NPT0045/25", "NPT0047/25", "NPT0050/25",
                        "NPT0008/26"],
    "Laboratory / IVD": ["NDP0852/26", "NDP0903/26", "NPT0001/23", "NPT0013/23", "NPT0014/26", "NPT0020/23", "NPT0020/26", "NPT0021/26",
                         "NPT0022/26", "NPT0023/26", "NPT0025/26", "NPT0026/26", "NPT0027/26", "NPT0030/26", "NPT0031/26", "NPT0039/23",
                         "NPT0050/24", "NPT0064/25"],
    "Services": ["", "NPT0021/25", "NPT0042/23", "NPT0042/26", "NPT0043/26", "NPT0049/24", "NPT0054/25", "NPT0055/24", "NPT0055/25",
                 "NPT0059/25", "NPT0060/25", "NPT0062/25", "NPT0063/25"],
    "Dental / maxillofacial": ["NPT0019/25", "NPT0028/26", "NPT0032/26", "NPT0043/25", "NPT0009/25"],
    "Ophthalmology / dermatology": ["NDP0889/26", "NPT0030/25", "NPT0049/25"],
    "Rehabilitation / prosthetics / physiotherapy": ["NDP0890/26", "NPT0003/26", "NPT0004/26", "NPT0014/23", "NPT0015/26", "NPT0018/26"],
    "Imaging / radiology": ["NPT0010/23", "NPT0013/26"],
    "Dialysis / IVF / nutrition devices": ["NPT0001/26", "NPT0010/26", "NDP0885/26"],
    "Sterilisation / infection control": ["NPT0019/26", "NPT0029/25", "NPT0031/25"],
    "Furniture / equipment / vehicles": ["NPT0002/25", "NPT0022/24", "NPT0023/25", "NPT0027/24", "NPT0033/24", "NDP0888/26"],
    "Infusion pumps / diabetes (no matching OU in taxonomy)": ["NPT0046/24", "NPT0047/24"],
}
for cat, ids in OUT_CAT.items():
    for i in ids:
        assert i not in NUPCO_SCOPE, i
        NUPCO_SCOPE[i] = ("Out", [], cat)
NUPCO = json.load(open(os.path.join(SRC, "nupco_parsed.json"), encoding="utf-8"))
for n in NUPCO:
    assert n["id"] in NUPCO_SCOPE, n["id"]
    assert n["id"] in NUPCO_EN, n["id"]


def nupco_stage(status):
    s = status.lower()
    if "cancel" in s:
        return "n/a - cancelled"
    if "final" in s:
        return "Awarded"
    return "Pipeline"


def expand(pairs):
    """[(OU, 'sub; sub' | 'all')] -> [(OU, sub)]"""
    out = []
    for ou, subs in pairs:
        if subs == "all":
            out += [(ou, s) for s in SUBOUS[ou]]
        else:
            for s in subs.split("; "):
                assert s in SUBOUS[ou], (ou, s)
                out.append((ou, s))
    return out


# ---------------------------------------------------------------- master table definition (from the "Tenders" doc)
MASTER = [
    ("Source", ["Row ID", "Country", "Portal", "Source Type", "Source URL"]),
    ("Tender", ["Tender ID", "Tender Title", "Tender Type", "Procurement Authority", "Hospital/Buyer"]),
    ("Dates", ["Publication Date", "Opening Date", "Award/Contract Date", "Result Date", "Award Year"]),
    ("Original product", ["Original-language Description", "English Description", "Lot/Item No.", "CPV/OKAS/NUPCO Code"]),
    ("Taxonomy", ["OU", "Sub-OU", "Device Category", "Company", "Family Name", "Disclosed Brand"]),
    ("Commercial", ["Supplier/Distributor", "Manufacturer", "Volume", "UOM", "Currency", "Unit Award Price", "Total Award Value"]),
    ("Normalized", ["FX Date", "LCY per USD", "Value USD", "HBP USD"]),
    ("Evidence", ["Tender Stage", "Price Evidence", "Manufacturer Evidence", "Winner Disclosed?", "Manufacturer Explicit?",
                  "Model Explicit?", "Award vs Budget?", "Confidence"]),
    ("QC", ["Included/Review/Excluded", "Exclusion Reason", "Duplicate/Retender Flag", "Caveat"]),
]
NEW_COLS = {"Tender Stage", "Price Evidence", "Manufacturer Evidence"}
ADDED = {"Row ID", "Award Year"}
DICT = {
    "Row ID": ("Unique row key: <Country ISO2>-<Portal>-<Tender ID>-<Lot/Item No.>. Keeps tender-row traceability.", "Text", "Added (ACEE SOP traceability)"),
    "Country": ("Country of the buyer.", "KSA / Turkey / Pakistan / Kuwait / UAE", "Doc"),
    "Portal": ("Portal the row was harvested from.", "NUPCO / Etimad / EKAP / EKAP legacy / EPADS / PPRA / KP DGHS / Punjab DGHS / Sindh PPRA / CAPT / MoF DPP / ADGPG / Dubai eSupply", "Doc"),
    "Source Type": ("Document type the row comes from.", "Item list / Preliminary result / Final result / Award list / Contract / BER / Notice", "Doc"),
    "Source URL": ("Direct URL of the page or PDF the row was read from.", "URL", "Doc"),
    "Tender ID": ("Portal tender number (NUPCO NPT/NDP, EKAP İKN, CAPT tender no., EPADS id).", "Text", "Doc"),
    "Tender Title": ("Title as published, English translation if not English.", "Text", "Doc"),
    "Tender Type": ("Procedure: open framework, direct purchase, open tender, RFQ, etc.", "Text", "Doc"),
    "Procurement Authority": ("Body running the procurement (NUPCO, CAPT, Ministry, KİK idare).", "Text", "Doc"),
    "Hospital/Buyer": ("End buyer / health facility if disclosed.", "Text", "Doc"),
    "Publication Date": ("Notice publication date.", "Date", "Doc"),
    "Opening Date": ("Bid opening / tender opening date.", "Date", "Doc"),
    "Award/Contract Date": ("Award date, or contract signature date where published (EKAP sözleşme tarihi).", "Date", "Doc"),
    "Result Date": ("Date the final result was published.", "Date", "Doc"),
    "Award Year": ("Calendar year used for the market view: Award/Contract Date year, else Result Date year. Never the tender-ID year.", "Formula", "Added (doc: keep 2024/25/26 separate)"),
    "Original-language Description": ("Item / lot description exactly as published.", "Text", "Doc"),
    "English Description": ("English translation of the item description.", "Text", "Doc"),
    "Lot/Item No.": ("Lot (kısım) or item line number within the tender.", "Text", "Doc"),
    "CPV/OKAS/NUPCO Code": ("Item-level class code: EKAP OKAS code, NUPCO item code, or other portal code.", "Text", "Doc"),
    "OU": ("Operating Unit from Mapping_Taxonomy.xlsx.", "List", "Doc"),
    "Sub-OU": ("Sub-OU (Clearstate mapped) from Mapping_Taxonomy.xlsx.", "List", "Doc"),
    "Device Category": ("Product Category from Mapping_Taxonomy.xlsx.", "Text", "Doc"),
    "Company": ("Manufacturer company as named in the taxonomy (OEM, not distributor).", "Text", "Doc"),
    "Family Name": ("Product family / model line from the taxonomy.", "Text", "Doc"),
    "Disclosed Brand": ("Brand or model text exactly as disclosed in the source.", "Text", "Doc"),
    "Supplier/Distributor": ("Winning bidder / contractor as named in the source. Not assumed to be the manufacturer.", "Text", "Doc"),
    "Manufacturer": ("Manufacturer, from the source or inferred (see Manufacturer Evidence).", "Text", "Doc"),
    "Volume": ("Awarded quantity (or requested quantity if that is all the source has - flag in Caveat).", "Number", "Doc"),
    "UOM": ("Unit of measure.", "Text", "Doc"),
    "Currency": ("Local currency code.", "SAR / TRY / PKR / KWD / AED", "Doc"),
    "Unit Award Price": ("Awarded unit price in local currency. Fill only when Price Evidence = Line price.", "Number", "Doc"),
    "Total Award Value": ("Award / contract value in local currency at the level the source gives it.", "Number", "Doc"),
    "FX Date": ("Date of the FX rate used.", "Date", "Doc"),
    "LCY per USD": ("Local currency units per 1 USD on FX Date.", "Number", "Doc"),
    "Value USD": ("Total Award Value / LCY per USD.", "Formula", "Doc"),
    "HBP USD": ("Hospital buying price per unit in USD = Unit Award Price / LCY per USD. Blank unless Price Evidence = Line price.", "Formula", "Doc"),
    "Tender Stage": ("Pipeline = notice, under study, item list or preliminary result only. Awarded = final award/result published with winner. "
                     "Contracted = signed contract with contract date and value (e.g. EKAP sözleşme).", "Pipeline / Awarded / Contracted", "NEW"),
    "Price Evidence": ("Line price = a unit price or value attributable to this single item line. Contract-level only = only a tender, lot or "
                       "contract total covering several items. None = no price published.", "Line price / Contract-level only / None", "NEW"),
    "Manufacturer Evidence": ("Explicit = the source names manufacturer / brand / model for the line. Inferred from distributor = derived from a "
                              "supplier -> OEM mapping (supplier is a distributor). None = no basis to name the manufacturer.",
                              "Explicit / Inferred from distributor / None", "NEW"),
    "Winner Disclosed?": ("Does the source name the winning supplier?", "Yes / No", "Doc"),
    "Manufacturer Explicit?": ("Derived: Yes when Manufacturer Evidence = Explicit.", "Formula", "Doc (now derived)"),
    "Model Explicit?": ("Does the source name the model / catalogue number?", "Yes / No", "Doc"),
    "Award vs Budget?": ("Is a budget / approximate cost (EKAP yaklaşık maliyet) published alongside the award?", "Yes / No", "Doc"),
    "Confidence": ("Overall confidence in the taxonomy mapping of the row.", "High / Medium / Low", "Doc"),
    "Included/Review/Excluded": ("QC decision for market sizing.", "Included / Review / Excluded", "Doc"),
    "Exclusion Reason": ("Why the row is excluded or held for review.", "Text", "Doc"),
    "Duplicate/Retender Flag": ("Duplicate of another row, or a re-tender of a cancelled tender.", "Duplicate / Retender / blank", "Doc"),
    "Caveat": ("Free-text caveat (e.g. framework multi-award, requested qty used as volume).", "Text", "Doc"),
}

EVIDENCE_RULES = [
    # country, portal, record type, stage, price, manufacturer, basis
    ("KSA", "NUPCO", "NPT tender - Available/New or Under Studying (item list only)", "Pipeline", "None", "None",
     "Item List has item code, description, UOM, INITIAL QTY; no supplier."),
    ("KSA", "NUPCO", "NPT tender - Initial Results (Preliminary Result PDF)", "Pipeline", "None", "Explicit (where the PDF lists it)",
     "Preliminary result is not final; keep out of realised market value."),
    ("KSA", "NUPCO", "NPT tender - Final Results (Final Result PDF)", "Awarded", "None", "Explicit",
     "Final results list vendor, manufacturer, model/catalogue no., country; no price or value anywhere. Framework: several winners per item."),
    ("KSA", "NUPCO", "NDP direct purchase", "Pipeline", "None", "None", "Only item list + T&C published; no result documents."),
    ("KSA", "Etimad", "Award announced (TenderCategory=6)", "Awarded", "Contract-level only", "Inferred from distributor (if mapped) / None",
     "Tender-level award value + supplier; no item lines."),
    ("Turkey", "EKAP v2", "Contract record (sozlesmeBilgiList), lot with several items", "Contracted", "Contract-level only",
     "Inferred from distributor (if mapped) / None", "sozlesmeBedeli is per contract; do not split across items."),
    ("Turkey", "EKAP v2", "Contract record, lot = single item", "Contracted", "Line price", "Inferred from distributor (if mapped) / None",
     "Lot value is attributable to one item; unit price = value / qty only if qty is published."),
    ("Turkey", "EKAP v2 / legacy", "Notice without result", "Pipeline", "None", "None", ""),
    ("Turkey", "EKAP legacy (BirBakistaIhale)", "Result page with contract date", "Contracted", "Contract-level only", "None",
     "Winner, contract value, approx. cost, hi/lo bid; lot items without qty/price/manufacturer."),
    ("Pakistan", "KP DGHS / healthkp", "Bid Evaluation Report (BER)", "Awarded", "Line price", "Explicit",
     "Item, make, model, origin, unit price; no quantity."),
    ("Pakistan", "PPRA (legacy contract_org)", "Contract >= Rs50M", "Contracted", "Contract-level only", "None", ""),
    ("Pakistan", "EPADS / EPMS", "Notice / bidding documents", "Pipeline", "None", "None", "Attached evaluations are often scans."),
    ("Kuwait", "CAPT", "Winning bids list", "Awarded", "Contract-level only", "None", "Tender-level only; item popup has price but no description."),
    ("UAE", "MoF DPP / ADGPG / Dubai eSupply", "Notice / opportunity", "Pipeline", "None", "None", "No award data published; eSupply has item lines + qty."),
]


# ---------------------------------------------------------------- build master workbook
def build_master():
    wb = openpyxl.Workbook()

    # README
    rd = wb.active
    rd.title = "README"
    n_cpv_in = sum(1 for c, _ in CPV_ROWS if c in CPV)
    n_nupco_in = sum(1 for n in NUPCO if NUPCO_SCOPE[n["id"]][0] in (C, P))
    lines = [
        ("TURMECA tender master", f"Built {TODAY} from the 'Tenders' Google Doc (<TENDERS_DOC_ID>) and "
                                  "Mapping_Taxonomy.xlsx (Google Sheet <TAXONOMY_SHEET_ID>)."),
        ("Master_Table", "Common 2024-2026 tender master: the doc's column groups, plus Tender Stage, Price Evidence and Manufacturer Evidence "
                         "(green headers) with dropdowns. Row ID and Award Year added for traceability; Value USD, HBP USD, Award Year and "
                         "Manufacturer Explicit? are formulas."),
        ("Field_Dictionary", "Definition, allowed values and origin of every master column."),
        ("Evidence_Rules", "Default Tender Stage / Price Evidence / Manufacturer Evidence per portal and record type, based on what each portal "
                           "actually publishes (checked live 2026-10-01). Override per row when the document says more."),
        ("Codes_By_Area", "The in-scope code list per Medtronic area: one row per OU x Sub-OU x code, for NUPCO tender families, "
                          "CPV/OKAS 331xxxxx codes and EKAP product-class (OKAS) codes."),
        ("NUPCO_Tender_Families", f"All {len(NUPCO)} tenders listed on nupco.com (pages 1-5) with English title, status, stage, documents and "
                                  f"Medtronic scope. {n_nupco_in} are Core/Partial."),
        ("CPV_OKAS_331", f"All {len(CPV_ROWS)} CPV 2008 codes in 331xxxxx (EU Publications Office), with the Turkish OKAS name EKAP uses "
                         f"where seen. {n_cpv_in} are in scope (Core / Partial / Catch-all)."),
        ("EKAP_Product_Class", "EKAP has no separate product-class scheme: the item-level class (ihtiyacKalemiOkasList.kodu) and the search filter "
                               "(okasBransKodList) both use OKAS codes. This tab is the filter set per area, with how often EKAP used each code "
                               "in 9,551 previously pulled tenders and the Turkish keyword screen to apply to lot names."),
        ("TR_Keywords", "Suggested Turkish keyword screen per Sub-OU for EKAP lot names (kısım adı). Suggested, not yet validated against EKAP lots."),
        ("Relevance", "Core = code/tender mainly holds products of that area. Partial = shared with non-Medtronic products, keep only lines that "
                      "match the keyword screen. Catch-all = parent class, pull only with the keyword screen. Review = needs a manual look. "
                      "Out = no product in the taxonomy."),
        ("Caveats", "EKAP OKAS names/counts come from earlier IVD-scoped EKAP pulls, so device codes are under-represented in the counts; "
                    "codes not seen there show the EU English label only. EKAP v2 currently needs human verification (Turnstile) before "
                    "extraction. NUPCO publishes no prices. Sub-OU assignments on CPV codes are judgement calls from the code definitions."),
    ]
    for i, (a, b) in enumerate(lines, 1):
        rd.cell(i, 1, a).font = BOLD
        rd.cell(i, 2, b).alignment = WRAP
    rd.column_dimensions["A"].width = 24
    rd.column_dimensions["B"].width = 130

    # Lists (for validation)
    ls = wb.create_sheet("Lists")
    lists = {
        "Country": ["KSA", "Turkey", "Pakistan", "Kuwait", "UAE"],
        "Portal": ["NUPCO", "Etimad", "EKAP v2", "EKAP legacy", "EPADS", "PPRA", "KP DGHS", "Punjab DGHS", "Sindh PPRA", "CAPT",
                   "MoF DPP", "ADGPG", "Dubai eSupply"],
        "Source Type": ["Item list", "Preliminary result", "Final result", "Award list", "Contract", "BER", "Notice"],
        "Tender Stage": ["Pipeline", "Awarded", "Contracted"],
        "Price Evidence": ["Line price", "Contract-level only", "None"],
        "Manufacturer Evidence": ["Explicit", "Inferred from distributor", "None"],
        "YesNo": ["Yes", "No"],
        "Confidence": ["High", "Medium", "Low"],
        "QC": ["Included", "Review", "Excluded"],
        "DupFlag": ["Duplicate", "Retender"],
        "Currency": ["SAR", "TRY", "PKR", "KWD", "AED", "USD"],
        "OU": OUS,
        "Sub-OU": sorted({t[2] for t in TAX}),
    }
    rng = {}
    for ci, (k, vals) in enumerate(lists.items(), 1):
        ls.cell(1, ci, k).font = BOLD
        for ri, v in enumerate(vals, 2):
            ls.cell(ri, ci, v)
        col = get_column_letter(ci)
        rng[k] = f"Lists!${col}$2:${col}${len(vals)+1}"
        ls.column_dimensions[col].width = max(14, max(len(str(v)) for v in vals) + 2)

    # Master_Table
    mt = wb.create_sheet("Master_Table", 1)
    cols, groups = [], []
    for g, fs in MASTER:
        groups.append((g, len(cols) + 1, len(cols) + len(fs)))
        cols += fs
    for g, a, b in groups:
        mt.merge_cells(start_row=1, start_column=a, end_row=1, end_column=b)
        x = mt.cell(1, a, g)
        x.font, x.fill, x.alignment = BOLD, BAND, Alignment(horizontal="center")
    header(mt, 2, cols)
    green = PatternFill("solid", fgColor="548235")
    for i, c in enumerate(cols, 1):
        mt.column_dimensions[get_column_letter(i)].width = max(12, min(30, len(c) + 4))
        if c in NEW_COLS:
            mt.cell(2, i).fill = green
    ci = {c: get_column_letter(i) for i, c in enumerate(cols, 1)}
    N = 2000
    first, last = 3, N + 2
    for r in range(first, last + 1):
        mt[f"{ci['Award Year']}{r}"] = (f'=IF({ci["Award/Contract Date"]}{r}<>"",YEAR({ci["Award/Contract Date"]}{r}),'
                                        f'IF({ci["Result Date"]}{r}<>"",YEAR({ci["Result Date"]}{r}),""))')
        mt[f"{ci['Value USD']}{r}"] = (f'=IF(OR({ci["Total Award Value"]}{r}="",{ci["LCY per USD"]}{r}=""),"",'
                                       f'{ci["Total Award Value"]}{r}/{ci["LCY per USD"]}{r})')
        mt[f"{ci['HBP USD']}{r}"] = (f'=IF(OR({ci["Price Evidence"]}{r}<>"Line price",{ci["Unit Award Price"]}{r}="",'
                                     f'{ci["LCY per USD"]}{r}=""),"",{ci["Unit Award Price"]}{r}/{ci["LCY per USD"]}{r})')
        mt[f"{ci['Manufacturer Explicit?']}{r}"] = (f'=IF({ci["Manufacturer Evidence"]}{r}="","",'
                                                    f'IF({ci["Manufacturer Evidence"]}{r}="Explicit","Yes","No"))')
    for c in ("Publication Date", "Opening Date", "Award/Contract Date", "Result Date", "FX Date"):
        for r in range(first, last + 1):
            mt[f"{ci[c]}{r}"].number_format = "yyyy-mm-dd"
    for c in ("Unit Award Price", "Total Award Value", "Value USD", "HBP USD"):
        for r in range(first, last + 1):
            mt[f"{ci[c]}{r}"].number_format = "#,##0.00"
    val_map = {"Country": "Country", "Portal": "Portal", "Source Type": "Source Type", "Tender Stage": "Tender Stage",
               "Price Evidence": "Price Evidence", "Manufacturer Evidence": "Manufacturer Evidence", "Winner Disclosed?": "YesNo",
               "Model Explicit?": "YesNo", "Award vs Budget?": "YesNo", "Confidence": "Confidence", "Included/Review/Excluded": "QC",
               "Duplicate/Retender Flag": "DupFlag", "Currency": "Currency", "OU": "OU", "Sub-OU": "Sub-OU"}
    strict = {"Tender Stage", "Price Evidence", "Manufacturer Evidence", "Included/Review/Excluded", "OU"}
    for col, key in val_map.items():
        dv = DataValidation(type="list", formula1=rng[key], allow_blank=True, showErrorMessage=col in strict,
                            errorStyle="stop" if col in strict else "warning")
        dv.add(f"{ci[col]}{first}:{ci[col]}{last}")
        mt.add_data_validation(dv)
    mt.freeze_panes = "C3"
    mt.auto_filter.ref = f"A2:{get_column_letter(len(cols))}{last}"

    # Field_Dictionary
    fd = wb.create_sheet("Field_Dictionary", 2)
    header(fd, 1, ["Group", "Field", "Definition", "Allowed values / format", "Origin"], [16, 28, 90, 50, 26])
    rows = []
    for g, fs in MASTER:
        for f in fs:
            d, a, o = DICT[f]
            rows.append([g, f, d, a, o])
    put_rows(fd, 2, rows)
    for r in range(2, len(rows) + 2):
        if fd.cell(r, 5).value == "NEW":
            for c in range(1, 6):
                fd.cell(r, c).fill = PatternFill("solid", fgColor="E2EFDA")

    # Evidence_Rules
    er = wb.create_sheet("Evidence_Rules", 3)
    header(er, 1, ["Country", "Portal", "Record type", "Tender Stage", "Price Evidence", "Manufacturer Evidence", "Basis (verified 2026-10-01)"],
           [10, 24, 50, 14, 20, 30, 80])
    put_rows(er, 2, [list(x) for x in EVIDENCE_RULES])

    # NUPCO_Tender_Families
    nt = wb.create_sheet("NUPCO_Tender_Families")
    header(nt, 1, ["Tender ID", "Type", "English title", "Arabic title", "NUPCO status", "Tender Stage", "Opening date", "Documents on page",
                   "Has Final Result", "Medtronic relevance", "Medtronic OU(s)", "Sub-OU(s)", "Note / reason", "Tender page URL", "Final Result URL"],
           [14, 10, 55, 45, 16, 14, 13, 40, 9, 12, 40, 50, 45, 30, 30])
    rows = []
    order = {"Core": 0, "Partial": 1, "Review": 2, "Out": 3}
    for n in sorted(NUPCO, key=lambda n: (order[NUPCO_SCOPE[n["id"]][0]], n["id"])):
        rel, pairs, note = NUPCO_SCOPE[n["id"]]
        prefix = n["id"][:3] if n["id"] else ""
        typ = {"NPT": "Tender", "NDP": "Direct purchase"}.get(prefix, "Service")
        final = next((p["h"] for p in n["pdfs"] if p["t"] == "Tender Final Result"), "")
        rows.append([n["id"] or "(no ID)", typ, NUPCO_EN[n["id"]], n["title"], n["type"].replace("ResultsUnder", "Results / Under"),
                     nupco_stage(n["type"]), n["opening"] if "/" in n["opening"] else "", ", ".join(n["docs"]),
                     "Yes" if final else "No", rel, "; ".join(dict.fromkeys(o for o, _ in pairs)),
                     " | ".join(f"{o}: {s}" for o, s in pairs), note, n["href"], final])
    put_rows(nt, 2, rows, rel_col=10)

    # CPV_OKAS_331
    cp = wb.create_sheet("CPV_OKAS_331")
    header(cp, 1, ["CPV/OKAS code", "Level", "Parent code", "English (EU CPV 2008)", "Turkish OKAS name (as used in EKAP)",
                   "EKAP items with code (prior pull)", "Medtronic relevance", "Medtronic OU(s)", "Sub-OU(s)", "Note / out-of-scope reason"],
           [13, 7, 12, 50, 55, 12, 12, 40, 60, 60])
    codes = [c for c, _ in CPV_ROWS]
    rows = []
    for c, en in CPV_ROWS:
        level = len(c.rstrip("0"))
        parent = ""
        for k in range(level - 1, 2, -1):
            cand = c[:k].ljust(8, "0")
            if cand in codes and cand != c:
                parent = cand
                break
        tr, cnt = OKAS.get(c, ["", 0])
        if c in CPV:
            rel, pairs, note = CPV[c]
            ous = "; ".join(dict.fromkeys(o for o, _ in pairs))
            subs = " | ".join(f"{o}: {s}" for o, s in pairs)
        else:
            rel, ous, subs, note = "Out", "", "", out_reason(c)
        rows.append([c, level, parent, en, tr, cnt or None, rel, ous, subs, note])
    put_rows(cp, 2, rows, rel_col=7)

    # EKAP_Product_Class
    ek = wb.create_sheet("EKAP_Product_Class")
    header(ek, 1, ["OKAS code (okasBransKodList)", "Turkish OKAS name (EKAP)", "English (EU CPV)", "EKAP items with code (prior pull)",
                   "Medtronic relevance", "Pull rule", "Medtronic OU(s)", "Turkish keyword screen (union of mapped Sub-OUs)"],
           [16, 55, 45, 12, 12, 45, 40, 90])
    rule = {C: "Pull all lots; screen out obvious non-device lines.", P: "Pull all lots; keep lines matching the keyword screen.",
            K: "Pull only with the keyword screen (very large, mostly non-Medtronic)."}
    en_map = dict(CPV_ROWS)
    ek_rows = []
    extra = [("33000000", K, [(o, "all") for o in OUS], "Top of medical devices/pharma/personal care. Seen 90x in EKAP; keyword screen only.")]
    for c, rel, pairs, note in [(c, *CPV[c]) for c in sorted(CPV)] + extra:
        tr, cnt = OKAS.get(c, ["", 0])
        kws = []
        for ou, s in expand(pairs):
            for kw in TR_KW.get((ou, s), "").split("; "):
                if kw and kw not in kws:
                    kws.append(kw)
        if rel == K:
            kwtxt = "Use the full TR_Keywords list for the target Sub-OU."
        else:
            kwtxt = "; ".join(kws)
        ek_rows.append([c, tr or "(not seen in prior EKAP pull)", en_map.get(c, "Medical equipments, pharmaceuticals and personal care products"),
                        cnt or None, rel, rule[rel], "; ".join(dict.fromkeys(o for o, _ in pairs)), kwtxt])
    ek_rows.sort(key=lambda r: ({C: 0, P: 1, K: 2}[r[4]], -(r[3] or 0), r[0]))
    put_rows(ek, 2, ek_rows, rel_col=5)

    # TR_Keywords
    tk = wb.create_sheet("TR_Keywords")
    header(tk, 1, ["OU", "Sub-OU", "Turkish keywords for EKAP lot names (suggested, validate)"], [36, 55, 120])
    put_rows(tk, 2, [[ou, s, TR_KW.get((ou, s), "")] for ou in OUS for s in SUBOUS[ou]])

    # Codes_By_Area (long)
    cb = wb.create_sheet("Codes_By_Area", 1)
    header(cb, 1, ["OU", "Sub-OU", "Code system", "Code", "Label", "Relevance", "Tender Stage / status", "Note"],
           [36, 50, 26, 14, 70, 11, 18, 60])
    rows = []
    for n in NUPCO:
        rel, pairs, note = NUPCO_SCOPE[n["id"]]
        if rel not in (C, P):
            continue
        for ou, s in expand(pairs):
            rows.append([ou, s, "NUPCO tender family", n["id"], NUPCO_EN[n["id"]], rel, nupco_stage(n["type"]), note])
    for c in sorted(CPV):
        rel, pairs, note = CPV[c]
        for ou, s in expand(pairs):
            rows.append([ou, s, "CPV/OKAS 331xxxxx", c, en_map[c], rel, "", note])
    for r in ek_rows:
        c = r[0]
        pairs = CPV[c][1] if c in CPV else extra[0][2]
        for ou, s in expand(pairs):
            rows.append([ou, s, "EKAP product class (OKAS)", c, r[1] if not r[1].startswith("(") else r[2], r[4], "",
                         f"EKAP items in prior pull: {r[3] or 0}. {r[5]}"])
    sys_order = {"NUPCO tender family": 0, "CPV/OKAS 331xxxxx": 1, "EKAP product class (OKAS)": 2}
    rel_order = {C: 0, P: 1, K: 2}
    rows.sort(key=lambda r: (OUS.index(r[0]), SUBOUS[r[0]].index(r[1]), sys_order[r[2]], rel_order[r[5]], r[3]))
    put_rows(cb, 2, rows, rel_col=6)

    # Area summary at top of Codes_By_Area? keep a separate compact summary
    sm = wb.create_sheet("Area_Summary", 1)
    header(sm, 1, ["OU", "Sub-OU", "NUPCO tender families (Core)", "NUPCO tender families (Partial)", "CPV/OKAS 331 codes (Core)",
                   "CPV/OKAS 331 codes (Partial)", "Catch-all parents", "Turkish keyword screen"], [34, 45, 40, 40, 45, 55, 40, 70])
    agg = collections.defaultdict(lambda: collections.defaultdict(list))
    for r in rows:
        key = (r[0], r[1])
        if r[2] == "NUPCO tender family":
            agg[key]["n" + r[5]].append(r[3])
        elif r[2] == "CPV/OKAS 331xxxxx":
            agg[key]["c" + r[5]].append(r[3])
        elif r[5] == K:
            agg[key]["k"].append(r[3])
    srows = []
    for ou in OUS:
        for s in SUBOUS[ou]:
            a = agg[(ou, s)]
            u = lambda xs: ", ".join(dict.fromkeys(xs))
            srows.append([ou, s, u(a["n" + C]), u(a["n" + P]), u(a["c" + C]), u(a["c" + P]), u(a["k"]), TR_KW.get((ou, s), "")])
    put_rows(sm, 2, srows)

    wb.move_sheet("Lists", offset=len(wb.sheetnames))
    wb.save(OUT_MASTER)
    return len(rows), n_cpv_in, n_nupco_in


if __name__ == "__main__":
    write_taxonomy()
    n, ncpv, nnup = build_master()
    print("taxonomy rows", len(TAX), "| codes_by_area rows", n, "| CPV in scope", ncpv, "/", len(CPV_ROWS), "| NUPCO in scope", nnup, "/", len(NUPCO))
