"""
Federal Pakistan (PPRA EPMS) medical evaluation reports -> PK_Federal_Evaluations_Mapped.xlsx

Two formats in federal/ber_text.json (text layer or OCR):
  A. EPADS-generated item reports: "<ITEM> (INDIVIDUAL ITEM) / UNSPSC: <class> / BID OFFER / <bidder> / <type> / Marks: n / <offer PKR>".
     Offers are line totals (qty x rate, qty not published) -> Price Evidence = Line price (value), no unit price.
     Winner = lowest offer among bidders with marks >= 70 (least-cost based selection).
  B. Package reports (mostly scanned): bidders, evaluated cost (PKR) per package, "Most Advantageous Bidder".
     -> Price Evidence = Contract-level only.
Category from item / package title only (vocab layer); Company is not published (bidders are distributors) -> None.
"""
import csv, json, os, re, sys
from collections import defaultdict

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(os.path.dirname(HERE)))
from vocab_rules import Vocab  # noqa: E402
from build_ref_prices import PKR_PER_USD  # noqa: E402

MED_TITLE = re.compile(r"medical|surgical|ventilator|anesth|anaesth|monitoring equipment|electrosurg|laparoscop|defibrill|ecg|suction|ot table|c-arm|"
                       r"radiology|urodynamic|lithotripter|endoscop|neurology|apheresis|blood gas|medicine|disposable|hospital prd|dental|mortuary", re.I)
MONEY = re.compile(r"^\s*([\d]{1,3}(?:,\d{3})+(?:\.\d+)?|\d+(?:\.\d+)?)\s*$")


def money(s):
    m = re.search(r"(\d{1,3}(?:[,.]\d{3})+(?:\.\d{1,2})?|\d{4,}(?:\.\d{1,2})?)", s or "")
    if not m:
        return None
    x = m.group(1)
    if x.count(".") > 1:              # OCR '168.519.229.30' -> 168519229.30
        head, tail = x.rsplit(".", 1)
        x = head.replace(".", "") + "." + tail
    try:
        return float(x.replace(",", ""))
    except ValueError:
        return None


def parse_epads(text):
    """Format A -> list of item dicts with bids."""
    text = re.sub(r"EPADS v2\.0 generated document.*?\n.*?Page \d+ of \d+\s*\n", "\n", text)
    blocks = re.split(r"\n(?=[^\n]+\(INDIVIDUAL\s*\n?\s*ITEM\)|[^\n]+\(INDIVIDUAL ITEM\))", text)
    items = {}
    for b in blocks:
        m = re.match(r"\s*(.+?)\(INDIVIDUAL\s*ITEM\)\s*\nUNSPSC:\s*([^\n]+)\n", b, re.S)
        if not m:
            continue
        name = re.sub(r"\s+", " ", m.group(1)).strip()
        it = items.setdefault(name, {"item": name, "unspsc": m.group(2).strip(), "bids": []})
        for bm in re.finditer(r"\n([^\n]+)\n([^\n]+)\nMarks:\s*([\d.]+)\s*\n\s*([\d,]+(?:\.\d+)?)\s*\n", b):
            it["bids"].append({"bidder": bm.group(1).strip(), "marks": float(bm.group(3)), "offer": float(bm.group(4).replace(",", ""))})
    return list(items.values())


def parse_package(text):
    """Format B -> bidders with evaluated cost + most advantageous bidder."""
    bidders = []
    for line in text.splitlines():
        m = re.match(r"\s*(?:\d+\s*)?(M/s\.?\s*[^|]+?)\s+(Qualified|Not Qualified|Disqualified|Responsive|Non[- ]Responsive)?.*?$", line, re.I)
        if m and "M/s" in line:
            bidders.append({"bidder": re.sub(r"\s+(Qualified|Not Qualified).*$", "", m.group(1)).strip(" |"),
                            "status": "Not qualified" if re.search(r"not qualified|disqualif|non[- ]responsive", line, re.I) else "Qualified",
                            "evaluated_cost_pkr": money(line)})
    best = re.search(r"Most Advantageous Bidder:\s*([^\n]+)", text)
    return bidders, (best.group(1).strip() if best else "")


def main():
    v = Vocab()
    texts = json.load(open(os.path.join(HERE, "ber_text.json"), encoding="utf8"))
    meta = {m["evaluation_no"]: m for m in csv.DictReader(open(os.path.join(HERE, "federal_medical_evaluations.csv"), encoding="utf-8-sig"))}
    item_rows, pkg_rows = [], []
    for k, d in texts.items():
        evl = k.split("_")[0]
        m = meta.get(evl, {})
        t = d["text"]
        if "(INDIVIDUAL ITEM)" in t or "(INDIVIDUAL\nITEM)" in t:
            for it in parse_epads(t):
                pm = v.product(it["item"]) or {}
                elig = [b for b in it["bids"] if b["marks"] >= 70]
                win = min(elig, key=lambda b: b["offer"]) if elig else None
                for b in it["bids"]:
                    item_rows.append({"Evaluation No": evl, "Procurement": m.get("title"), "Buyer": m.get("office") or m.get("organization"),
                                      "Published": m.get("published"), "Item": it["item"], "UNSPSC": it["unspsc"], "Bidder (distributor)": b["bidder"],
                                      "Marks": b["marks"], "Offer PKR (line total)": b["offer"], "Offer USD": round(b["offer"] / PKR_PER_USD, 2),
                                      "Winner": "Yes" if b is win else "", "OU": pm.get("ou"), "Sub-OU": pm.get("sub_ou"),
                                      "Device Category": pm.get("product_category"), "Tender Stage": "Awarded" if "final" in (m.get("evaluation_type") or "").lower() or "Evaluation Report (Under" in t[:200] else "Pipeline",
                                      "Price Evidence": "Line price (total, no qty)", "Manufacturer Evidence": "None"})
        else:
            if not MED_TITLE.search(m.get("title") or ""):
                continue
            bidders, best = parse_package(t)
            pm = v.product(m.get("title") or "") or {}
            for b in bidders or [{"bidder": "", "status": "", "evaluated_cost_pkr": None}]:
                pkg_rows.append({"Evaluation No": evl, "Procurement": m.get("title"), "Buyer": m.get("office") or m.get("organization"),
                                 "City": m.get("city"), "Evaluation Type": m.get("evaluation_type"), "Published": m.get("published"),
                                 "Bids Received": m.get("bids"), "Bidder (distributor)": b["bidder"], "Status": b["status"],
                                 "Evaluated Cost PKR": b["evaluated_cost_pkr"],
                                 "Evaluated Cost USD": round(b["evaluated_cost_pkr"] / PKR_PER_USD, 2) if b["evaluated_cost_pkr"] else None,
                                 "Most Advantageous / Lowest Bidder": best or m.get("lowest_bidder"), "OU": pm.get("ou"), "Sub-OU": pm.get("sub_ou"),
                                 "Device Category (from title)": pm.get("product_category"), "Text Source": d["src"],
                                 "Tender Stage": "Awarded" if "final" in (m.get("evaluation_type") or "").lower() else "Pipeline",
                                 "Price Evidence": "Contract-level only", "Manufacturer Evidence": "None"})
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    hdr = PatternFill("solid", fgColor="1F4E78")

    def sheet(title, rows, widths):
        ws = wb.create_sheet(title)
        if not rows:
            ws.append(["(none)"])
            return
        cols = list(rows[0].keys())
        ws.append(cols)
        for c in ws[1]:
            c.font, c.fill, c.alignment = Font(bold=True, color="FFFFFF"), hdr, Alignment(wrap_text=True)
        for r in rows:
            ws.append([r[c] for c in cols])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(rows)+1}"
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    ws = wb.create_sheet("README")
    n_items = len({(r["Evaluation No"], r["Item"]) for r in item_rows})
    for line in [
        "Pakistan federal (PPRA EPMS) evaluation reports for medical procurements, mapped to Mapping_Taxonomy.xlsx",
        f"EPMS lists 3,676 federal evaluations; {len(meta)} are medical-related and all reports were downloaded (scans OCR'd with Tesseract).",
        f"EPADS item-level reports: {n_items} items / {len(item_rows)} bids - offers are line totals (no quantity, no unit price, no make).",
        f"Package-level reports: {len({r['Evaluation No'] for r in pkg_rows})} evaluations - evaluated cost per package and the most advantageous bidder.",
        "EPMS public contracts register (694 contracts) holds almost no medical contracts; the legacy ppra.org.pk contract_org page now redirects to EPMS.",
        "Bidders are distributors; no manufacturer is published, so Company is not filled. OCR'd amounts should be spot-checked against the PDF.",
    ]:
        ws.append([line])
    ws.column_dimensions["A"].width = 150
    sheet("Item_Bids", item_rows, [16, 40, 30, 14, 50, 25, 30, 7, 14, 11, 7, 25, 30, 35, 10, 18, 12])
    sheet("Package_Bids", pkg_rows, [16, 60, 35, 12, 14, 14, 6, 40, 12, 16, 14, 35, 25, 30, 35, 8, 10, 18, 12])
    wb.save(os.path.join(HERE, "PK_Federal_Evaluations_Mapped.xlsx"))
    print(f"item bids {len(item_rows)} ({n_items} items, mapped {sum(1 for r in item_rows if r['Device Category'])}); "
          f"package rows {len(pkg_rows)} ({len({r['Evaluation No'] for r in pkg_rows})} evaluations, with cost {sum(1 for r in pkg_rows if r['Evaluated Cost PKR'])})")


if __name__ == "__main__":
    main()
