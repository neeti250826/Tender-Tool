"""Full pipeline: map all three data tabs of the VN medtech tenders file against
the Clearstate MedTech taxonomy (keyword + semantic/fuzzy) and write one workbook
with a combined 'Mapped' sheet and per-source 'Unmapped' tabs.
"""
from __future__ import annotations

import sys, time, datetime, json
from collections import Counter
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE


def sanitize(v):
    """Make a value safe for openpyxl (tz-aware dt, inf/nan, illegal chars, length)."""
    if v is None:
        return None
    if isinstance(v, bool) or isinstance(v, int):
        return v
    if isinstance(v, float):
        return str(v) if (v != v or v in (float("inf"), float("-inf"))) else v
    if isinstance(v, datetime.datetime):
        return v.replace(tzinfo=None) if v.tzinfo is not None else v
    if isinstance(v, datetime.date):
        return v
    s = str(v)
    s = ILLEGAL_CHARACTERS_RE.sub("", s)
    return s[:32000] if len(s) > 32000 else s

sys.path.insert(0, r"<WORKSPACE>/mapping")
from taxonomy import load_taxonomy
from matcher import match_row
from enrich import bilingual_matches


def extract_part_lot_name(detail: str) -> str:
    """Pull the specific (Vietnamese) line-item name out of awarded_value_detail JSON."""
    if not detail or "part_lot_name" not in detail:
        return ""
    try:
        d = json.loads(detail)
        return str(d.get("part_lot_name") or d.get("raw__Part/Lot name") or "")
    except Exception:
        return ""

SRC = r"<DOWNLOADS>/VN_Medtech_Player_Mapped_Enriched.xlsx"
TAXO = r"<DOWNLOADS>/Mapping_Taxonomy.xlsx"
OUT = r"<WORKSPACE>/VN_Medtech_Taxonomy_Mapped.xlsx"

SOURCE_TABS = ["Medtech Player Data", "Removed Rows", "Manual Review Sample"]

ITEM_FIELDS = ["item_description_english"]                      # -> Product Category, Family
CONTEXT_FIELDS = ["title_english", "item_description_english"]  # -> OU, Sub-OU
# Company is the OEM/manufacturer, so match only the plugin's OEM-signal fields.
# supplier_name_english is the local VN distributor/bidder and produces false
# positives (e.g. "MERCURY ELECTRIK VIETNAM" -> "Mercury Medical"), so it is excluded.
COMPANY_FIELDS = ["final_player", "mapped_oem", "likely_oem", "oem_match_terms"]

DIMS_ORDER = ["ou", "sub_ou", "product_category", "company", "family"]

# Mapping provenance columns added in front of each row.
MAP_COLS = ["mapping_status", "mapped_fields", "n_mapped_dims"]
for d in DIMS_ORDER:
    MAP_COLS += [f"map_{d}", f"map_{d}_term", f"map_{d}_method", f"map_{d}_confidence"]
MAP_COLS += ["map_family_category"]


def build_unified_header(wb):
    base = None
    extras = []
    seen = set()
    for tab in SOURCE_TABS:
        hdr = [c.value for c in next(wb[tab].iter_rows(min_row=1, max_row=1))]
        if base is None:
            base = hdr
            seen = set(hdr)
        else:
            for h in hdr:
                if h not in seen:
                    seen.add(h); extras.append(h)
    return base + extras


def main():
    t0 = time.time()
    dims = load_taxonomy(TAXO)
    print("Taxonomy:", {k: len(v) for k, v in dims.items()})

    wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)
    unified = build_unified_header(wb)
    out_header = ["source_tab"] + MAP_COLS + unified

    mapped_rows = []
    unmapped = {tab: [] for tab in SOURCE_TABS}
    stats = {tab: Counter() for tab in SOURCE_TABS}
    dim_value_counts = {d: Counter() for d in DIMS_ORDER}
    method_counts = Counter()

    for tab in SOURCE_TABS:
        ws = wb[tab]
        it = ws.iter_rows(values_only=True)
        hdr = list(next(it))
        idx = {h: i for i, h in enumerate(hdr)}

        def get(row, col):
            i = idx.get(col)
            if i is None or i >= len(row):
                return None
            return row[i]

        n = 0
        for row in it:
            n += 1
            itext = " ".join(str(get(row, c)) for c in ITEM_FIELDS if get(row, c) is not None)
            xtext = " ".join(str(get(row, c)) for c in CONTEXT_FIELDS if get(row, c) is not None)
            ctext = " ".join(str(get(row, c)) for c in COMPANY_FIELDS if get(row, c) is not None)
            supplier = str(get(row, "supplier_name_english") or "")
            res = match_row(dims, itext, xtext, ctext, supplier_text=supplier)

            # Additive bilingual (EN+VN) keyword layer over the Vietnamese line item.
            vn_part = extract_part_lot_name(str(get(row, "awarded_value_detail") or ""))
            all_text = " ".join([itext, str(get(row, "item_description") or ""),
                                  str(get(row, "title") or ""), vn_part])
            for d, m in bilingual_matches(all_text).items():
                if not res.get(d):
                    res[d] = m

            matched_fields = [d for d in DIMS_ORDER if res.get(d)]
            status = "Mapped" if matched_fields else "Unmapped"
            stats[tab][status] += 1
            for d in matched_fields:
                stats[tab][f"dim_{d}"] += 1
                dim_value_counts[d][res[d][0]] += 1
                method_counts[res[d][2]] += 1

            # Build mapping cell values
            mp = {"mapping_status": status,
                  "mapped_fields": ",".join(matched_fields),
                  "n_mapped_dims": len(matched_fields),
                  "map_family_category": ""}
            for d in DIMS_ORDER:
                r = res.get(d)
                if r:
                    mp[f"map_{d}"] = r[0]
                    mp[f"map_{d}_term"] = r[1]
                    mp[f"map_{d}_method"] = r[2]
                    mp[f"map_{d}_confidence"] = r[4]
                else:
                    mp[f"map_{d}"] = mp[f"map_{d}_term"] = mp[f"map_{d}_method"] = mp[f"map_{d}_confidence"] = ""
            if res.get("family"):
                mp["map_family_category"] = res["family"][5].get("category", "")

            # Assemble unified original values
            orig = [get(row, col) for col in unified]
            out_row = [sanitize(x) for x in ([tab] + [mp[c] for c in MAP_COLS] + orig)]

            if status == "Mapped":
                mapped_rows.append(out_row)
            else:
                unmapped[tab].append(out_row)
        print(f"  {tab}: {n} rows -> mapped {stats[tab]['Mapped']}, unmapped {stats[tab]['Unmapped']}")

    # ---- write workbook (single write-only pass) ----
    from openpyxl.cell import WriteOnlyCell
    out = openpyxl.Workbook(write_only=True)
    total = sum(stats[t]['Mapped'] + stats[t]['Unmapped'] for t in SOURCE_TABS)
    tot_mapped = sum(stats[t]['Mapped'] for t in SOURCE_TABS)

    # Summary sheet
    s = out.create_sheet("Mapping Summary")
    s.column_dimensions["A"].width = 42
    s.column_dimensions["B"].width = 95
    title_cell = WriteOnlyCell(s, value="VN Medtech Tenders — Taxonomy Mapping")
    title_cell.font = Font(bold=True, size=14)
    s.append([title_cell])
    lines = [
        ["Generated", time.strftime("%Y-%m-%d %H:%M")],
        ["Source file", "VN_Medtech_Player_Mapped_Enriched.xlsx"],
        ["Taxonomy file", "Mapping_Taxonomy.xlsx (dim_overview=OU, dim_subview=Sub-OU, dim_product=Product Category, dim_company=Company, dim_procedure=Family)"],
        ["Method", "Keyword (exact distinctive-token match, locality-checked) + semantic fuzzy word-similarity (rapidfuzz, no embeddings)"],
        ["'Mapped' definition", "row matched at least one of OU / Sub-OU / Product Category / Company / Family"],
        ["", ""],
        ["Total rows processed", total],
        ["Mapped rows (>=1 dimension)", tot_mapped],
        ["Unmapped rows", total - tot_mapped],
        ["Mapped %", round(100 * tot_mapped / total, 2)],
        ["", ""],
        ["Per source tab", "rows / mapped / unmapped"],
    ]
    for tab in SOURCE_TABS:
        tot = stats[tab]['Mapped'] + stats[tab]['Unmapped']
        lines.append([tab, f"{tot} / {stats[tab]['Mapped']} / {stats[tab]['Unmapped']}"])
    lines.append(["", ""])
    lines.append(["Coverage by dimension (rows with a match)", ""])
    for d in DIMS_ORDER:
        c = sum(stats[t][f"dim_{d}"] for t in SOURCE_TABS)
        lines.append([d, f"{c}  ({round(100*c/total,2)}%)"])
    lines.append(["", ""])
    lines.append(["Matches by method (dimension-level)", ""])
    for m, c in method_counts.most_common():
        lines.append([m, c])
    lines.append(["", ""])
    lines.append(["Top mapped values per dimension", ""])
    for d in DIMS_ORDER:
        top = dim_value_counts[d].most_common(10)
        lines.append([d, "; ".join(f"{v} ({n})" for v, n in top)])
    for ln in lines:
        s.append(ln)

    def write_sheet(title, header, rows):
        ws = out.create_sheet(title[:31])
        ws.freeze_panes = "A2"
        hcells = []
        for h in header:
            c = WriteOnlyCell(ws, value=h)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F4E78")
            hcells.append(c)
        ws.append(hcells)
        for r in rows:
            ws.append(r)

    write_sheet("Mapped", out_header, mapped_rows)
    for tab in SOURCE_TABS:
        write_sheet("Unmapped - " + tab, out_header, unmapped[tab])

    out.save(OUT)

    print(f"\nDONE in {time.time()-t0:.1f}s  ->  {OUT}")
    print(f"Total {total} | Mapped {tot_mapped} ({100*tot_mapped/total:.1f}%) | Unmapped {total-tot_mapped}")


if __name__ == "__main__":
    main()
