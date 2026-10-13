"""
Turkey / EKAP item-level build.

Sources: every EKAP tender-detail JSON on disk - ekap/details/ (MedTech pull) plus the earlier IVD pull's
cached details under <TURKEY_IVD_PROJECT>/outputs/**/raw*details/ - de-duplicated by tender id.

One row per contract x lot (kisim). Rules:
  * Award Year = contract date year (sozlesmeTarih); only 2024-2026 kept as Included.
  * Contract values cover several lots -> never split or summed across lots. "Line Value TRY" is filled only
    when the contract has exactly one lot (Price Evidence = Line price); otherwise Contract-level only.
  * Product Category / Sub-OU / OU from the lot name only: Turkish rules first, then the Google-translated
    English lot name through the same vocabulary.
  * Company only when the supplier's registered name is an OEM / OEM affiliate (alias table) ->
    Manufacturer Evidence = Inferred from distributor; otherwise None (no guessing).
Output: ekap/TR_EKAP_MedTech_Mapped.xlsx and ekap/tr_ekap_lots_all.csv
"""
import csv, glob, json, os, re, sys
from collections import Counter, defaultdict

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))
from vocab_rules import Vocab  # noqa: E402
from translate import translate_many  # noqa: E402

OLD = "<TURKEY_IVD_PROJECT>/outputs"
OUT_X = os.path.join(HERE, "TR_EKAP_MedTech_Mapped.xlsx")
OUT_C = os.path.join(HERE, "tr_ekap_lots_all.csv")


def money(s):
    """'21.540,00 TRY' (Turkish) or '21,540.00 TRY' -> float."""
    if not s or s in ("-",):
        return None
    m = re.search(r"[\d.,]+", str(s))
    if not m:
        return None
    x = m.group(0)
    if re.search(r",\d{1,2}$", x):
        x = x.replace(".", "").replace(",", ".")
    else:
        x = x.replace(",", "")
    try:
        return float(x)
    except ValueError:
        return None


def load_details():
    paths = glob.glob(os.path.join(HERE, "details", "*.json"))
    for p in glob.glob(OLD + "/**/*.json", recursive=True):
        if "details" in os.path.dirname(p):
            paths.append(p)
    seen = {}
    for p in paths:
        tid = os.path.basename(p)[:-5]
        if tid not in seen:
            seen[tid] = p
    return seen


QTY = re.compile(r"(\d[\d.]*)\s*(adet|ad\.|kutu|paket|takim|set|tane|kit)\b", re.I)


def main():
    v = Vocab()
    det = load_details()
    rows = []
    for tid, p in det.items():
        try:
            d = json.load(open(p, encoding="utf-8"))
        except Exception:
            continue
        b = d.get("ihaleBilgi") or {}
        okas = ";".join(str(o.get("kodu")) for o in d.get("ihtiyacKalemiOkasList") or [])
        idare = d.get("idare") or {}
        for s in d.get("sozlesmeBilgiList") or []:
            k = s.get("kisimItemDto") or {}
            lots = k.get("kisimList") or [{"kisimAdi": d.get("ihaleAdi")}]
            cdate = (s.get("sozlesmeTarih") or k.get("sozlesmeTarih") or "")[:10]
            cval = money(s.get("sozlesmeBedeli") or k.get("sozlesmeBedeli"))
            for lot in lots:
                name = (lot.get("kisimAdi") or "").strip()
                rows.append({
                    "Country": "Turkey", "Portal": "EKAP v2", "Source Type": "Contract", "Tender Detail ID": tid,
                    "Source URL": f"https://ekapv2.kik.gov.tr/ekap/search/{tid}", "Tender ID": d.get("ikn"),
                    "Tender Title": d.get("ihaleAdi"), "Procurement Authority": d.get("idareAdi"),
                    "City": (idare.get("il") or {}).get("adi"), "Tender Date": b.get("ihaleTarihSaat"),
                    "Award/Contract Date": cdate, "Award Year": int(cdate[:4]) if cdate[:4].isdigit() else None,
                    "CPV/OKAS Codes": okas, "Lot Name (TR)": name, "Lots in Contract": len(lots),
                    "Supplier/Distributor": (s.get("yukleniciAdi") or k.get("yuklenici") or "").strip(),
                    "Contract Value TRY": cval, "Line Value TRY": cval if len(lots) == 1 else None,
                    "Approx. Cost TRY (tender/lot)": money(s.get("yaklasikMaliyet")),
                    "Terminated": s.get("fesihString"),
                })
    print("details", len(det), "contract-lot rows", len(rows))

    names = sorted({r["Lot Name (TR)"] for r in rows if r["Lot Name (TR)"]})
    print("unique lot names", len(names))
    en = dict(zip(names, translate_many(names, src="tr")))

    pm_cache, co_cache = {}, {}
    for r in rows:
        n = r["Lot Name (TR)"]
        if n not in pm_cache:
            m = v.product(n)
            if m:
                m["method"] = "tr_rule"
            else:
                m = v.product(en.get(n, ""))
                if m:
                    m["method"] = "en_translation_rule"
            pm_cache[n] = m
        sup = r["Supplier/Distributor"]
        if sup not in co_cache:
            co_cache[sup] = v.company(sup)
        pm, cm = pm_cache[n] or {}, co_cache[sup] or {}
        q = QTY.search(n or "")
        qty = float(q.group(1).replace(".", "")) if q else None
        r.update({
            "Lot Name (EN)": en.get(n, ""), "OU": pm.get("ou"), "Sub-OU": pm.get("sub_ou"), "Device Category": pm.get("product_category"),
            "Mapping Method": pm.get("method"), "Company": cm.get("company"),
            "Family Name": v.family(cm.get("company"), pm.get("product_category"), n, en.get(n)) if cm and pm else None,
            "Qty in Lot Name": qty, "Unit Price TRY": (r["Line Value TRY"] / qty) if (qty and r["Line Value TRY"]) else None,
            "Tender Stage": "Contracted", "Price Evidence": "Line price" if r["Line Value TRY"] else "Contract-level only",
            "Manufacturer Evidence": "Inferred from distributor" if cm else "None",
            "Included/Review/Excluded": ("Included" if pm and r["Award Year"] in (2024, 2025, 2026) and r["Terminated"] in (None, "Yok")
                                         else ("Review" if pm else "Excluded")),
        })

    cols = ["Country", "Portal", "Source Type", "Source URL", "Tender ID", "Tender Title", "Procurement Authority", "City", "Tender Date",
            "Award/Contract Date", "Award Year", "CPV/OKAS Codes", "Lot Name (TR)", "Lot Name (EN)", "Lots in Contract",
            "OU", "Sub-OU", "Device Category", "Mapping Method", "Company", "Family Name", "Supplier/Distributor",
            "Contract Value TRY", "Line Value TRY", "Qty in Lot Name", "Unit Price TRY", "Approx. Cost TRY (tender/lot)",
            "Tender Stage", "Price Evidence", "Manufacturer Evidence", "Included/Review/Excluded", "Terminated", "Tender Detail ID"]
    with open(OUT_C, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=cols, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)

    mapped = [r for r in rows if r["Device Category"]]
    inc = [r for r in mapped if r["Included/Review/Excluded"] == "Included"]
    # category summary: lots, single-lot value, by year
    agg = defaultdict(lambda: {"lots": 0, "val": 0.0, "val_lots": 0, 2024: 0, 2025: 0, 2026: 0, "sup": Counter(), "co": Counter()})
    for r in inc:
        a = agg[(r["OU"], r["Sub-OU"], r["Device Category"])]
        a["lots"] += 1
        a[r["Award Year"]] += 1
        if r["Line Value TRY"]:
            a["val"] += r["Line Value TRY"]
            a["val_lots"] += 1
        a["sup"][r["Supplier/Distributor"][:50]] += 1
        if r["Company"]:
            a["co"][r["Company"]] += 1
    cat_rows = [[*k, a["lots"], a[2024], a[2025], a[2026], a["val_lots"], round(a["val"], 2),
                 "; ".join(f"{s} ({n})" for s, n in a["sup"].most_common(5)), "; ".join(f"{c} ({n})" for c, n in a["co"].most_common(5))]
                for k, a in sorted(agg.items(), key=lambda x: (x[0][0] or "", x[0][1] or "", -x[1]["lots"]))]
    sup = Counter(r["Supplier/Distributor"] for r in inc)
    sup_rows = []
    for s, n in sup.most_common(300):
        cats = Counter(r["Device Category"] for r in inc if r["Supplier/Distributor"] == s)
        sup_rows.append([s, n, (co_cache.get(s) or {}).get("company"), "; ".join(f"{c} ({k})" for c, k in cats.most_common(4))])
    unm = Counter((r["Lot Name (TR)"], r["Lot Name (EN)"]) for r in rows if not r["Device Category"] and r["Award Year"] in (2024, 2025, 2026))

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    hdr = PatternFill("solid", fgColor="1F4E78")

    def sheet(title, header, data, widths):
        ws = wb.create_sheet(title)
        ws.append(header)
        for c in ws[1]:
            c.font, c.fill, c.alignment = Font(bold=True, color="FFFFFF"), hdr, Alignment(wrap_text=True)
        for r in data:
            ws.append(r)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(header))}{len(data)+1}"
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    years = Counter(r["Award Year"] for r in rows)
    ws = wb.create_sheet("README")
    for line in [
        "Turkey EKAP contract lots, mapped to Mapping_Taxonomy.xlsx",
        f"Tender details on disk: {len(det)} (MedTech pull in progress + earlier IVD pull cache). Contract-lot rows: {len(rows)}; "
        f"by contract year: { {k: years[k] for k in sorted(years, key=lambda x: x or 0)} }.",
        f"Mapped lots: {len(mapped)}; Included (mapped, contract year 2024-26, not terminated): {len(inc)}.",
        "Contract value is per contract and may cover several lots: Line Value is filled only for single-lot contracts. Never sum Contract Value across lots.",
        "EKAP publishes no item quantities; Qty/Unit Price are filled only where the lot name states a quantity (e.g. '100 ADET').",
        "Company is filled only when the supplier's own name is the OEM or its Turkish affiliate (Manufacturer Evidence = Inferred from distributor).",
        "Lot names were translated with translate.googleapis.com (cached in turmeca/translation_cache.sqlite) and mapped TR-first, then EN.",
    ]:
        ws.append([line])
    ws.column_dimensions["A"].width = 160
    sheet("Mapped_Lots", cols, [[r[c] for c in cols] for r in mapped], [8, 8, 9, 30, 13, 40, 35, 12, 15, 11, 7, 18, 45, 45, 7, 26, 30, 36, 14, 18, 18, 35, 14, 14, 8, 12, 14, 10, 14, 18, 10, 6, 20])
    sheet("By_Category", ["OU", "Sub-OU", "Device Category", "Lots (incl.)", "2024", "2025", "2026", "Single-lot contracts", "Single-lot value TRY",
                          "Top suppliers (lots)", "OEM-affiliate companies (lots)"], cat_rows, [28, 35, 40, 10, 7, 7, 7, 10, 16, 80, 40])
    sheet("Suppliers_Mapped_Lots", ["Supplier/Distributor", "Included lots", "Company (if OEM affiliate)", "Top categories"], sup_rows, [70, 10, 22, 90])
    sheet("Unmapped_Top_Lots", ["Lot Name (TR)", "Lot Name (EN)", "Rows (2024-26)"], [[a, b, n] for (a, b), n in unm.most_common(3000)], [60, 60, 10])
    wb.save(OUT_X)
    print(f"mapped {len(mapped)} included {len(inc)}; single-lot priced included {sum(1 for r in inc if r['Line Value TRY'])}; "
          f"with qty {sum(1 for r in inc if r['Unit Price TRY'])}")


if __name__ == "__main__":
    main()
